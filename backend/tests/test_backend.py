from __future__ import annotations

import csv
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.service import Conflict, InvoiceService, Settings, ValidationFailure


def make_service(tmp_path: Path) -> InvoiceService:
    return InvoiceService(
        Settings(
            database_path=tmp_path / "invoices.sqlite3",
            source_dir=tmp_path / "sources",
            export_dir=tmp_path / "exports",
            workers=0,
        )
    )


def test_explicit_confirmation_learns_supplier_alias_and_approval_is_versioned(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    invoice = service.get_invoice("demo-invoice-002")

    with pytest.raises(ValidationFailure) as unresolved:
        service.approve(invoice["id"], invoice["version"])
    assert any(error["code"] == "unmapped" for error in unresolved.value.errors)

    line = invoice["lines"][0]
    wrong_supplier_line = {
        **line,
        "rms_item_id": "RMS-1001",
        "match_status": "confirmed",
        "confidence": 100,
    }
    with pytest.raises(ValidationFailure) as wrong_supplier:
        service.update_invoice(
            invoice["id"], invoice["version"], {"lines": [wrong_supplier_line]}
        )
    assert wrong_supplier.value.errors[0]["code"] == "supplier_item_mismatch"

    line.update(
        {
            "rms_item_id": "RMS-2001",
            "match_status": "confirmed",
            "confidence": 100,
        }
    )
    corrected = service.update_invoice(
        invoice["id"], invoice["version"], {"lines": [line]}
    )
    assert corrected["status"] == "needs_review"
    aliases = service.list_aliases()
    assert aliases["total"] == 1
    assert aliases["items"][0]["supplier_id"] == "DEMO-SECOND"
    assert aliases["items"][0]["uom"] == "ea"
    assert aliases["items"][0]["rms_item_id"] == "RMS-2001"

    approved = service.approve(corrected["id"], corrected["version"])
    assert approved["status"] == "ready"
    assert approved["version"] == corrected["version"] + 1

    with pytest.raises(Conflict) as stale:
        service.update_invoice(approved["id"], corrected["version"], {"po_number": "STALE"})
    assert stale.value.current_version == approved["version"]

    edited = service.update_invoice(
        approved["id"], approved["version"], {"po_number": "DEMO-PO-CORRECTED"}
    )
    assert edited["status"] == "needs_review"
    events = service.audit(approved["id"])["items"]
    assert {event["event_type"] for event in events} >= {
        "fictional_demo_seeded",
        "invoice_edited",
        "invoice_approved",
    }


def test_export_uses_one_approved_set_decimal_checks_and_formula_safe_cells(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    invoice = service.get_invoice("demo-invoice-001")
    line = invoice["lines"][0]
    line.update(
        {
            "description": "=2+2",
            "quantity": "3",
            "unit_price": "0.10",
            "line_total": "0.30",
            "tax_rate": "10",
            "rms_item_id": "RMS-1001",
            "match_status": "confirmed",
            "confidence": 100,
        }
    )
    edited = service.update_invoice(
        invoice["id"],
        invoice["version"],
        {
            "supplier_name": "+Fictional formula-like supplier",
            "document": "=2+2",
            "comment": "+Formula-like comment",
            "subtotal": "0.30",
            "tax_total": "0.03",
            "total": "0.33",
            "lines": [line],
        },
    )
    approved = service.approve(
        edited["id"],
        edited["version"],
        acknowledge_target_cost_variance=True,
    )

    with pytest.raises(ValidationFailure) as mixed_set:
        service.create_export([approved["id"], "demo-invoice-002"])
    assert any(error["code"] == "not_approved" for error in mixed_set.value.errors)
    assert service.list_exports()["items"] == []
    assert service.get_invoice(approved["id"])["status"] == "ready"

    exported = service.create_export([approved["id"]])
    workbook = load_workbook(exported["path"], data_only=False)
    assert workbook.sheetnames == ["Header", "Tax_Breakdown", "Details"]
    assert workbook["Header"]["B2"].value == "'=2+2"
    assert workbook["Header"]["M2"].value == "'+Formula-like comment"
    assert workbook["Header"]["H2"].value == pytest.approx(0.30)
    assert workbook["Header"]["I2"].value == pytest.approx(0.03)
    assert workbook["Details"]["D2"].value == pytest.approx(0.10)
    assert workbook["Details"]["E2"].value == 3
    assert workbook["Header"]["A2"].data_type == "s"
    assert workbook["Details"]["B2"].data_type == "s"
    assert service.get_invoice(approved["id"])["status"] == "exported"
    assert service.list_exports()["items"][0]["sha256"] == exported["sha256"]


def test_catalog_import_accepts_csv_and_xlsx_with_precise_costs(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    csv_result = service.import_catalog(
        "catalog.csv",
        b"rms_item_id,description,supplier_id,uom,unit_cost\n"
        b"CSV-1,=Formula-like catalog text,SUP-1,EA,0.10\n",
    )
    assert csv_result == {"imported": 1, "skipped": 0, "warnings": []}

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["item id", "item description", "vendor id", "unit", "cost"])
    sheet.append(["XLSX-1", "Workbook item", "SUP-2", "EA", "0.20"])
    output = io.BytesIO()
    workbook.save(output)
    xlsx_result = service.import_catalog("catalog.xlsx", output.getvalue())
    assert xlsx_result == {"imported": 1, "skipped": 0, "warnings": []}

    catalog = service.list_catalog(search=None, limit=20)
    assert catalog["total"] == 2
    by_id = {item["rms_item_id"]: item for item in catalog["items"]}
    assert by_id["CSV-1"]["description"] == "=Formula-like catalog text"
    assert by_id["CSV-1"]["unit_cost"] == pytest.approx(0.10)
    assert by_id["XLSX-1"]["unit_cost"] == pytest.approx(0.20)


def test_approval_rejects_inconsistent_line_arithmetic_even_when_headers_balance(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    invoice = service.get_invoice("demo-invoice-001")
    line = invoice["lines"][0]
    line.update(
        {
            "quantity": "3",
            "unit_price": "0.10",
            "line_total": "0.34",
            "rms_item_id": "RMS-1001",
            "match_status": "confirmed",
        }
    )
    edited = service.update_invoice(
        invoice["id"],
        invoice["version"],
        {"subtotal": "0.34", "tax_total": "0.03", "total": "0.37", "lines": [line]},
    )
    with pytest.raises(ValidationFailure) as mismatch:
        service.approve(edited["id"], edited["version"])
    assert any(error["code"] == "line_calculation_mismatch" for error in mismatch.value.errors)


def test_restart_recovers_processing_job_and_records_audit(tmp_path: Path) -> None:
    settings = Settings(
        database_path=tmp_path / "invoices.sqlite3",
        source_dir=tmp_path / "sources",
        export_dir=tmp_path / "exports",
        workers=0,
    )
    first = InvoiceService(settings)
    accepted, duplicate = first.ingest_bytes("restart.txt", b"unique queued document")
    assert not duplicate
    assert first._claim_job() == accepted["id"]
    assert first.get_invoice(accepted["id"])["status"] == "processing"

    restarted = InvoiceService(settings)
    assert restarted.start() == 1
    recovered = restarted.get_invoice(accepted["id"])
    assert recovered["status"] == "queued"
    assert restarted.audit(accepted["id"])["items"][0]["event_type"] == "processing_recovered"


def test_exception_csv_is_separate_safe_and_explains_noninvoice_and_unresolved(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)
    empty = service.create_exception_report()
    assert empty["count"] == 0
    assert empty["content"].decode("utf-8-sig").splitlines() == [
        "Invoice ID,Filename,Document Type,Status,Reasons"
    ]

    service.seed_demo()
    invoice = service.get_invoice("demo-invoice-002")
    changed = service.update_invoice(
        invoice["id"], invoice["version"], {"document_type": "delivery_note"}
    )
    queued, duplicate = service.ingest_bytes("=FORMULA.txt", b"safe exception source")
    assert not duplicate

    report = service.create_exception_report()
    rows = list(
        csv.DictReader(io.StringIO(report["content"].decode("utf-8-sig")))
    )
    assert list(rows[0]) == [
        "Invoice ID",
        "Filename",
        "Document Type",
        "Status",
        "Reasons",
    ]
    by_id = {row["Invoice ID"]: row for row in rows}
    review_reasons = by_id[changed["id"]]["Reasons"]
    assert "[document_type.not_invoice]" in review_reasons
    assert "[lines.0.rms_item_id.unmapped]" in review_reasons
    assert by_id[queued["id"]]["Filename"] == "'=FORMULA.txt"
    assert by_id[queued["id"]]["Reasons"] == "[status_queued] awaiting extraction"


def test_real_text_extraction_and_catalog_matching_flow_to_review(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    service.update_settings(
        1,
        {
            "brand_label": "Synthetic test brand",
            "location": "00001",
            "location_type": "Store (S)",
            "supplier_rules": [
                {
                    "supplier_id": "SUP-TXT",
                    "supplier_site": "00009",
                    "tax_code": "TEST05",
                }
            ],
            "target_cost_policy": {
                "mode": "invoice_only",
                "maximum_absolute_difference_aed": 10,
            },
        },
    )
    service.import_catalog(
        "catalog.csv",
        b"rms_item_id,description,supplier_id,uom,unit_cost\n"
        b"RMS-TXT-1,Fictional Widget 250ml,SUP-TXT,EA,10.00\n",
    )
    content = b"""Supplier name: Fictional Text Supplier
Supplier ID: SUP-TXT
Supplier site: TEST-SITE
Invoice number: TXT-INV-1
Invoice date: 2026-09-24
PO number: TXT-PO-1
Currency: AED

Item Description Qty UOM Unit Price Amount
Fictional Widget 250ml 2 EA 10.00 20.00

Subtotal 20.00
Tax 1.00
Total 21.00
"""
    uploaded, duplicate = service.ingest_bytes("invoice.txt", content)
    assert not duplicate
    assert service._claim_job() == uploaded["id"]
    service._process_job(uploaded["id"])

    invoice = service.get_invoice(uploaded["id"])
    assert invoice["status"] == "needs_review"
    assert invoice["extraction_method"] == "text"
    assert invoice["invoice_number"] == "TXT-INV-1"
    assert invoice["total"] == 21
    assert len(invoice["lines"]) == 1
    assert invoice["lines"][0]["rms_item_id"] == "RMS-TXT-1"
    assert invoice["lines"][0]["match_status"] == "auto"
    classified = service.update_invoice(
        invoice["id"], invoice["version"], {"document_type": "invoice"}
    )
    approved = service.approve(classified["id"], classified["version"])
    assert approved["status"] == "ready"


def test_one_thousand_distinct_jobs_persist_duplicate_and_paginate(tmp_path: Path) -> None:
    """Persistence volume test: no OCR is mocked or invoked; workers stay disabled."""
    settings = Settings(
        database_path=tmp_path / "invoices.sqlite3",
        source_dir=tmp_path / "sources",
        export_dir=tmp_path / "exports",
        workers=0,
    )
    service = InvoiceService(settings)
    first_content = b"synthetic queue document 0"
    first_id = None
    for index in range(1000):
        item, duplicate = service.ingest_bytes(
            f"batch-{index:04d}.txt", f"synthetic queue document {index}".encode()
        )
        assert not duplicate
        if index == 0:
            first_id = item["id"]

    duplicate_item, is_duplicate = service.ingest_bytes("renamed.txt", first_content)
    assert is_duplicate
    assert duplicate_item["id"] == first_id

    reopened = InvoiceService(settings)
    assert reopened.stats()["total"] == 1000
    assert reopened.stats()["queued"] == 1000
    page = reopened.list_invoices(status="queued", search=None, limit=25, offset=975)
    assert page["total"] == 1000
    assert page["limit"] == 25
    assert page["offset"] == 975
    assert len(page["items"]) == 25
    assert len(list((tmp_path / "sources").glob("*.txt"))) == 1000


def test_http_contract_upload_duplicate_and_safe_source_headers(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("INVOICE_DATA_DIR", str(tmp_path / "global-app"))
    from fastapi.testclient import TestClient

    from backend.app import create_app

    service_settings = Settings(
        database_path=tmp_path / "http.sqlite3",
        source_dir=tmp_path / "http-sources",
        export_dir=tmp_path / "http-exports",
        workers=0,
    )
    app = create_app(service_settings)
    contract = {
        (method, route.path)
        for route in app.routes
        for method in (getattr(route, "methods", None) or set())
    }
    assert {
        ("GET", "/api/health"),
        ("GET", "/api/stats"),
        ("GET", "/api/settings"),
        ("PUT", "/api/settings"),
        ("GET", "/api/invoices"),
        ("POST", "/api/invoices/upload"),
        ("GET", "/api/invoices/{invoice_id}"),
        ("GET", "/api/invoices/{invoice_id}/source"),
        ("PUT", "/api/invoices/{invoice_id}"),
        ("POST", "/api/invoices/{invoice_id}/approve"),
        ("POST", "/api/invoices/{invoice_id}/rematch"),
        ("POST", "/api/invoices/{invoice_id}/retry"),
        ("GET", "/api/invoices/{invoice_id}/audit"),
        ("GET", "/api/catalog"),
        ("POST", "/api/catalog/import"),
        ("GET", "/api/aliases"),
        ("POST", "/api/aliases/import"),
        ("POST", "/api/exports"),
        ("GET", "/api/exports"),
        ("GET", "/api/reports/exceptions.csv"),
        ("POST", "/api/demo"),
    } <= contract

    with TestClient(app) as client:
        health = client.get("/api/health")
        assert health.status_code == 200
        assert {"pdf", "png", "csv", "xlsx", "docx", "txt"} <= set(
            health.json()["supported_formats"]
        )
        brand_settings = client.get("/api/settings")
        assert brand_settings.status_code == 200
        assert brand_settings.json()["include_upc_in_export"] is False
        settings_payload = {
            **brand_settings.json(),
            "expected_version": brand_settings.json()["version"],
            "include_upc_in_export": True,
        }
        saved_settings = client.put("/api/settings", json=settings_payload)
        assert saved_settings.status_code == 200
        assert saved_settings.json()["include_upc_in_export"] is True
        first = client.post(
            "/api/invoices/upload",
            files=[("files", ("one.txt", b"API source content", "text/html"))],
        )
        assert first.status_code == 200
        assert len(first.json()["accepted"]) == 1
        invoice_id = first.json()["accepted"][0]["id"]
        duplicate = client.post(
            "/api/invoices/upload",
            files=[("files", ("renamed.txt", b"API source content", "text/plain"))],
        )
        assert duplicate.json()["duplicates"][0]["id"] == invoice_id
        rejected = client.post(
            "/api/invoices/upload",
            files=[("files", ("active.html", b"<script>bad()</script>", "text/html"))],
        )
        assert rejected.json()["rejected"][0]["filename"] == "active.html"
        source = client.get(f"/api/invoices/{invoice_id}/source")
        assert source.headers["content-type"].startswith("text/plain")
        assert source.headers["x-content-type-options"] == "nosniff"
        assert source.headers["content-security-policy"].startswith("sandbox")
        report = client.get("/api/reports/exceptions.csv")
        assert report.status_code == 200
        assert report.headers["content-type"].startswith("text/csv")
        assert report.headers["content-disposition"].startswith(
            'attachment; filename="invoice-exceptions-'
        )
        assert report.headers["x-exception-count"] == "1"
        report_rows = list(
            csv.DictReader(io.StringIO(report.content.decode("utf-8-sig")))
        )
        assert report_rows[0]["Invoice ID"] == invoice_id
        assert report_rows[0]["Status"] == "queued"
