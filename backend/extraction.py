"""Bounded, evidence-preserving invoice document extraction.

The extractor intentionally uses generic labels and table heuristics instead of
supplier templates.  Every returned value comes from source text or a
structured cell; missing values remain ``None``.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import shutil
import subprocess
import tempfile
import threading
import time
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from xml.etree import ElementTree


class DocumentExtractionError(RuntimeError):
    """A document could not be safely or credibly extracted."""

    retryable = False


class UnsupportedDocumentError(DocumentExtractionError):
    """The file type is outside the deliberately supported set."""


class ExtractionLimitError(DocumentExtractionError):
    """A document exceeded a configured resource bound."""


class ExtractionTimeBudgetError(ExtractionLimitError):
    """OCR ran out of its time budget.

    The time budget bounds the host's work, not the document's content: the
    same document can clear it on a quiet host and miss it on a busy one.  The
    failure is therefore retryable, unlike the content bounds (bytes, pages,
    pixels, rows), which stay plain ``ExtractionLimitError`` and are final.
    """

    retryable = True


_RETRY_HINT = "; try again when the host is quieter."


@dataclass(frozen=True)
class ExtractionLimits:
    """Hard parser limits.  Limits fail the whole document; nothing is truncated."""

    max_file_bytes: int = 50 * 1024 * 1024
    max_pages: int = 50
    max_pixels_per_page: int = 30_000_000
    max_total_pixels: int = 150_000_000
    max_spreadsheet_rows: int = 20_000
    max_spreadsheet_columns: int = 100
    max_archive_members: int = 10_000
    max_archive_uncompressed_bytes: int = 100 * 1024 * 1024
    max_archive_compression_ratio: int = 200
    max_embedded_xml_bytes: int = 10 * 1024 * 1024
    max_xml_elements: int = 100_000
    ocr_timeout_seconds_per_page: float = 45.0
    ocr_timeout_seconds_total: float = 180.0
    pdf_text_min_characters: int = 30


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
    ".csv",
    ".xlsx",
    ".txt",
    ".docx",
}

_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
_CURRENCIES = {"AED", "USD", "EUR", "GBP", "SAR", "QAR", "KWD", "BHD", "OMR"}
_EMPTY_HEADER = {
    "supplier_name": None,
    "supplier_id": None,
    "supplier_site": None,
    "invoice_number": None,
    "invoice_date": None,
    "po_number": None,
    "currency": None,
    "subtotal": None,
    "tax_total": None,
    "total": None,
}

_DOCUMENT_TYPES = {"invoice", "credit_note", "purchase_order", "delivery_note", "unknown"}


@dataclass(frozen=True)
class _ColumnAnchor:
    """One explicit table heading and its fixed-width horizontal position."""

    field: str
    start: int
    end: int
    priority: int
    basis: str | None = None


# Longer and more specific labels deliberately win over generic words.  These
# are document vocabulary, not supplier templates; a row is never accepted
# unless the same document also exposes description, quantity and value
# headings at distinct horizontal positions.
_COLUMN_PATTERNS: tuple[tuple[str, re.Pattern[str], int, str | None], ...] = (
    ("unit_price", re.compile(r"\bnet\s+unit\s+(?:price|cost|rate)\b", re.I), 130, "net"),
    ("unit_price", re.compile(r"\bnet\s+(?:price|cost|rate)\b", re.I), 125, "net"),
    (
        "line_total",
        re.compile(r"\b(?:amount|value)\s+(?:excl\.?|excluding|before)\s+(?:vat|tax)\b", re.I),
        140,
        "net",
    ),
    ("line_total", re.compile(r"\bnet\s+(?:amount|value|total)\b", re.I), 125, "net"),
    (
        "gross_line_total",
        re.compile(r"\b(?:amount|value)\s+(?:incl\.?|including|after)\s+(?:vat|tax)\b", re.I),
        115,
        "gross",
    ),
    (
        "gross_line_total",
        re.compile(r"\bgross(?:\s+(?:amount|value|total))?(?:\s+(?:AED|USD|EUR|GBP|SAR|QAR|KWD|BHD|OMR))?\b", re.I),
        100,
        "gross",
    ),
    ("unit_price", re.compile(r"\bunit\s+(?:price|cost|rate)\b", re.I), 115, "unit"),
    ("unit_price", re.compile(r"\bgross\s+(?:price|cost|rate)\b", re.I), 105, "gross"),
    ("line_total", re.compile(r"\b(?:line|extended)\s+(?:amount|total|value)\b", re.I), 115, "line"),
    ("line_total", re.compile(r"\btotal\s+(?:amount|value)\b", re.I), 105, "total"),
    (
        "line_total",
        re.compile(r"\bnet(?:\s+(?:amount|value|total))?(?:\s+(?:AED|USD|EUR|GBP|SAR|QAR|KWD|BHD|OMR))?\b", re.I),
        100,
        "net",
    ),
    ("tax_amount", re.compile(r"\b(?:vat|tax)[ \t]{1,3}(?:amount|value)\b", re.I), 120, "tax"),
    ("tax_rate", re.compile(r"\b(?:vat|tax)\s*(?:rate|%)", re.I), 120, "tax"),
    ("discount_amount", re.compile(r"\b(?:discount|disc\.?)\s+(?:amount|amt\.?|value)\b", re.I), 115, "discount"),
    ("discount_rate", re.compile(r"\b(?:discount|disc\.?)\s*(?:rate|%)", re.I), 120, "discount"),
    ("discount_amount", re.compile(r"\b(?:amt|ami)\b", re.I), 60, "discount"),
    ("discount_rate", re.compile(r"\b(?:disc|dis|die)\b", re.I), 55, "discount"),
    ("description", re.compile(r"\b(?:item|product|goods|material|article)\s+description\b", re.I), 130, None),
    ("row_number", re.compile(r"\b(?:sr\.?\s*(?:no\.?|#)|line\s*(?:no\.?|#))\b", re.I), 120, None),
    ("item_code", re.compile(r"\b(?:item|product|material|article)\s+(?:code|no\.?)\b", re.I), 125, None),
    ("upc", re.compile(r"\b(?:bar\s*code|barcode|ean|upc)\b", re.I), 125, None),
    ("quantity", re.compile(r"\b(?:invoice\s+)?(?:quantity|q?ty|oty)\b", re.I), 120, None),
    ("uom", re.compile(r"\b(?:uom|unit\s+of\s+measure)\b", re.I), 120, None),
    ("uom", re.compile(r"\b(?:pack(?:ing)?|packin|bckin|lackin|bekin)\b", re.I), 80, None),
    ("description", re.compile(r"\b(?:description|desc\.?|particulars?|details?)\b", re.I), 115, None),
    ("description", re.compile(r"\b(?:product|item)\b", re.I), 70, None),
    ("item_code", re.compile(r"\bsku\b", re.I), 110, None),
    ("item_code", re.compile(r"code\b", re.I), 65, None),
    ("uom", re.compile(r"\bunit\b(?!\s*(?:price|cost|rate))", re.I), 75, None),
    ("unit_price", re.compile(r"\b(?:price|rate|cost)\b", re.I), 70, "unspecified"),
    ("line_total", re.compile(r"\b(?:amount|value|total)\b", re.I), 65, "unspecified"),
    ("tax_rate", re.compile(r"\b(?:vat|tax)\s*%", re.I), 100, "tax"),
    ("discount_rate", re.compile(r"\b(?:discount|disc\.?)\s*%", re.I), 100, "discount"),
    ("tax_rate", re.compile(r"\b(?:vat|tax)\b", re.I), 55, "tax"),
    ("discount_rate", re.compile(r"\b(?:discount|disc\.?)\b", re.I), 55, "discount"),
)


def extract_document(
    path: str | Path,
    filename: str | None = None,
    *,
    limits: ExtractionLimits | None = None,
) -> dict[str, Any]:
    """Extract one invoice from a supported document.

    ``filename`` is used to determine the source type because upload services
    often store files under extensionless temporary names.  Parser failures and
    limit violations are explicit exceptions; a successful result is never a
    silently shortened document.
    """

    limits = limits or ExtractionLimits()
    source_path = Path(path)
    source_name = filename or source_path.name
    extension = Path(source_name).suffix.lower() or source_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedDocumentError(
            f"Unsupported document type {extension or '(no extension)'}. Supported types: {supported}."
        )
    if not source_path.exists() or not source_path.is_file():
        raise DocumentExtractionError(f"Document does not exist or is not a regular file: {source_path}")
    try:
        file_size = source_path.stat().st_size
    except OSError as exc:
        raise DocumentExtractionError(f"Could not inspect document: {exc}") from exc
    if file_size > limits.max_file_bytes:
        raise ExtractionLimitError(
            f"Document is {file_size} bytes; the limit is {limits.max_file_bytes} bytes."
        )

    try:
        if extension == ".pdf":
            evidence = _extract_pdf(source_path, limits)
        elif extension in _IMAGE_EXTENSIONS:
            evidence = _extract_image(source_path, limits)
        elif extension == ".csv":
            return _extract_csv(source_path, source_name, limits)
        elif extension == ".xlsx":
            return _extract_xlsx(source_path, source_name, limits)
        elif extension == ".txt":
            evidence = _extract_text_file(source_path)
        elif extension == ".docx":
            evidence = _extract_docx(source_path, limits)
        else:  # guarded by SUPPORTED_EXTENSIONS
            raise UnsupportedDocumentError(f"Unsupported document type: {extension}")
    except (DocumentExtractionError, ExtractionLimitError, UnsupportedDocumentError):
        raise
    except Exception as exc:  # normalize third-party parser errors
        raise DocumentExtractionError(f"Could not parse {source_name}: {exc}") from exc

    parsed = evidence.get("structured_invoice")
    if parsed is None:
        parsed = _parse_invoice_text(
            evidence["raw_text"],
            table_pages=evidence.get("table_pages"),
        )
    else:
        parsed = dict(parsed)
    warnings = _dedupe([*evidence.get("warnings", []), *parsed.pop("warnings", [])])
    if not parsed["lines"]:
        warnings.append("No credible invoice line rows were detected; manual review is required.")
    result = {
        **parsed,
        "source_filename": source_name,
        "raw_text": evidence["raw_text"],
        "extraction_method": evidence["extraction_method"],
        "pages": evidence["pages"],
        "page_texts": evidence["page_texts"],
        "warnings": _dedupe(warnings),
    }
    result["document_confidence"] = _document_confidence(result)
    return result


def _extract_pdf(path: Path, limits: ExtractionLimits) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:
        raise DocumentExtractionError("PDF extraction requires the 'pypdf' package.") from exc

    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            raise DocumentExtractionError(
                "Encrypted PDF files are not supported; provide a decrypted copy."
            )
        page_count = len(reader.pages)
    except DocumentExtractionError:
        raise
    except (PdfReadError, ValueError, OSError) as exc:
        raise DocumentExtractionError(f"The PDF is corrupt or unreadable: {exc}") from exc

    _check_page_count(page_count, limits)
    structured_invoice, structured_warnings = _extract_facturx_attachment(reader, limits)
    native_texts: list[str] = []
    native_layouts: list[str] = []
    needs_ocr: list[int] = []
    layout_failures: list[int] = []
    for index, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise DocumentExtractionError(f"Could not read PDF page {index + 1}: {exc}") from exc
        text = _clean_source_text(text)
        native_texts.append(text)
        if len(re.sub(r"\s+", "", text)) < limits.pdf_text_min_characters:
            needs_ocr.append(index)
            native_layouts.append("")
        else:
            try:
                layout_text = page.extract_text(
                    extraction_mode="layout",
                    layout_mode_space_vertically=False,
                ) or ""
                native_layouts.append(_clean_source_text(layout_text))
            except Exception:
                # Reading-order text still preserves source evidence.  Layout
                # extraction is an enhancement and must not turn a readable
                # invoice into a whole-document failure.
                native_layouts.append(text)
                layout_failures.append(index)

    page_evidence: list[dict[str, Any]] = []
    warnings: list[str] = list(structured_warnings)
    if layout_failures:
        pages = ", ".join(str(index + 1) for index in layout_failures)
        warnings.append(
            f"Fixed-position PDF layout extraction failed on page(s) {pages}; reading-order text was used for table review."
        )
    ocr_results: dict[int, tuple[str, str, float | None]] = {}
    if needs_ocr and structured_invoice is None:
        ocr_results = _ocr_pdf_pages(path, needs_ocr, limits)
        warnings.append(
            "OCR used English language data ('eng'); non-English text may be incomplete and requires review."
        )
    elif needs_ocr:
        warnings.append(
            "Image-only PDF page text was not OCR processed because a valid embedded Factur-X invoice supplied the structured fields."
        )

    for index, native_text in enumerate(native_texts):
        if index in ocr_results:
            ocr_text, ocr_layout_text, confidence = ocr_results[index]
            chosen = ocr_text if len(re.sub(r"\s+", "", ocr_text)) >= len(re.sub(r"\s+", "", native_text)) else native_text
            method = "ocr" if chosen == ocr_text else "pdf_text"
            if method == "ocr" and confidence is not None and confidence < 65:
                warnings.append(f"OCR confidence on page {index + 1} is low ({confidence:.0f}/100).")
        else:
            chosen, ocr_layout_text, confidence, method = native_text, native_text, None, "pdf_text"
        layout_text = ocr_layout_text if method == "ocr" else (native_layouts[index] or chosen)
        page_evidence.append(
            {
                "page": index + 1,
                "raw_text": chosen,
                "layout_text": layout_text,
                "method": method,
                "ocr_confidence": confidence,
            }
        )

    methods = {page["method"] for page in page_evidence}
    if structured_invoice is not None:
        extraction_method = "pdf_facturx_xml"
    elif methods == {"pdf_text"}:
        extraction_method = "pdf_text"
    elif methods == {"ocr"}:
        extraction_method = "pdf_ocr"
    else:
        extraction_method = "pdf_text+ocr"
    raw_text = "\n\n".join(page["raw_text"] for page in page_evidence)
    if not raw_text.strip():
        warnings.append("The document contains no readable text; manual review is required.")
    return {
        "raw_text": raw_text,
        "extraction_method": extraction_method,
        "pages": page_count,
        "page_texts": page_evidence,
        "table_pages": [
            {"page": page["page"], "text": page["layout_text"], "method": page["method"]}
            for page in page_evidence
        ],
        "structured_invoice": structured_invoice,
        "warnings": warnings,
    }


def _extract_facturx_attachment(
    reader: Any,
    limits: ExtractionLimits,
) -> tuple[dict[str, Any] | None, list[str]]:
    """Read a bounded Factur-X/ZUGFeRD XML attachment when one is present."""

    try:
        attachments = reader.attachments
    except Exception:
        return None, ["PDF embedded-file metadata could not be read; text extraction was used."]

    recognized_names = {"factur-x.xml", "zugferd-invoice.xml"}
    payloads: list[bytes] = []
    for name, values in attachments.items():
        if Path(str(name)).name.casefold() not in recognized_names:
            continue
        for value in values:
            if not isinstance(value, bytes):
                continue
            if len(value) > limits.max_embedded_xml_bytes:
                return None, [
                    "An embedded Factur-X XML file exceeded the configured size limit; PDF text was used."
                ]
            payloads.append(value)

    if not payloads:
        return None, []
    unique_payloads = {hashlib.sha256(payload).digest(): payload for payload in payloads}
    if len(unique_payloads) != 1:
        return None, [
            "Conflicting embedded Factur-X XML files were found; their values were withheld and PDF text was used."
        ]

    payload = next(iter(unique_payloads.values()))
    if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", payload, re.I):
        return None, [
            "Embedded Factur-X XML declared a document type or entity and was not parsed; PDF text was used."
        ]
    try:
        root = ElementTree.fromstring(payload)
    except ElementTree.ParseError:
        return None, ["Embedded Factur-X XML was malformed; PDF text was used."]
    if sum(1 for _ in root.iter()) > limits.max_xml_elements:
        return None, [
            "Embedded Factur-X XML exceeded the configured element limit; PDF text was used."
        ]
    try:
        parsed = _parse_facturx_tree(root)
    except (InvalidOperation, ValueError):
        return None, [
            "Embedded Factur-X XML contained invalid structured values; PDF text was used."
        ]
    if parsed is None:
        return None, [
            "An XML attachment used a Factur-X filename but did not contain a supported CrossIndustryInvoice; PDF text was used."
        ]
    return parsed, [
        "Structured invoice fields were read from the embedded Factur-X XML; the rendered PDF remains available for review."
    ]


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _xml_child(element: ElementTree.Element | None, name: str) -> ElementTree.Element | None:
    if element is None:
        return None
    return next((child for child in element if _xml_local_name(child.tag) == name), None)


def _xml_path(
    element: ElementTree.Element | None,
    *names: str,
) -> ElementTree.Element | None:
    current = element
    for name in names:
        current = _xml_child(current, name)
        if current is None:
            break
    return current


def _xml_descendant(element: ElementTree.Element | None, name: str) -> ElementTree.Element | None:
    if element is None:
        return None
    return next(
        (candidate for candidate in element.iter() if _xml_local_name(candidate.tag) == name),
        None,
    )


def _xml_text(element: ElementTree.Element | None) -> str | None:
    if element is None or element.text is None:
        return None
    value = element.text.strip()
    return value or None


def _xml_number(element: ElementTree.Element | None) -> float | None:
    return _parse_number(_xml_text(element))


def _facturx_date(element: ElementTree.Element | None) -> str | None:
    value = _xml_text(element)
    if value is None:
        return None
    format_code = str(element.attrib.get("format") or "")
    formats = {"102": "%Y%m%d", "203": "%Y%m%d%H%M"}
    date_format = formats.get(format_code)
    if date_format is None:
        return None
    try:
        return datetime.strptime(value, date_format).date().isoformat()
    except ValueError:
        return None


def _facturx_uom(code: str | None) -> str | None:
    if not code:
        return None
    aliases = {
        "C62": "EA",
        "H87": "PC",
        "KGM": "KG",
        "LTR": "L",
        "MLT": "ML",
    }
    cleaned = code.strip().upper()
    return aliases.get(cleaned, cleaned[:20] or None)


def _parse_facturx_tree(root: ElementTree.Element) -> dict[str, Any] | None:
    if _xml_local_name(root.tag) != "CrossIndustryInvoice":
        return None
    exchanged = _xml_child(root, "ExchangedDocument")
    transaction = _xml_child(root, "SupplyChainTradeTransaction")
    if exchanged is None or transaction is None:
        return None

    type_code = _xml_text(_xml_child(exchanged, "TypeCode"))
    if type_code == "380":
        document_type = "invoice"
    elif type_code == "381":
        document_type = "credit_note"
    else:
        document_type = "unknown"

    result: dict[str, Any] = dict(_EMPTY_HEADER)
    result["invoice_number"] = _xml_text(_xml_child(exchanged, "ID"))
    result["invoice_date"] = _facturx_date(
        _xml_path(exchanged, "IssueDateTime", "DateTimeString")
    )

    agreement = _xml_child(transaction, "ApplicableHeaderTradeAgreement")
    settlement = _xml_child(transaction, "ApplicableHeaderTradeSettlement")
    seller = _xml_child(agreement, "SellerTradeParty")
    result["supplier_name"] = _xml_text(_xml_child(seller, "Name"))
    tax_registration = _xml_child(seller, "SpecifiedTaxRegistration")
    result["supplier_id"] = _xml_text(_xml_child(tax_registration, "ID"))
    order_reference = _xml_path(agreement, "BuyerOrderReferencedDocument", "IssuerAssignedID")
    result["buyer_order_reference"] = _xml_text(order_reference)
    result["po_number"] = _explicit_po_token(result["buyer_order_reference"])
    result["currency"] = _parse_currency(
        _xml_text(_xml_child(settlement, "InvoiceCurrencyCode")) or ""
    )

    monetary = _xml_child(settlement, "SpecifiedTradeSettlementHeaderMonetarySummation")
    result["subtotal"] = _xml_number(_xml_child(monetary, "TaxBasisTotalAmount"))
    if result["subtotal"] is None:
        result["subtotal"] = _xml_number(_xml_child(monetary, "LineTotalAmount"))
    result["tax_total"] = _xml_number(_xml_child(monetary, "TaxTotalAmount"))
    result["total"] = _xml_number(_xml_child(monetary, "GrandTotalAmount"))
    if result["total"] is None:
        result["total"] = _xml_number(_xml_child(monetary, "DuePayableAmount"))

    lines: list[dict[str, Any]] = []
    rejected_lines = 0
    for line_element in transaction:
        if _xml_local_name(line_element.tag) != "IncludedSupplyChainTradeLineItem":
            continue
        product = _xml_child(line_element, "SpecifiedTradeProduct")
        line_agreement = _xml_child(line_element, "SpecifiedLineTradeAgreement")
        delivery = _xml_child(line_element, "SpecifiedLineTradeDelivery")
        line_settlement = _xml_child(line_element, "SpecifiedLineTradeSettlement")
        description = _xml_text(_xml_child(product, "Name")) or _xml_text(
            _xml_child(product, "Description")
        )
        quantity_element = _xml_child(delivery, "BilledQuantity")
        quantity = _xml_number(quantity_element)
        gross_price = _xml_number(
            _xml_path(line_agreement, "GrossPriceProductTradePrice", "ChargeAmount")
        )
        net_price = _xml_number(
            _xml_path(line_agreement, "NetPriceProductTradePrice", "ChargeAmount")
        )
        line_total = _xml_number(
            _xml_path(
                line_settlement,
                "SpecifiedTradeSettlementLineMonetarySummation",
                "LineTotalAmount",
            )
        )
        if not description or quantity is None or (net_price is None and gross_price is None and line_total is None):
            rejected_lines += 1
            continue

        derived_fields: list[str] = []
        unit_price = net_price
        unit_price_source = "facturx_net_price" if net_price is not None else None
        if unit_price is None and line_total is not None:
            unit_price = _divide_money_by_quantity(line_total, quantity)
            if unit_price is not None:
                net_price = unit_price
                unit_price_source = "derived_facturx_net_line_total_divided_by_quantity"
                derived_fields = ["unit_price", "net_unit_price"]
        if unit_price is None:
            unit_price = gross_price
            unit_price_source = "facturx_gross_price" if gross_price is not None else None

        global_id = _xml_child(product, "GlobalID")
        upc = None
        if global_id is not None and str(global_id.attrib.get("schemeID") or "") in {"0088", "0160"}:
            upc = _clean_upc(_xml_text(global_id))
        tax = _xml_child(line_settlement, "ApplicableTradeTax")
        line = {
            "description": description,
            "quantity": quantity,
            "uom": _facturx_uom(quantity_element.attrib.get("unitCode") if quantity_element is not None else None),
            "unit_price": unit_price,
            "net_unit_price": net_price,
            "gross_unit_price": gross_price,
            "printed_unit_price": net_price if net_price is not None else gross_price,
            "printed_unit_price_basis": "net" if net_price is not None else "gross",
            "unit_price_basis": "net" if net_price is not None else "gross",
            "unit_price_source": unit_price_source,
            "line_total": line_total,
            "line_total_basis": "net",
            "tax_rate": _xml_number(_xml_child(tax, "RateApplicablePercent")),
            "tax_amount": None,
            "discount_rate": None,
            "discount_amount": None,
            "upc": upc,
            "item_code": _clean_item_code(_xml_text(_xml_child(product, "SellerAssignedID"))),
            "source_page": None,
            "source_rows": None,
            "source_attachment": "factur-x.xml",
            "extraction_method": "facturx_xml",
            "derived_fields": derived_fields,
            "confidence": 0.99,
        }
        line["id"] = _line_id(len(lines), line)
        lines.append(line)

    provenance = {
        field: "facturx_xml"
        for field, value in result.items()
        if value is not None
    }
    warnings: list[str] = []
    if rejected_lines:
        warnings.append(
            f"{rejected_lines} Factur-X line item(s) lacked required structured values and require review."
        )
    if not lines:
        warnings.append("Factur-X XML contained no complete invoice line items; manual review is required.")
    if document_type == "unknown":
        warnings.append("Factur-X document type was not a supported invoice or credit note; route it to review.")
    return {
        **result,
        "field_confidence": {field: 0.99 for field in provenance},
        "field_provenance": provenance,
        "lines": lines,
        "document_type": document_type,
        "document_type_confidence": 0.99 if document_type != "unknown" else 0.0,
        "document_type_evidence": [f"facturx:type_code:{type_code or 'missing'}"],
        "table_extraction_method": "facturx_xml",
        "table_diagnostics": {
            "recognized_table_sections": 1,
            "sections_without_rows": int(not lines),
            "possible_unparsed_rows": rejected_lines,
            "emitted_rows": len(lines),
            "joined_continuation_lines": 0,
            "legacy_fallback": False,
            "header_schemas": ["facturx_xml"],
            "paired_adjustment_rows": 0,
            "unpaired_adjustment_rows": 0,
            "source_data_rows": len(lines) + rejected_lines,
        },
        "warnings": warnings,
    }


def _explicit_po_token(reference: Any) -> str | None:
    if reference in (None, ""):
        return None
    value = unicodedata.normalize("NFKC", str(reference)).strip()
    token = r"([A-Z0-9][A-Z0-9._/-]*)"
    patterns = (
        rf"\bP\.?\s*O\.?(?!\s*BOX\b)\s*(?:NUMBER|NO\.?|#)\s*[:#-]?\s*{token}",
        rf"\bP\.?\s*O\.?(?!\s*BOX\b)\s*[:#-]+\s*{token}",
        rf"^\s*P\.?\s*O\.?(?!\s*BOX\b)\s+{token}\s*$",
    )
    for pattern in patterns:
        match = re.search(pattern, value, re.I)
        if match:
            return match.group(1)
    return None


# PDFium is not thread-safe.  Extraction workers are threads in one process,
# so every pypdfium2 call (load, page access, render, close) is serialised
# here and no PDFium-owned memory is accessed outside the lock; Tesseract
# runs outside the lock on a copied image so OCR still parallelises.
_PDFIUM_LOCK = threading.Lock()


def _ocr_pdf_pages(
    path: Path, page_indices: Sequence[int], limits: ExtractionLimits
) -> dict[int, tuple[str, str, float | None]]:
    if shutil.which("tesseract") is None:
        raise DocumentExtractionError(
            "This PDF contains pages without usable embedded text, and the native 'tesseract' OCR binary is not installed."
        )
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise DocumentExtractionError(
            "Scanned PDF extraction requires the 'pypdfium2' package for rasterization."
        ) from exc

    started = time.monotonic()
    total_pixels = 0
    results: dict[int, tuple[str, str, float | None]] = {}
    # Multi-pass OCR is bounded to short documents.  On large scanned packets,
    # spending the document-wide timeout on optional second and third passes
    # can prevent later pages from receiving the original PSM 6 pass at all.
    # Keeping the baseline-only path for larger packets preserves complete,
    # reviewable page coverage within the same resource limits.
    use_enhancement_passes = len(page_indices) <= 12
    try:
        with _PDFIUM_LOCK:
            document = pdfium.PdfDocument(str(path))
        for index in page_indices:
            page_started = time.monotonic()
            remaining = limits.ocr_timeout_seconds_total - (time.monotonic() - started)
            if remaining <= 0:
                raise ExtractionTimeBudgetError(
                    f"OCR exceeded the {limits.ocr_timeout_seconds_total:.0f} second document limit"
                    + _RETRY_HINT
                )
            with _PDFIUM_LOCK:
                page = document[index]
                width, height = page.get_size()
            render_scale = 2.0
            pixels = int(width * render_scale) * int(height * render_scale)
            _check_page_pixels(pixels, index + 1, limits)
            total_pixels += pixels
            if total_pixels > limits.max_total_pixels:
                raise ExtractionLimitError(
                    f"PDF rasterization would exceed the {limits.max_total_pixels} total pixel limit."
                )
            with _PDFIUM_LOCK:
                bitmap = page.render(scale=render_scale)
                # to_pil() shares memory with the PDFium-owned bitmap buffer;
                # copy so nothing outside the lock touches PDFium memory.
                image = bitmap.to_pil().copy()
            page_budget = min(limits.ocr_timeout_seconds_per_page, remaining)
            budget_message = _ocr_budget_message(page_budget, limits)
            # Preserve the former OCR path and its full timeout first.  The
            # enhancement passes may use only time that remains; a slow
            # enhancement can never turn a formerly readable page into a
            # whole-document timeout.
            baseline_text, baseline_confidence = _run_tesseract(
                image,
                page_budget,
                limits.max_pixels_per_page,
                psm=6,
                autocontrast=False,
                dpi=round(72 * render_scale),
                timeout_message=budget_message,
            )
            enhanced_text, enhanced_confidence = baseline_text, baseline_confidence
            page_remaining = page_budget - (time.monotonic() - page_started)
            enhancement_completed = False
            if use_enhancement_passes and page_remaining > 2:
                try:
                    enhanced_text, enhanced_confidence = _run_tesseract(
                        image,
                        max(1.0, page_remaining * 0.62),
                        limits.max_pixels_per_page,
                        psm=4,
                        autocontrast=True,
                        dpi=round(72 * render_scale),
                        timeout_message=budget_message,
                    )
                    enhancement_completed = True
                except ExtractionLimitError:
                    enhanced_text, enhanced_confidence = baseline_text, baseline_confidence
            if _ocr_candidate_score(enhanced_text, enhanced_confidence) >= _ocr_candidate_score(
                baseline_text, baseline_confidence
            ):
                layout_text, confidence = enhanced_text, enhanced_confidence
            else:
                layout_text, confidence = baseline_text, baseline_confidence
            page_remaining = page_budget - (time.monotonic() - page_started)
            sparse_text = ""
            if enhancement_completed and page_remaining > 1:
                try:
                    sparse_text, _ = _run_tesseract(
                        image,
                        page_remaining,
                        limits.max_pixels_per_page,
                        psm=11,
                        autocontrast=True,
                        dpi=round(72 * render_scale),
                        timeout_message=budget_message,
                    )
                except ExtractionLimitError:
                    sparse_text = ""
            # Sparse segmentation is useful for isolated header labels and
            # values, while PSM 4 preserves the row geometry needed for table
            # parsing.  Keep the evidence streams separate and put the sparse
            # header evidence first for conservative labelled-field matching.
            raw_text = "\n\n".join(
                part for part in (sparse_text, layout_text) if part.strip()
            )
            results[index] = (raw_text, layout_text, confidence)
            try:
                image.close()
                with _PDFIUM_LOCK:
                    bitmap.close()
                    page.close()
            except Exception:
                pass
    except (DocumentExtractionError, ExtractionLimitError):
        raise
    except Exception as exc:
        raise DocumentExtractionError(f"Could not rasterize scanned PDF pages: {exc}") from exc
    finally:
        try:
            with _PDFIUM_LOCK:
                document.close()  # type: ignore[possibly-undefined]
        except Exception:
            pass
    return results


def _ocr_candidate_score(text: str, confidence: float | None) -> float:
    """Choose between source OCR passes without document-specific values."""

    normalized = _normalized_words(text)
    evidence_terms = (
        "invoice",
        "invoice no",
        "description",
        "qty",
        "quantity",
        "unit price",
        "gross",
        "net",
        "subtotal",
        "total",
        "currency",
    )
    evidence = sum(term in normalized for term in evidence_terms)
    numeric_rows = sum(
        len(re.findall(r"\d[\d,.]*", line)) >= 3 for line in text.splitlines()
    )
    header_strength = 0
    for _, _, anchors in _find_table_headers(text.splitlines()):
        fields = {anchor.field for anchor in anchors}
        header_strength = max(
            header_strength,
            len(fields) * 2 + (5 if "unit_price" in fields else 0),
        )
    return (
        float(confidence or 0)
        + evidence * 3
        + min(numeric_rows, 20) * 0.25
        + header_strength
    )


def _extract_image(path: Path, limits: ExtractionLimits) -> dict[str, Any]:
    if shutil.which("tesseract") is None:
        raise DocumentExtractionError(
            "Image extraction requires the native 'tesseract' OCR binary, which is not installed."
        )
    try:
        from PIL import Image, ImageOps, ImageSequence, UnidentifiedImageError
    except ImportError as exc:
        raise DocumentExtractionError("Image extraction requires the 'Pillow' package.") from exc

    started = time.monotonic()
    page_evidence: list[dict[str, Any]] = []
    total_pixels = 0
    try:
        with Image.open(path) as source:
            page_count = int(getattr(source, "n_frames", 1))
            _check_page_count(page_count, limits)
            for index, frame in enumerate(ImageSequence.Iterator(source)):
                frame = ImageOps.exif_transpose(frame.copy())
                pixels = frame.width * frame.height
                _check_page_pixels(pixels, index + 1, limits)
                total_pixels += pixels
                if total_pixels > limits.max_total_pixels:
                    raise ExtractionLimitError(
                        f"Image frames exceed the {limits.max_total_pixels} total pixel limit."
                    )
                remaining = limits.ocr_timeout_seconds_total - (time.monotonic() - started)
                if remaining <= 0:
                    raise ExtractionTimeBudgetError(
                        f"OCR exceeded the {limits.ocr_timeout_seconds_total:.0f} second document limit"
                        + _RETRY_HINT
                    )
                page_budget = min(limits.ocr_timeout_seconds_per_page, remaining)
                text, confidence = _run_tesseract(
                    frame,
                    page_budget,
                    limits.max_pixels_per_page,
                    psm=4,
                    autocontrast=True,
                    timeout_message=_ocr_budget_message(page_budget, limits),
                )
                page_evidence.append(
                    {
                        "page": index + 1,
                        "raw_text": text,
                        "layout_text": text,
                        "method": "ocr",
                        "ocr_confidence": confidence,
                    }
                )
                frame.close()
    except (DocumentExtractionError, ExtractionLimitError):
        raise
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise DocumentExtractionError(f"The image is corrupt or unreadable: {exc}") from exc

    warnings = [
        "OCR used English language data ('eng'); non-English text may be incomplete and requires review."
    ]
    for page in page_evidence:
        confidence = page["ocr_confidence"]
        if confidence is not None and confidence < 65:
            warnings.append(
                f"OCR confidence on page {page['page']} is low ({confidence:.0f}/100)."
            )
    return {
        "raw_text": "\n\n".join(page["raw_text"] for page in page_evidence),
        "extraction_method": "image_ocr",
        "pages": len(page_evidence),
        "page_texts": page_evidence,
        "table_pages": [
            {"page": page["page"], "text": page["layout_text"], "method": page["method"]}
            for page in page_evidence
        ],
        "warnings": warnings,
    }


def _ocr_budget_message(page_budget: float, limits: ExtractionLimits) -> str:
    """Name the budget an OCR pass ran under.

    A page runs under ``min(per-page limit, remaining document budget)``.  When
    the document budget clips the page, the number that ran out is the
    document limit; reporting the clipped remainder ("exceeded its 0 second
    limit") is the 180 second document budget wearing another face.
    """
    if page_budget < limits.ocr_timeout_seconds_per_page:
        return (
            f"OCR exceeded the {limits.ocr_timeout_seconds_total:.0f} second document limit."
        )
    return f"OCR page exceeded its {limits.ocr_timeout_seconds_per_page:.0f} second limit."


def _run_tesseract(
    image: Any,
    timeout: float,
    max_pixels: int,
    *,
    psm: int = 6,
    autocontrast: bool = False,
    dpi: int | None = None,
    timeout_message: str | None = None,
) -> tuple[str, float | None]:
    from PIL import ImageOps

    prepared = ImageOps.grayscale(image)
    if autocontrast:
        contrasted = ImageOps.autocontrast(prepared, cutoff=1)
        prepared.close()
        prepared = contrasted
    # A minimum resolution materially improves OCR of screenshots and small scans.
    if prepared.width < 1200:
        desired_scale = min(3.0, 1200 / max(1, prepared.width))
        pixel_scale = math.sqrt(max_pixels / max(1, prepared.width * prepared.height))
        scale = min(desired_scale, pixel_scale)
        if scale > 1.05:
            original = prepared
            prepared = prepared.resize(
                (int(prepared.width * scale), int(prepared.height * scale))
            )
            if original is not image:
                original.close()
    # Tesseract 5.5 on Debian can take tens of seconds to decode a mode-L PNG
    # while the equivalent RGB PNG completes in under two seconds.  Keep the
    # grayscale pixels used for OCR, but store them in an RGB container.
    rgb_prepared = prepared.convert("RGB")
    if prepared is not image:
        prepared.close()
    prepared = rgb_prepared
    temp_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
            temp_name = handle.name
        prepared.save(temp_name, format="PNG", optimize=False)
        try:
            command = [
                "tesseract",
                temp_name,
                "stdout",
                "-l",
                "eng",
            ]
            if dpi is not None:
                command.extend(["--dpi", str(dpi)])
            command.extend(["--psm", str(psm), "tsv"])
            completed = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=max(1.0, timeout),
            )
        except subprocess.TimeoutExpired as exc:
            # main names the budget that actually ran out (the document budget can
            # clip the page); PR #24 makes a missed time budget retryable.  Keep
            # both: main's message, stripped of its full stop, plus the hint.
            raise ExtractionTimeBudgetError(
                (timeout_message or f"OCR page exceeded its {timeout:.0f} second limit.").rstrip(".")
                + _RETRY_HINT
            ) from exc
        if completed.returncode != 0:
            detail = (completed.stderr or "unknown Tesseract error").strip().replace("\n", " ")[:500]
            raise DocumentExtractionError(f"Tesseract OCR failed: {detail}")
        return _text_from_tesseract_tsv(completed.stdout)
    finally:
        if temp_name:
            try:
                Path(temp_name).unlink(missing_ok=True)
            except OSError:
                pass
        try:
            if prepared is not image:
                prepared.close()
        except Exception:
            pass


def _text_from_tesseract_tsv(tsv_text: str) -> tuple[str, float | None]:
    try:
        rows = list(csv.DictReader(io.StringIO(tsv_text), delimiter="\t"))
    except csv.Error as exc:
        raise DocumentExtractionError(f"Tesseract returned invalid TSV output: {exc}") from exc

    groups: dict[tuple[int, int, int, int], list[dict[str, Any]]] = {}
    confidence_values: list[tuple[float, int]] = []
    character_widths: list[float] = []
    for row in rows:
        word = (row.get("text") or "").strip()
        if not word or row.get("level") != "5":
            continue
        try:
            key = tuple(int(row[name]) for name in ("page_num", "block_num", "par_num", "line_num"))
            item = {
                "text": word,
                "left": int(row["left"]),
                "top": int(row["top"]),
                "width": int(row["width"]),
                "height": int(row["height"]),
            }
            confidence = float(row.get("conf", "-1"))
        except (KeyError, TypeError, ValueError):
            continue
        groups.setdefault(key, []).append(item)
        character_widths.append(max(1.0, item["width"] / max(1, len(word))))
        if confidence >= 0:
            confidence_values.append((confidence, len(word)))

    lines: list[tuple[int, int, str]] = []
    ordered_widths = sorted(character_widths)
    median_character_width = (
        ordered_widths[len(ordered_widths) // 2] if ordered_widths else 8.0
    )
    for words in groups.values():
        words.sort(key=lambda word: word["left"])
        pieces: list[str] = []
        cursor = 0
        for word in words:
            target = max(0, int(round(word["left"] / median_character_width)))
            pieces.append(" " * max(1, min(120, target - cursor)))
            pieces.append(word["text"])
            cursor = target + len(word["text"])
        lines.append((min(word["top"] for word in words), min(word["left"] for word in words), "".join(pieces)))
    lines.sort(key=lambda item: (item[0], item[1]))
    text = "\n".join(item[2] for item in lines)
    if confidence_values:
        denominator = sum(weight for _, weight in confidence_values)
        confidence = sum(value * weight for value, weight in confidence_values) / denominator
    else:
        confidence = None
    return _clean_source_text(text), round(confidence, 1) if confidence is not None else None


def _extract_text_file(path: Path) -> dict[str, Any]:
    data = path.read_bytes()
    text = _decode_text(data)
    cleaned = _clean_source_text(text)
    return {
        "raw_text": cleaned,
        "extraction_method": "text",
        "pages": 1,
        "page_texts": [{"page": 1, "raw_text": cleaned, "layout_text": cleaned, "method": "text", "ocr_confidence": None}],
        "table_pages": [{"page": 1, "text": cleaned, "method": "text"}],
        "warnings": [],
    }


def _extract_docx(path: Path, limits: ExtractionLimits) -> dict[str, Any]:
    _check_zip_archive(path, limits)
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentExtractionError("DOCX extraction requires the 'python-docx' package.") from exc
    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentExtractionError(f"The DOCX file is corrupt or unreadable: {exc}") from exc
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        table_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not table_rows:
            continue
        column_count = max(len(row) for row in table_rows)
        widths = [
            max((len(row[index]) if index < len(row) else 0) for row in table_rows)
            for index in range(column_count)
        ]
        for row in table_rows:
            parts.append(
                "  ".join(
                    (row[index] if index < len(row) else "").ljust(widths[index])
                    for index in range(column_count)
                ).rstrip()
            )
    text = _clean_source_text("\n".join(parts))
    return {
        "raw_text": text,
        "extraction_method": "docx_text",
        "pages": 1,
        "page_texts": [{"page": 1, "raw_text": text, "layout_text": text, "method": "docx_text", "ocr_confidence": None}],
        "table_pages": [{"page": 1, "text": text, "method": "docx_text"}],
        "warnings": ["Embedded images in DOCX files are not OCR processed."],
    }


def _extract_csv(path: Path, source_name: str, limits: ExtractionLimits) -> dict[str, Any]:
    text = _decode_text(path.read_bytes())
    try:
        sample = text[:8192]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error as exc:
        raise DocumentExtractionError(f"The CSV is malformed: {exc}") from exc
    return _extract_structured_rows(rows, source_name, "csv_structured", limits)


def _extract_xlsx(path: Path, source_name: str, limits: ExtractionLimits) -> dict[str, Any]:
    _check_zip_archive(path, limits)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentExtractionError("XLSX extraction requires the 'openpyxl' package.") from exc
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise DocumentExtractionError(f"The XLSX file is corrupt or unreadable: {exc}") from exc
    try:
        nonempty_sheets: list[tuple[str, list[list[Any]]]] = []
        for worksheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if row_index > limits.max_spreadsheet_rows:
                    raise ExtractionLimitError(
                        f"Worksheet '{worksheet.title}' exceeds the {limits.max_spreadsheet_rows} row limit."
                    )
                values = list(row)
                if len(values) > limits.max_spreadsheet_columns and any(
                    value not in (None, "") for value in values[limits.max_spreadsheet_columns :]
                ):
                    raise ExtractionLimitError(
                        f"Worksheet '{worksheet.title}' exceeds the {limits.max_spreadsheet_columns} column limit."
                    )
                values = values[: limits.max_spreadsheet_columns]
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            if rows:
                nonempty_sheets.append((worksheet.title, rows))
    finally:
        workbook.close()
    if not nonempty_sheets:
        raise DocumentExtractionError("The XLSX workbook has no non-empty cells.")
    if len(nonempty_sheets) > 1:
        names = ", ".join(name for name, _ in nonempty_sheets)
        raise DocumentExtractionError(
            f"The workbook has multiple non-empty sheets ({names}); split it into one invoice per file."
        )
    return _extract_structured_rows(nonempty_sheets[0][1], source_name, "xlsx_structured", limits)


_COLUMN_ALIASES = {
    "invoicenumber": "invoice_number",
    "invoiceno": "invoice_number",
    "invoice": "invoice_number",
    "invoicedate": "invoice_date",
    "suppliername": "supplier_name",
    "vendorname": "supplier_name",
    "supplier": "supplier_name",
    "vendor": "supplier_name",
    "supplierid": "supplier_id",
    "vendorid": "supplier_id",
    "suppliersite": "supplier_site",
    "vendorsite": "supplier_site",
    "ponumber": "po_number",
    "purchaseordernumber": "po_number",
    "po": "po_number",
    "currency": "currency",
    "description": "description",
    "itemdescription": "description",
    "productdescription": "description",
    "itemname": "description",
    "quantity": "quantity",
    "qty": "quantity",
    "unitprice": "unit_price",
    "unitcost": "unit_price",
    "price": "unit_price",
    "netunitprice": "net_unit_price",
    "netunitcost": "net_unit_price",
    "grossunitprice": "gross_unit_price",
    "linetotal": "line_total",
    "lineamount": "line_total",
    "extendedamount": "line_total",
    "taxrate": "tax_rate",
    "vatrate": "tax_rate",
    "discount": "discount_rate",
    "discountrate": "discount_rate",
    "discountamount": "discount_amount",
    "taxamount": "tax_amount",
    "uom": "uom",
    "unitofmeasure": "uom",
    "upc": "upc",
    "barcode": "upc",
    "ean": "upc",
    "itemcode": "item_code",
    "productcode": "item_code",
    "sku": "item_code",
    "subtotal": "subtotal",
    "taxtotal": "tax_total",
    "vattotal": "tax_total",
    "invoiceamount": "total",
    "invoicetotal": "total",
    "grandtotal": "total",
    "total": "total",
}


def _extract_structured_rows(
    rows: Sequence[Sequence[Any]],
    source_name: str,
    method: str,
    limits: ExtractionLimits,
) -> dict[str, Any]:
    if len(rows) > limits.max_spreadsheet_rows:
        raise ExtractionLimitError(
            f"Spreadsheet has {len(rows)} rows; the limit is {limits.max_spreadsheet_rows}."
        )
    if any(len(row) > limits.max_spreadsheet_columns for row in rows):
        raise ExtractionLimitError(
            f"Spreadsheet exceeds the {limits.max_spreadsheet_columns} column limit."
        )

    header_index = -1
    canonical_headers: list[str | None] = []
    for index, row in enumerate(rows[:25]):
        candidate = [_canonical_column(value) for value in row]
        recognized = [value for value in candidate if value]
        if "description" in recognized and len(set(recognized)) >= 2:
            header_index, canonical_headers = index, candidate
            break

    raw_text = "\n".join(
        "  ".join("" if value is None else str(value) for value in row).rstrip() for row in rows
    )
    if header_index < 0:
        evidence = {
            "raw_text": _clean_source_text(raw_text),
            "extraction_method": method.replace("structured", "layout"),
            "pages": 1,
            "page_texts": [{"page": 1, "raw_text": _clean_source_text(raw_text), "layout_text": _clean_source_text(raw_text), "method": method, "ocr_confidence": None}],
            "warnings": ["No structured column header was found; layout text heuristics were used."],
        }
        parsed = _parse_invoice_text(
            evidence["raw_text"],
            table_pages=[{"page": 1, "text": evidence["raw_text"], "method": method}],
        )
        warnings = _dedupe([*evidence["warnings"], *parsed.pop("warnings", [])])
        if not parsed["lines"]:
            warnings.append("No credible invoice line rows were detected; manual review is required.")
        result = {
            **parsed,
            "source_filename": source_name,
            "raw_text": evidence["raw_text"],
            "extraction_method": evidence["extraction_method"],
            "pages": 1,
            "page_texts": evidence["page_texts"],
            "warnings": warnings,
        }
        result["document_confidence"] = _document_confidence(result)
        return result

    records: list[dict[str, Any]] = []
    for row in rows[header_index + 1 :]:
        record: dict[str, Any] = {}
        for column_index, canonical in enumerate(canonical_headers):
            if canonical and column_index < len(row):
                value = row[column_index]
                if value not in (None, ""):
                    record[canonical] = value
        if record:
            records.append(record)
    if not records:
        raise DocumentExtractionError("The spreadsheet header has no data rows.")

    _reject_mixed_invoices(records)
    result: dict[str, Any] = dict(_EMPTY_HEADER)
    warnings: list[str] = []
    field_confidence: dict[str, float] = {}
    for field in _EMPTY_HEADER:
        values = _unique_nonempty(record.get(field) for record in records)
        if len(values) == 1:
            value = values[0]
            if field in {"subtotal", "tax_total", "total"}:
                result[field] = _parse_number(value)
            elif field == "invoice_date":
                parsed_date, date_warning = _parse_date(str(value))
                result[field] = parsed_date
                if date_warning:
                    warnings.append(date_warning)
            elif field == "currency":
                result[field] = _parse_currency(str(value))
            else:
                result[field] = str(value).strip()
            if result[field] is not None:
                field_confidence[field] = 0.99
        elif len(values) > 1:
            # Identity conflicts were rejected above.  Conflicting totals are
            # withheld rather than guessed (often they are accidental line totals).
            warnings.append(f"Conflicting values for {field} were withheld.")

    lines: list[dict[str, Any]] = []
    for record in records:
        description = str(record.get("description", "")).strip()
        if not description or _is_summary_label(description):
            continue
        structured_net_price = _parse_number(record.get("net_unit_price"))
        structured_gross_price = _parse_number(record.get("gross_unit_price"))
        structured_printed_price = _parse_number(record.get("unit_price"))
        structured_unit_price = structured_net_price
        structured_price_basis = "net" if structured_net_price is not None else "unspecified"
        if structured_unit_price is None:
            structured_unit_price = structured_printed_price
        if structured_unit_price is None:
            structured_unit_price = structured_gross_price
            structured_price_basis = "gross" if structured_gross_price is not None else None
        line = {
            "description": description,
            "quantity": _parse_number(record.get("quantity")),
            "unit_price": structured_unit_price,
            "net_unit_price": structured_net_price,
            "gross_unit_price": structured_gross_price,
            "printed_unit_price": structured_printed_price,
            "printed_unit_price_basis": "unspecified" if structured_printed_price is not None else None,
            "unit_price_basis": structured_price_basis,
            "unit_price_source": "structured_source_column" if structured_unit_price is not None else None,
            "line_total": _parse_number(record.get("line_total")),
            "tax_rate": _parse_percentage(record.get("tax_rate")),
            "uom": _clean_uom(record.get("uom")),
            "upc": _clean_upc(record.get("upc")),
            "item_code": _clean_item_code(record.get("item_code")),
            "discount_rate": _parse_percentage(record.get("discount_rate")),
            "discount_amount": _parse_number(record.get("discount_amount")),
            "tax_amount": _parse_number(record.get("tax_amount")),
            "source_page": 1,
            "source_rows": [header_index + 2 + len(lines), header_index + 2 + len(lines)],
            "extraction_method": "structured_columns",
            "derived_fields": [],
            "confidence": 0.99,
        }
        line["id"] = _line_id(len(lines), line)
        lines.append(line)
    if not lines:
        warnings.append("No credible invoice line rows were detected; manual review is required.")

    if result.get("invoice_number") and lines:
        structured_type = "invoice"
        structured_type_confidence = 0.95
        structured_type_evidence = ["structured:invoice_number", "structured:invoice_rows"]
    else:
        (
            structured_type,
            structured_type_confidence,
            structured_type_evidence,
            structured_type_warning,
        ) = _classify_document_type(raw_text)
        if structured_type_warning:
            warnings.append(structured_type_warning)

    result.update(
        {
            "field_confidence": field_confidence,
            "lines": lines,
            "document_type": structured_type,
            "document_type_confidence": structured_type_confidence,
            "document_type_evidence": structured_type_evidence,
            "table_extraction_method": "structured_columns",
            "table_diagnostics": {
                "recognized_table_sections": 1,
                "sections_without_rows": 0,
                "possible_unparsed_rows": 0,
                "emitted_rows": len(lines),
                "joined_continuation_lines": 0,
                "legacy_fallback": False,
                "header_schemas": ["structured_columns"],
                "paired_adjustment_rows": 0,
                "unpaired_adjustment_rows": 0,
                "source_data_rows": len(lines),
            },
            "source_filename": source_name,
            "raw_text": _clean_source_text(raw_text),
            "extraction_method": method,
            "pages": 1,
            "page_texts": [
                {
                    "page": 1,
                    "raw_text": _clean_source_text(raw_text),
                    "layout_text": _clean_source_text(raw_text),
                    "method": method,
                    "ocr_confidence": None,
                }
            ],
            "warnings": _dedupe(warnings),
        }
    )
    result["document_confidence"] = _document_confidence(result)
    return result


def _reject_mixed_invoices(records: Sequence[Mapping[str, Any]]) -> None:
    identity_fields = ("invoice_number", "supplier_id", "supplier_name", "invoice_date")
    conflicts = []
    for field in identity_fields:
        values = _unique_nonempty(record.get(field) for record in records)
        if len(values) > 1:
            conflicts.append(field)
    if conflicts:
        raise DocumentExtractionError(
            "The spreadsheet appears to contain multiple invoices (conflicting "
            + ", ".join(conflicts)
            + "); split it into one invoice per file."
        )


def _parse_invoice_text(
    text: str,
    *,
    table_pages: Sequence[Mapping[str, Any]] | None = None,
) -> dict[str, Any]:
    result: dict[str, Any] = dict(_EMPTY_HEADER)
    warnings: list[str] = []
    confidence: dict[str, float] = {}
    provenance: dict[str, str] = {}

    document_type, type_confidence, type_evidence, type_warning = _classify_document_type(text)
    result["document_type"] = document_type
    result["document_type_confidence"] = type_confidence
    result["document_type_evidence"] = type_evidence
    if type_warning:
        warnings.append(type_warning)

    label_patterns = {
        "supplier_name": [r"(?im)^\s*(?:supplier|vendor)(?:\s+name)?\s*[:#-]\s*([^\r\n|]+)"],
        "supplier_id": [r"(?im)^\s*(?:supplier|vendor)\s+(?:id|code)\s*[:#-]\s*([^\r\n|]+)"],
        "supplier_site": [r"(?im)^\s*(?:supplier|vendor)\s+site\s*[:#-]\s*([^\r\n|]+)"],
        "invoice_number": [
            r"(?im)^\s*invoice\s*(?:number|no\.?|#)\s*[:#-]?\s*([A-Z0-9][A-Z0-9._/-]*)\s*$",
            r"(?im)^\s*invoice\s*[:#-]\s*([A-Z0-9][A-Z0-9._/-]*)\s*$",
            r"(?i)\binv(?:oice)?\.?\s*(?:number|no\.?|#)\s*[:#;-]?\s*([A-Z0-9][A-Z0-9._/-]*)\b",
        ],
        "po_number": [
            r"(?im)^\s*(?:purchase\s+order|p\.?\s*o\.?)\s*(?:number|no\.?|#)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9._/-]*)\s*$",
            r"(?i)\b(?:purchase\s+order|p\.?\s*o\.?)(?!\s*box\b)\s*(?:number|no\.?|#)?\s*[:#;-]?\s*([A-Z0-9][A-Z0-9._/-]*)\b",
        ],
    }
    for field, patterns in label_patterns.items():
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                result[field] = match.group(1).strip()
                confidence[field] = 0.94
                provenance[field] = "source_label"
                break

    if result["supplier_name"] is None:
        seller_name = _extract_seller_before_bill_to(text)
        if seller_name:
            result["supplier_name"] = seller_name
            confidence["supplier_name"] = 0.86
            provenance["supplier_name"] = "legal_name_immediately_before_bill_to"
    if result["supplier_name"] is None:
        seller_name = _extract_seller_above_invoice_title(text)
        if seller_name:
            result["supplier_name"] = seller_name
            confidence["supplier_name"] = 0.82
            provenance["supplier_name"] = "company_name_immediately_above_invoice_title"

    date_match = re.search(
        r"(?im)^\s*(?:invoice\s+date|date\s+of\s+invoice|date)\s*[:#-]\s*([^\r\n|]+)", text
    )
    if date_match:
        invoice_date, date_warning = _parse_date(date_match.group(1).strip())
        result["invoice_date"] = invoice_date
        if invoice_date:
            confidence["invoice_date"] = 0.92
            provenance["invoice_date"] = "source_label"
        if date_warning:
            warnings.append(date_warning)

    currency, currency_warning = _extract_currency(text)
    result["currency"] = currency
    if currency:
        confidence["currency"] = 0.85
        provenance["currency"] = "source_currency_evidence"
    if currency_warning:
        warnings.append(currency_warning)

    grid_fields, grid_warnings = _extract_layout_header_fields(table_pages)
    warnings.extend(grid_warnings)
    for field, value in grid_fields.items():
        if result.get(field) is None and value is not None:
            result[field] = value
            confidence[field] = 0.9
            provenance[field] = "fixed_position_header_grid"

    amounts, amount_warnings = _extract_totals_with_layout(text, table_pages)
    warnings.extend(amount_warnings)
    for field, value in amounts.items():
        result[field] = value
        if value is not None:
            confidence[field] = 0.9
            provenance[field] = "source_total_label"

    lines, table_method, table_warnings, table_diagnostics = _extract_lines(
        text, table_pages=table_pages
    )
    warnings.extend(table_warnings)
    schemas = table_diagnostics.get("header_schemas") or []
    complete_table_diagnostics = bool(
        int(table_diagnostics.get("recognized_table_sections") or 0) > 0
        and int(table_diagnostics.get("sections_without_rows") or 0) == 0
        and int(table_diagnostics.get("possible_unparsed_rows") or 0) == 0
        and int(table_diagnostics.get("unpaired_adjustment_rows") or 0) == 0
        and not bool(table_diagnostics.get("legacy_fallback"))
    )
    if (
        result["tax_total"] is None
        and result["total"] is not None
        and lines
        and complete_table_diagnostics
        and any("tax_amount:tax" in str(schema) for schema in schemas)
        and all(line.get("tax_amount") is not None for line in lines)
    ):
        try:
            derived_tax_total = sum(
                (Decimal(str(line["tax_amount"])) for line in lines),
                start=Decimal("0"),
            ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except InvalidOperation:
            derived_tax_total = None
        if derived_tax_total is not None:
            result["tax_total"] = float(derived_tax_total)
            confidence["tax_total"] = 0.82
            provenance["tax_total"] = "derived_sum_complete_explicit_line_tax_amounts"
            warnings.append(
                "Tax total was derived from all explicit line tax amounts in a complete recognized table using HALF_UP cents."
            )

    if (
        result["tax_total"] is None
        and result["subtotal"] is None
        and result["total"] is not None
        and _explicit_absent_tax_evidence(
            text,
            lines=lines,
            schemas=schemas,
            complete_table_diagnostics=complete_table_diagnostics,
            printed_total=result["total"],
        )
    ):
        result["tax_total"] = 0.0
        confidence["tax_total"] = 0.78
        provenance["tax_total"] = "absent_tax_treated_as_zero"
        warnings.append(
            "No VAT/tax field was printed in a complete Gross/Net table whose explicit net extensions reconcile to the printed total; tax total was treated as zero for review."
        )

    if (
        result["subtotal"] is None
        and result["total"] is not None
        and result["tax_total"] is not None
    ):
        try:
            derived_subtotal = (
                Decimal(str(result["total"])) - Decimal(str(result["tax_total"]))
            ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except InvalidOperation:
            derived_subtotal = None
        if derived_subtotal is not None:
            result["subtotal"] = float(derived_subtotal)
            confidence["subtotal"] = 0.84
            provenance["subtotal"] = "derived_grand_total_minus_tax_total"
            warnings.append(
                "Subtotal was derived from explicit grand total minus explicit or source-derived tax total using HALF_UP cents."
            )
    result["field_confidence"] = confidence
    result["field_provenance"] = provenance
    result["lines"] = lines
    result["table_extraction_method"] = table_method
    result["table_diagnostics"] = table_diagnostics
    result["warnings"] = warnings
    return result


def _explicit_absent_tax_evidence(
    text: str,
    *,
    lines: Sequence[Mapping[str, Any]],
    schemas: Sequence[Any],
    complete_table_diagnostics: bool,
    printed_total: Any,
) -> bool:
    """Require independent source evidence before treating omitted tax as zero."""

    if not complete_table_diagnostics or not lines or re.search(r"\b(?:vat|tax)\b", text, re.I):
        return False
    normalized_schemas = [str(schema) for schema in schemas]
    if not any(
        "line_total:net" in schema
        and "gross_line_total:gross" in schema
        and ("discount_amount:discount" in schema or "discount_rate:discount" in schema)
        for schema in normalized_schemas
    ):
        return False
    if not all(
        line.get("line_total") is not None
        and line.get("gross_line_total") is not None
        and line.get("tax_amount") is None
        and line.get("tax_rate") is None
        for line in lines
    ):
        return False
    try:
        extension_total = sum(
            (Decimal(str(line["line_total"])) for line in lines),
            start=Decimal("0"),
        ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        source_total = Decimal(str(printed_total)).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )
    except (InvalidOperation, TypeError, ValueError):
        return False
    return extension_total == source_total


def _classify_document_type(
    text: str,
) -> tuple[str, float, list[str], str | None]:
    """Classify document workflow without using filenames or supplier rules.

    Evidence contains rule names only; it is safe to retain without copying
    private document text.  Weak or conflicting evidence remains ``unknown``.
    """

    normalized_lines = [_normalized_words(line) for line in text.splitlines()[:80]]
    scores = {kind: 0 for kind in _DOCUMENT_TYPES if kind != "unknown"}
    evidence: dict[str, list[str]] = {kind: [] for kind in scores}

    title_rules = (
        ("credit_note", re.compile(r"^(?:tax\s+)?credit\s+(?:note|memo)(?:\s+(?:original|copy))?$"), "title:credit_note", 8),
        ("delivery_note", re.compile(r"^(?:proof\s+of\s+delivery|delivery\s+(?:note|challan))(?:\s+(?:original|copy))?$"), "title:delivery_note", 7),
        ("purchase_order", re.compile(r"^(?:purchase\s+order|p\s*o)(?:\s+(?:original|copy))?$"), "title:purchase_order", 7),
        ("invoice", re.compile(r"^(?:(?:tax|commercial|sales|credit)\s+)?invoice(?:\s+(?:original|copy))?$"), "title:invoice", 7),
    )
    for line in normalized_lines:
        for kind, pattern, rule, weight in title_rules:
            if pattern.search(line):
                scores[kind] += weight
                if rule not in evidence[kind]:
                    evidence[kind].append(rule)

    phrase_rules = (
        ("credit_note", re.compile(r"\b(?:tax\s+)?credit\s+(?:note|memo)\b"), "phrase:credit_note", 8),
        ("delivery_note", re.compile(r"\b(?:proof\s+of\s+delivery|delivery\s+(?:note|challan))\b"), "phrase:delivery_note", 6),
        ("invoice", re.compile(r"\b(?:tax|commercial|sales|credit)\s+invoice\b"), "phrase:qualified_invoice", 7),
    )
    for line in normalized_lines:
        for kind, pattern, rule, weight in phrase_rules:
            if pattern.search(line):
                scores[kind] += weight
                if rule not in evidence[kind]:
                    evidence[kind].append(rule)

    field_rules = (
        ("invoice", re.compile(r"\binvoice\s+(?:number|no)\b"), "field:invoice_number", 2),
        ("invoice", re.compile(r"\b(?:invoice\s+date|date\s+of\s+invoice)\b"), "field:invoice_date", 1),
        ("credit_note", re.compile(r"\bcredit\s+(?:note|memo)\s+(?:number|no)\b"), "field:credit_note_number", 3),
        ("delivery_note", re.compile(r"\bdelivery\s+(?:note\s+)?(?:number|no)\b"), "field:delivery_note_number", 3),
        ("purchase_order", re.compile(r"\bpurchase\s+order\s+(?:number|no)\b"), "field:purchase_order_number", 3),
    )
    for line in normalized_lines:
        for kind, pattern, rule, weight in field_rules:
            if pattern.search(line):
                scores[kind] += weight
                if rule not in evidence[kind]:
                    evidence[kind].append(rule)

    ranked = sorted(scores.items(), key=lambda item: item[1], reverse=True)
    best_kind, best_score = ranked[0]
    second_score = ranked[1][1]
    if best_score < 3 or (second_score >= 3 and best_score - second_score < 2):
        present = sorted(kind for kind, score in scores.items() if score >= 3)
        rule_evidence = sorted({rule for kind in present for rule in evidence[kind]})
        warning = (
            "Document type evidence conflicts; route the document to review."
            if len(present) > 1
            else "Document type could not be established from source labels; route the document to review."
        )
        return "unknown", 0.0, rule_evidence, warning

    confidence = 0.98 if best_score >= 8 else 0.9 if best_score >= 7 else 0.75
    warning = None
    if best_kind != "invoice":
        warning = (
            f"Document classified as {best_kind.replace('_', ' ')}; keep it outside the invoice workflow pending review."
        )
    return best_kind, confidence, evidence[best_kind], warning


def _extract_seller_before_bill_to(text: str) -> str | None:
    """Capture a legal seller name only when a Bill To boundary is explicit."""

    lines = [line.strip() for line in text.splitlines()[:50]]
    legal_suffix = re.compile(
        r"\b(?:L\.?L\.?C\.?|LTD\.?|LIMITED|INC\.?|CORP(?:ORATION)?\.?|PLC|FZE|FZC|WLL|GMBH|S\.?A\.?)\b",
        re.I,
    )
    for index, line in enumerate(lines):
        ascii_line = re.sub(r"[^\x00-\x7F]+", " ", line)
        if not re.search(r"\b(?:bill|invoice|sold)\s+to\b", ascii_line, re.I):
            continue
        for candidate in reversed(lines[max(0, index - 4) : index]):
            ascii_candidate = re.sub(r"[^\x00-\x7F]+", " ", candidate)
            ascii_candidate = re.sub(r"\s+", " ", ascii_candidate).strip(" |:-")
            ascii_candidate = re.sub(r"(?:\s+\.)+\s*$", "", ascii_candidate)
            if legal_suffix.search(ascii_candidate) and len(ascii_candidate) <= 160:
                return ascii_candidate
        return None
    return None


def _extract_seller_above_invoice_title(text: str) -> str | None:
    """Capture a company line only when it directly captions an invoice title.

    This is deliberately narrower than a first-line guess: the candidate must
    immediately precede a qualified invoice heading and contain an explicit
    legal or company marker.  Buyer/client labels are rejected.
    """

    lines = [re.sub(r"\s+", " ", line).strip(" |:-") for line in text.splitlines()[:60]]
    invoice_title = re.compile(r"^(?:(?:tax|commercial|sales|credit)\s+)?invoice$", re.I)
    company_marker = re.compile(
        r"\b(?:company|co\.?|l\.?l\.?c\.?|ltd\.?|limited|inc\.?|corp(?:oration)?\.?|plc|fze|fzc|wll|gmbh|s\.?a\.?)\b",
        re.I,
    )
    for index, line in enumerate(lines):
        if not invoice_title.fullmatch(_normalized_words(line)):
            continue
        for candidate in reversed(lines[max(0, index - 3) : index]):
            ascii_candidate = re.sub(r"[^\x00-\x7F]+", " ", candidate)
            ascii_candidate = re.sub(r"\s+", " ", ascii_candidate).strip(" |:-")
            if not ascii_candidate:
                continue
            if re.search(r"\b(?:client|customer|buyer|bill|deliver|ship)\b", ascii_candidate, re.I):
                continue
            if company_marker.search(ascii_candidate) and 3 <= len(ascii_candidate) <= 160:
                return ascii_candidate
        return None
    return None


def _extract_currency(text: str) -> tuple[str | None, str | None]:
    label = re.search(r"(?im)^\s*currency\s*[:#-]\s*([A-Z]{3})\b", text)
    if label:
        code = label.group(1).upper()
        if code in _CURRENCIES:
            return code, None
        supported_codes = sorted(
            {match.upper() for match in re.findall(r"\b(?:AED|USD|EUR|GBP|SAR|QAR|KWD|BHD|OMR)\b", text, re.I)}
        )
        if len(supported_codes) == 1:
            return (
                supported_codes[0],
                f"The currency label was OCR-read as unsupported code '{code}'; the one explicit supported currency code elsewhere in the source was retained for review.",
            )
        return None, f"Currency code '{code}' is not in the supported currency set and was withheld."
    codes = sorted({match.upper() for match in re.findall(r"\b(?:AED|USD|EUR|GBP|SAR|QAR|KWD|BHD|OMR)\b", text, re.I)})
    if len(codes) == 1:
        return codes[0], None
    if len(codes) > 1:
        return None, "Multiple currency codes were found; currency was withheld."
    if "€" in text:
        return "EUR", None
    if "£" in text:
        return "GBP", None
    if "$" in text:
        return None, "The '$' symbol is ambiguous; currency was withheld."
    return None, None


def _extract_layout_header_fields(
    table_pages: Sequence[Mapping[str, Any]] | None,
) -> tuple[dict[str, str], list[str]]:
    if not table_pages:
        return {}, []
    patterns = {
        "invoice_number": re.compile(r"\binvoice\s+(?:number|no\.?)\b", re.I),
        "invoice_date": re.compile(r"\binvoice\s+date\b", re.I),
        "currency": re.compile(r"\binvoice\s+currency\b|\bcurrency\b", re.I),
    }
    candidates: dict[str, set[str]] = {field: set() for field in patterns}
    for page in table_pages:
        lines = str(page.get("text") or "").splitlines()[:60]
        for index, header in enumerate(lines):
            anchors: list[tuple[str, int, int]] = []
            for field, pattern in patterns.items():
                match = pattern.search(header)
                if match:
                    anchors.append((field, match.start(), match.end()))
            if len(anchors) < 2:
                continue
            anchors.sort(key=lambda item: item[1])
            centers = [(start + end) / 2 for _, start, end in anchors]
            boundaries = [0]
            boundaries.extend(
                int(round((left + right) / 2))
                for left, right in zip(centers, centers[1:])
            )
            boundaries.append(max(len(header), int(centers[-1] + 80)))
            for value_line in lines[index + 1 : index + 5]:
                row_values: dict[str, str] = {}
                for anchor_index, (field, _, _) in enumerate(anchors):
                    cell = value_line[boundaries[anchor_index] : boundaries[anchor_index + 1]].strip()
                    value = _parse_header_grid_cell(field, cell)
                    if value is not None:
                        row_values[field] = value
                if len(row_values) >= 2:
                    for field, value in row_values.items():
                        candidates[field].add(value)

    result: dict[str, str] = {}
    warnings: list[str] = []
    for field, values in candidates.items():
        if len(values) == 1:
            result[field] = next(iter(values))
        elif len(values) > 1:
            warnings.append(
                f"Fixed-position header rows disagree on {field}; the value was withheld."
            )
    return result, warnings


def _parse_header_grid_cell(field: str, cell: str) -> str | None:
    if not cell:
        return None
    if field == "currency":
        matches = sorted(set(re.findall(r"\b(?:AED|USD|EUR|GBP|SAR|QAR|KWD|BHD|OMR)\b", cell, re.I)))
        return matches[0].upper() if len(matches) == 1 else None
    if field == "invoice_date":
        date_patterns = (
            r"\b\d{4}-\d{1,2}-\d{1,2}\b",
            r"\b\d{1,2}[- ](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[A-Za-z]*[- ]\d{4}\b",
            r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{4}\b",
        )
        for pattern in date_patterns:
            match = re.search(pattern, cell, re.I)
            if match:
                parsed, _ = _parse_date(match.group())
                if parsed:
                    return parsed
        return None
    match = re.search(r"\b[A-Z0-9][A-Z0-9._/-]{3,}\b", cell, re.I)
    return match.group() if match else None


def _extract_totals_with_layout(
    text: str,
    table_pages: Sequence[Mapping[str, Any]] | None,
) -> tuple[dict[str, float | None], list[str]]:
    sources = [text]
    if table_pages:
        layout_text = "\n".join(str(page.get("text") or "") for page in table_pages)
        if layout_text.strip() and layout_text != text:
            sources.append(layout_text)
    parsed_sources = [_extract_totals(source) for source in sources]
    result: dict[str, float | None] = {"subtotal": None, "tax_total": None, "total": None}
    warnings: list[str] = []
    for field in result:
        values = {
            round(float(parsed[field]), 6)
            for parsed in parsed_sources
            if parsed.get(field) is not None
        }
        if len(values) == 1:
            result[field] = values.pop()
        elif len(values) > 1:
            warnings.append(
                f"Reading-order and fixed-position text disagree on {field}; the value was withheld."
            )
    return result, warnings


def _extract_totals(text: str) -> dict[str, float | None]:
    result: dict[str, float | None] = {"subtotal": None, "tax_total": None, "total": None}
    patterns = {
        "subtotal": re.compile(
            r"^\s*[^A-Za-z0-9]{0,3}(?:sub[ -]?total|untaxed\s+amount|total(?:\s+amount)?\s+(?:before|excluding|excl\.?)\s+(?:vat|tax)|net\s+(?:amount|total)\s+(?:before|excluding|excl\.?)\s+(?:vat|tax)|taxable\s+amount)\b",
            re.I,
        ),
        "tax_total": re.compile(
            r"^\s*[^A-Za-z0-9]{0,3}(?:(?:total\s+)?(?:vat|tax)(?:\s+(?:total|amount|value))?|(?:vat|tax)\s+total)\b",
            re.I,
        ),
        "total": re.compile(
            r"^\s*[^A-Za-z0-9]{0,3}(?:grand\s+total|invoice\s+total|amount\s+due|total\s+(?:amount\s+)?due(?:\s*\([^)]*\))?|net\s+payable|total\s+payable|total(?:\s+amount)?\s+(?:including|inclusive\s+of|incl\.?)\s+(?:vat|tax)|total\s+amount|total|otal)\b",
            re.I,
        ),
    }
    lines = text.splitlines()
    for field, pattern in patterns.items():
        for line_index in range(len(lines) - 1, -1, -1):
            line = lines[line_index]
            match = pattern.search(line)
            if not match:
                continue
            # Never treat paid/settled/balance-forward values as an invoice total.
            if re.search(r"\b(?:paid|settled|payment|balance\s+forward)\b", line, re.I):
                continue
            if field == "total" and re.search(
                r"\b(?:sub[ -]?total|vat|tax|qty|quantity|pieces?|cartons?)\b",
                line,
                re.I,
            ) and not re.search(
                r"\b(?:including|inclusive\s+of|incl\.?)\s+(?:vat|tax)\b",
                line,
                re.I,
            ):
                continue
            tail = line[match.end() :]
            value = _last_number(tail)
            if field == "tax_total" and "%" in tail:
                numeric_values = re.findall(r"-?\d[\d,.']*", tail)
                if len(numeric_values) < 2:
                    value = None
            if value is None:
                value = _following_standalone_amount(lines, line_index)
            if value is not None:
                result[field] = value
                break
    return result


def _following_standalone_amount(lines: Sequence[str], index: int) -> float | None:
    for following in lines[index + 1 : index + 4]:
        stripped = following.strip()
        if not stripped:
            continue
        value = _parse_cell_number(stripped)
        if value is not None:
            return value
        # A translated label may sit between the English label and amount.
        # An ASCII field label starts a new field and ends the search.
        if re.search(r"[A-Za-z]", stripped):
            break
    return None


def _extract_lines(
    text: str,
    *,
    table_pages: Sequence[Mapping[str, Any]] | None = None,
) -> tuple[list[dict[str, Any]], str, list[str], dict[str, Any]]:
    pages = list(table_pages or [{"page": 1, "text": text, "method": "text"}])
    extracted: list[dict[str, Any]] = []
    header_sections = 0
    empty_sections = 0
    possible_unparsed_rows = 0
    continuation_lines = 0
    paired_adjustment_rows = 0
    unpaired_adjustment_rows = 0
    source_data_rows = 0
    header_schemas: list[str] = []
    for page in pages:
        (
            page_lines,
            page_headers,
            page_empty,
            page_possible,
            page_continuations,
            page_schemas,
            page_paired_adjustments,
            page_unpaired_adjustments,
            page_source_rows,
        ) = _extract_layout_page_lines(
            str(page.get("text") or ""),
            page_number=int(page.get("page") or 1),
        )
        extracted.extend(page_lines)
        header_sections += page_headers
        empty_sections += page_empty
        possible_unparsed_rows += page_possible
        continuation_lines += page_continuations
        header_schemas.extend(page_schemas)
        paired_adjustment_rows += page_paired_adjustments
        unpaired_adjustment_rows += page_unpaired_adjustments
        source_data_rows += page_source_rows

    warnings: list[str] = []
    if extracted:
        for index, line in enumerate(extracted):
            line["id"] = _line_id(index, line)
        if empty_sections:
            warnings.append(
                f"{empty_sections} table section(s) had recognizable headers but no source rows meeting the conservative numeric checks."
            )
        if unpaired_adjustment_rows:
            warnings.append(
                f"{unpaired_adjustment_rows} explicit discount adjustment row(s) could not be paired one-to-one and require review."
            )
        return extracted, "layout_header", warnings, {
            "recognized_table_sections": header_sections,
            "sections_without_rows": empty_sections,
            "possible_unparsed_rows": possible_unparsed_rows,
            "emitted_rows": len(extracted),
            "joined_continuation_lines": continuation_lines,
            "legacy_fallback": False,
            "header_schemas": header_schemas,
            "paired_adjustment_rows": paired_adjustment_rows,
            "unpaired_adjustment_rows": unpaired_adjustment_rows,
            "source_data_rows": source_data_rows,
        }

    legacy_allowed = header_sections > 0 or all(
        str(page.get("method") or "")
        in {"text", "docx_text", "csv_structured", "xlsx_structured"}
        for page in pages
    )
    legacy = _extract_legacy_lines(text) if legacy_allowed else []
    if legacy:
        return legacy, "legacy_pattern", warnings, {
            "recognized_table_sections": header_sections,
            "sections_without_rows": empty_sections,
            "possible_unparsed_rows": possible_unparsed_rows,
            "emitted_rows": len(legacy),
            "joined_continuation_lines": continuation_lines,
            "legacy_fallback": True,
            "header_schemas": header_schemas,
            "paired_adjustment_rows": paired_adjustment_rows,
            "unpaired_adjustment_rows": unpaired_adjustment_rows,
            "source_data_rows": source_data_rows,
        }
    if header_sections:
        warnings.append(
            "Table headings were found, but no row had an explicit description, quantity, and source price or total."
        )
    return [], "none", warnings, {
        "recognized_table_sections": header_sections,
        "sections_without_rows": empty_sections,
        "possible_unparsed_rows": possible_unparsed_rows,
        "emitted_rows": 0,
        "joined_continuation_lines": continuation_lines,
        "legacy_fallback": False,
        "header_schemas": header_schemas,
        "paired_adjustment_rows": paired_adjustment_rows,
        "unpaired_adjustment_rows": unpaired_adjustment_rows,
        "source_data_rows": source_data_rows,
    }


def _extract_layout_page_lines(
    text: str,
    *,
    page_number: int,
) -> tuple[list[dict[str, Any]], int, int, int, int, list[str], int, int, int]:
    source_lines = [line.rstrip() for line in text.splitlines()]
    headers = _find_table_headers(source_lines)
    extracted: list[dict[str, Any]] = []
    empty_sections = 0
    possible_unparsed_rows = 0
    continuation_lines = 0
    header_schemas: list[str] = []
    paired_adjustment_rows = 0
    unpaired_adjustment_rows = 0
    source_data_rows = 0
    for header_index, (start, end, anchors) in enumerate(headers):
        header_schemas.append(
            "|".join(
                f"{anchor.field}:{anchor.basis}" if anchor.basis else anchor.field
                for anchor in anchors
            )
        )
        section_end = headers[header_index + 1][0] if header_index + 1 < len(headers) else len(source_lines)
        (
            section,
            section_possible,
            section_continuations,
            section_paired_adjustments,
            section_unpaired_adjustments,
            section_source_rows,
        ) = _parse_layout_section(
            source_lines,
            start=end + 1,
            end=section_end,
            anchors=anchors,
            page_number=page_number,
        )
        if section:
            extracted.extend(section)
        else:
            empty_sections += 1
        possible_unparsed_rows += section_possible
        continuation_lines += section_continuations
        paired_adjustment_rows += section_paired_adjustments
        unpaired_adjustment_rows += section_unpaired_adjustments
        source_data_rows += section_source_rows
    return (
        extracted,
        len(headers),
        empty_sections,
        possible_unparsed_rows,
        continuation_lines,
        header_schemas,
        paired_adjustment_rows,
        unpaired_adjustment_rows,
        source_data_rows,
    )


def _find_table_headers(
    lines: Sequence[str],
) -> list[tuple[int, int, list[_ColumnAnchor]]]:
    headers: list[tuple[int, int, list[_ColumnAnchor]]] = []
    index = 0
    while index < len(lines):
        best: tuple[int, int, list[_ColumnAnchor], int] | None = None
        for height in (1, 2, 3):
            end = index + height
            if end > len(lines):
                break
            anchors = _collect_column_anchors(lines[index:end])
            fields = {anchor.field for anchor in anchors}
            if not {
                "description",
                "quantity",
            }.issubset(fields) or not ({"unit_price", "line_total"} & fields):
                continue
            score = len(fields) * 10 + sum(anchor.priority for anchor in anchors)
            if best is None or score > best[3]:
                best = (index, end - 1, anchors, score)
        if best is None:
            index += 1
            continue
        headers.append((best[0], best[1], best[2]))
        index = best[1] + 1
    return headers


def _collect_column_anchors(lines: Sequence[str]) -> list[_ColumnAnchor]:
    candidates: list[_ColumnAnchor] = []
    for line in lines:
        # Unicode descriptions can be printed directly beside an English
        # heading.  Replacing only non-ASCII characters keeps the exact x
        # offsets while restoring the intended English word boundary.
        searchable = "".join(character if ord(character) < 128 else " " for character in line)
        for field, pattern, priority, basis in _COLUMN_PATTERNS:
            for match in pattern.finditer(searchable):
                candidates.append(
                    _ColumnAnchor(
                        field=field,
                        start=match.start(),
                        end=match.end(),
                        priority=priority,
                        basis=basis,
                    )
                )

    selected: list[_ColumnAnchor] = []
    for field in {candidate.field for candidate in candidates}:
        field_candidates = [candidate for candidate in candidates if candidate.field == field]
        field_candidates.sort(
            key=lambda candidate: (
                candidate.priority,
                candidate.start if field == "line_total" else -candidate.start,
                candidate.end - candidate.start,
            ),
            reverse=True,
        )
        selected.append(field_candidates[0])

    # Remove a lower-specificity word contained inside a stronger phrase such
    # as UNIT inside UNIT PRICE or AMOUNT inside TAX AMOUNT.
    resolved: list[_ColumnAnchor] = []
    for candidate in sorted(selected, key=lambda anchor: anchor.priority, reverse=True):
        overlaps = any(
            max(candidate.start, other.start) < min(candidate.end, other.end)
            for other in resolved
        )
        if not overlaps:
            resolved.append(candidate)
    resolved.sort(key=lambda anchor: (anchor.start, anchor.end))

    # Two headings printed at effectively the same horizontal position cannot
    # define independent cells, so keep the more specific one.
    distinct: list[_ColumnAnchor] = []
    for candidate in resolved:
        if distinct and candidate.start - distinct[-1].start < 2:
            if candidate.priority > distinct[-1].priority:
                distinct[-1] = candidate
            continue
        distinct.append(candidate)
    return distinct


def _parse_layout_section(
    lines: Sequence[str],
    *,
    start: int,
    end: int,
    anchors: Sequence[_ColumnAnchor],
    page_number: int,
) -> tuple[list[dict[str, Any]], int, int, int, int, int]:
    fields = {anchor.field: anchor for anchor in anchors}
    if "description" not in fields or "quantity" not in fields:
        return [], 0, 0, 0, 0, 0

    extracted: list[dict[str, Any]] = []
    pending_description: list[str] = []
    last_data_row: int | None = None
    possible_unparsed_rows = 0
    continuation_lines = 0
    annotation_block = False
    for row_index in range(start, end):
        raw_line = lines[row_index]
        stripped = raw_line.strip()
        if not stripped:
            continue
        if _is_table_summary(stripped):
            if extracted:
                break
            pending_description.clear()
            continue
        normalized_row = _normalized_words(stripped)
        if (
            extracted
            and not re.search(r"\d", stripped)
            and re.search(r"\b(?:logistics|quality|warehouse)\s+division\b", normalized_row)
        ):
            annotation_block = True
            continue
        if annotation_block:
            continue

        if (
            not extracted
            and row_index - start < 3
            and not re.search(r"[A-Za-z]", stripped)
            and sum(character.isalpha() for character in stripped) >= 3
        ):
            # A translated heading is often printed directly under the
            # English table heading.  Do not attach that heading to the first
            # product description.
            continue

        cells = _slice_layout_cells(raw_line, anchors)
        tail_cells = _slice_right_aligned_numeric_cells(
            raw_line,
            anchors,
            allow_uom=_layout_cells_misaligned(raw_line, anchors, cells),
        )
        if tail_cells is not None:
            cells.update(tail_cells)
        description = _clean_description(cells.get("description"))
        quantity, embedded_uom = _parse_quantity_cell(cells.get("quantity"))
        uom = _clean_uom(cells.get("uom")) or embedded_uom
        printed_unit_price = _parse_layout_number(cells.get("unit_price"))
        line_total = _parse_layout_number(cells.get("line_total"))
        gross_line_total = _parse_layout_number(cells.get("gross_line_total"))
        tax_rate = _parse_layout_number(cells.get("tax_rate"))
        discount_rate = _parse_layout_number(cells.get("discount_rate"))
        discount_amount = _parse_layout_number(cells.get("discount_amount"))
        tax_amount = _parse_layout_number(cells.get("tax_amount"))
        upc = _clean_upc(cells.get("upc"))
        item_code = _clean_item_code(cells.get("item_code"))

        corrected_quantity = _validated_ocr_quantity(
            cells.get("quantity"),
            cells.get("unit_price"),
            quantity,
            printed_unit_price,
            line_total,
            gross_line_total,
        )
        quantity_was_normalized = corrected_quantity is not None and corrected_quantity != quantity
        if corrected_quantity is not None:
            quantity = corrected_quantity

        reconciled_line_total = _line_extension_from_consistent_source_columns(
            cells.get("unit_price"),
            quantity,
            printed_unit_price,
            line_total,
            gross_line_total,
            discount_rate,
            discount_amount,
        )
        line_total_was_reconciled = (
            reconciled_line_total is not None and reconciled_line_total != line_total
        )
        if reconciled_line_total is not None:
            line_total = reconciled_line_total

        has_required_numbers = quantity is not None and (
            printed_unit_price is not None or line_total is not None
        )
        if has_required_numbers:
            pieces = [*pending_description]
            if description:
                pieces.append(description)
            pending_description.clear()
            full_description = _join_description_parts(pieces)
            if not full_description or _is_summary_label(full_description):
                possible_unparsed_rows += 1
                continue
            price_basis = fields.get("unit_price").basis if fields.get("unit_price") else None
            total_basis = fields.get("line_total").basis if fields.get("line_total") else None
            unit_price = printed_unit_price
            net_unit_price = printed_unit_price if price_basis == "net" else None
            gross_unit_price = printed_unit_price if price_basis == "gross" else None
            unit_price_basis = price_basis
            unit_price_source = (
                "printed_net_unit_price"
                if price_basis == "net"
                else "printed_gross_unit_price"
                if price_basis == "gross"
                else "printed_unit_price"
                if printed_unit_price is not None
                else None
            )
            derived_fields: list[str] = []
            if quantity_was_normalized:
                derived_fields.append("quantity")
            if line_total_was_reconciled:
                derived_fields.append("line_total")
            if total_basis == "net" and line_total is not None and quantity not in (None, 0):
                derived_net_unit_price = _divide_money_by_quantity(line_total, quantity)
                if derived_net_unit_price is not None:
                    unit_price = derived_net_unit_price
                    net_unit_price = derived_net_unit_price
                    unit_price_basis = "net"
                    unit_price_source = "derived_net_line_amount_divided_by_quantity"
                    derived_fields.extend(["unit_price", "net_unit_price"])
            line: dict[str, Any] = {
                "description": full_description,
                "quantity": quantity,
                "quantity_source": (
                    "ocr_numeric_cell_validated_by_explicit_rate_and_extension"
                    if quantity_was_normalized
                    else "printed_quantity"
                    if quantity is not None
                    else None
                ),
                "uom": uom,
                "unit_price": unit_price,
                "net_unit_price": net_unit_price,
                "gross_unit_price": gross_unit_price,
                "printed_unit_price": printed_unit_price,
                "printed_unit_price_basis": price_basis,
                "unit_price_basis": unit_price_basis,
                "unit_price_source": unit_price_source,
                "line_total": line_total,
                "line_total_basis": total_basis,
                "line_total_source": (
                    "gross_extension_validated_by_zero_discount_and_printed_rate"
                    if line_total_was_reconciled
                    else "printed_line_total"
                    if line_total is not None
                    else None
                ),
                "gross_line_total": gross_line_total,
                "tax_rate": tax_rate,
                "tax_amount": tax_amount,
                "discount_rate": discount_rate,
                "discount_amount": discount_amount,
                "upc": upc,
                "item_code": item_code,
                "source_page": page_number,
                "source_rows": [row_index + 1, row_index + 1],
                "source_sequence": _parse_cell_number(cells.get("row_number")),
                "extraction_method": "layout_header",
                "derived_fields": derived_fields,
                "confidence": 0.94 if unit_price is not None and line_total is not None else 0.87,
            }
            extracted.append(line)
            last_data_row = row_index
            continue

        # Some printed tables put an unlabeled GTIN directly below the item
        # description.  Treat it as a continuation only when it is adjacent
        # to an emitted row and its source digits have a valid GS1 check
        # digit.  Other numeric-only text remains reviewable as unparsed.
        continuation_upc = _clean_upc(cells.get("description"))
        if (
            continuation_upc
            and extracted
            and last_data_row is not None
            and row_index - last_data_row <= 3
        ):
            if _valid_gtin(continuation_upc) and extracted[-1].get("upc") is None:
                extracted[-1]["upc"] = continuation_upc
            elif not _valid_gtin(continuation_upc):
                extracted[-1].setdefault("ocr_identifier_candidates", []).append(
                    continuation_upc
                )
            extracted[-1]["source_rows"][1] = row_index + 1
            last_data_row = row_index
            continuation_lines += 1
            continue

        has_other_cell_evidence = any(
            value is not None
            for value in (
                quantity,
                printed_unit_price,
                line_total,
                gross_line_total,
                tax_rate,
                discount_rate,
                discount_amount,
                tax_amount,
                upc,
                item_code,
            )
        )
        numeric_continuation_values = [
            value
            for value in (
                quantity,
                printed_unit_price,
                line_total,
                gross_line_total,
                tax_rate,
                discount_rate,
                discount_amount,
                tax_amount,
            )
            if value is not None
        ]
        if (
            description
            and extracted
            and last_data_row is not None
            and row_index - last_data_row <= 3
            and quantity is None
            and printed_unit_price is None
            and gross_line_total is None
            and len(numeric_continuation_values) <= 1
        ):
            extracted[-1]["description"] = _join_description_parts(
                [extracted[-1]["description"], description]
            )
            if numeric_continuation_values:
                extracted[-1].setdefault("ocr_ignored_continuation_noise", []).extend(
                    numeric_continuation_values
                )
            extracted[-1]["source_rows"][1] = row_index + 1
            last_data_row = row_index
            continuation_lines += 1
            continue
        numeric_like = any(
            re.search(r"\d", cells.get(field, ""))
            for field in (
                "quantity",
                "unit_price",
                "line_total",
                "tax_rate",
                "tax_amount",
                "discount_rate",
                "discount_amount",
            )
        )
        if has_other_cell_evidence or numeric_like:
            possible_unparsed_rows += 1
            continue
        if not description:
            continue

        # Description-only physical lines immediately below a data row are a
        # wrapped continuation.  Before the first row they are held for the
        # next numeric row.  Values are never derived from either case.
        if extracted and last_data_row is not None and row_index - last_data_row <= 3:
            extracted[-1]["description"] = _join_description_parts(
                [extracted[-1]["description"], description]
            )
            extracted[-1]["source_rows"][1] = row_index + 1
            last_data_row = row_index
            continuation_lines += 1
        else:
            pending_description.append(description)
            pending_description = pending_description[-3:]

    source_data_rows = len(extracted)
    extracted, paired_adjustments, unpaired_adjustments = _pair_explicit_discount_rows(extracted)
    return (
        extracted,
        possible_unparsed_rows,
        continuation_lines,
        paired_adjustments,
        unpaired_adjustments,
        source_data_rows,
    )


# A pack or size token printed inside a description ("250ml", "1kg", "33cl",
# "6x330ml", "6 x 330") describes the product, never the quantity ordered.  A
# bare number followed by a unit word ("2 EA", "8 KG") is a quantity.  Size
# tokens are masked before numeric cells are located so they stay in the
# description text.
_SIZE_UNIT = r"(?:ml|cl|dl|l|ltr|litre|liter|g|gm|gr|kg|mg|oz|lb|lbs)"
_SIZE_TOKEN_PATTERN = re.compile(
    rf"(?<![A-Za-z0-9])(?:\d+\s*[xX\u00d7]\s*)?\d[\d.,]*{_SIZE_UNIT}(?![A-Za-z0-9])"
    rf"|(?<![A-Za-z0-9])\d+\s*[xX\u00d7]\s*\d[\d.,]*(?![A-Za-z0-9])",
    re.I,
)


def _mask_size_tokens(line: str) -> str:
    """Blank pack/size tokens with spaces so character offsets are preserved."""

    return _SIZE_TOKEN_PATTERN.sub(lambda match: " " * len(match.group()), line)


_LAYOUT_NUMERIC_FIELDS = {
    "quantity",
    "unit_price",
    "line_total",
    "gross_line_total",
    "tax_rate",
    "tax_amount",
    "discount_rate",
    "discount_amount",
}


def _slice_right_aligned_numeric_cells(
    line: str,
    anchors: Sequence[_ColumnAnchor],
    *,
    allow_uom: bool = False,
) -> dict[str, str] | None:
    """Recover tables whose centered headings do not align with overflow rows.

    The fallback is deliberately limited to a description plus numeric tail.
    Tables with UPC or item-code columns continue to use fixed positions.  A
    UOM column is accepted only when the caller passes ``allow_uom`` (fixed
    positions cut through a token on this row); the unit word must then sit
    alone in the UOM column's position between the numeric cells.
    """

    nonnumeric_fields = {
        anchor.field for anchor in anchors if anchor.field not in _LAYOUT_NUMERIC_FIELDS
    }
    allowed_nonnumeric = {"row_number", "description"}
    if allow_uom:
        allowed_nonnumeric.add("uom")
    if not nonnumeric_fields.issubset(allowed_nonnumeric):
        return None
    numeric_anchors = [
        anchor for anchor in anchors if anchor.field in _LAYOUT_NUMERIC_FIELDS
    ]
    if len(numeric_anchors) < 3:
        return None
    matches = list(
        re.finditer(
            r"(?<![A-Za-z0-9])(?:N\s*/\s*A|N\.?\s*A\.?|\(?-?\d[\d,.']*\)?%?)(?![A-Za-z0-9])",
            _mask_size_tokens(line),
            re.I,
        )
    )
    if len(matches) < len(numeric_anchors):
        return None
    selected = matches[-len(numeric_anchors) :]
    uom_text: str | None = None
    if "uom" in nonnumeric_fields:
        ordered = [
            anchor for anchor in anchors if anchor.field not in {"row_number", "description"}
        ]
        uom_position = next(index for index, anchor in enumerate(ordered) if anchor.field == "uom")
        if uom_position == 0:
            # UOM printed before the quantity: only a known unit word at the
            # end of the description prefix is taken; anything else stays in
            # the description.
            prefix_match = re.search(rf"\s({_UOM})\s*$", line[: selected[0].start()], re.I)
            uom_text = prefix_match.group(1) if prefix_match else ""
        else:
            span_start = selected[uom_position - 1].end()
            span_end = (
                selected[uom_position].start() if uom_position < len(selected) else len(line)
            )
            between = line[span_start:span_end].strip()
            if between and not re.fullmatch(r"[A-Za-z]{1,8}\.?", between):
                return None
            uom_text = between
    for anchor, match in zip(numeric_anchors, selected):
        if _parse_cell_number(match.group()) is None and anchor.field not in {
            "tax_rate",
            "tax_amount",
            "discount_rate",
            "discount_amount",
        }:
            return None

    cells = {
        anchor.field: match.group()
        for anchor, match in zip(numeric_anchors, selected)
    }
    prefix = line[: selected[0].start()].strip()
    if uom_text is not None:
        cells["uom"] = uom_text
        if uom_text and prefix.upper().endswith(uom_text.upper()):
            prefix = prefix[: -len(uom_text)].strip()
    if "row_number" in nonnumeric_fields:
        sequence = re.match(r"^\s*(\d{1,6})[.)]?\s+", prefix)
        if sequence:
            cells["row_number"] = sequence.group(1)
            prefix = prefix[sequence.end() :].strip()
    cells["description"] = prefix
    return cells


def _pair_explicit_discount_rows(
    lines: Sequence[Mapping[str, Any]],
) -> tuple[list[dict[str, Any]], int, int]:
    """Pair a directly adjacent explicit non-positive discount with one product.

    Pairing requires equal source quantity, compatible tax rate and explicit
    net/ex-tax line amounts on both rows.  The consumed adjustment is retained
    on the emitted product line so no source row silently disappears.
    """

    paired: list[dict[str, Any]] = []
    paired_adjustments = 0
    unpaired_adjustments = 0
    index = 0
    while index < len(lines):
        current = dict(lines[index])
        following = dict(lines[index + 1]) if index + 1 < len(lines) else None
        product: dict[str, Any] | None = None
        adjustment: dict[str, Any] | None = None
        if following is not None:
            if _is_explicit_discount_adjustment(current) and not _is_explicit_discount_adjustment(following):
                adjustment, product = current, following
            elif not _is_explicit_discount_adjustment(current) and _is_explicit_discount_adjustment(following):
                product, adjustment = current, following
        if product is not None and adjustment is not None and _discount_pair_is_safe(product, adjustment):
            paired.append(_combine_discount_pair(product, adjustment))
            paired_adjustments += 1
            index += 2
            continue

        if _is_explicit_discount_adjustment(current):
            current["unpaired_adjustment"] = True
            unpaired_adjustments += 1
        paired.append(current)
        index += 1
    return paired, paired_adjustments, unpaired_adjustments


def _is_explicit_discount_adjustment(line: Mapping[str, Any]) -> bool:
    description = _normalized_words(str(line.get("description") or ""))
    line_total = line.get("line_total")
    return bool(
        re.search(r"\bdiscount(?:ed)?\b", description)
        and line_total is not None
        and float(line_total) <= 0
    )


def _numbers_equal(left: Any, right: Any, tolerance: float = 1e-8) -> bool:
    if left is None or right is None:
        return False
    return abs(float(left) - float(right)) <= tolerance


def _discount_pair_is_safe(
    product: Mapping[str, Any],
    adjustment: Mapping[str, Any],
) -> bool:
    if product.get("line_total_basis") != "net" or adjustment.get("line_total_basis") != "net":
        return False
    if product.get("line_total") is None or float(product["line_total"]) <= 0:
        return False
    if not _numbers_equal(product.get("quantity"), adjustment.get("quantity")):
        return False
    product_tax = product.get("tax_rate")
    adjustment_tax = adjustment.get("tax_rate")
    if product_tax is not None or adjustment_tax is not None:
        if not _numbers_equal(product_tax, adjustment_tax):
            return False
    return True


def _add_money(left: Any, right: Any) -> float | None:
    if left is None or right is None:
        return None
    try:
        value = (Decimal(str(left)) + Decimal(str(right))).quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP,
        )
    except InvalidOperation:
        return None
    return float(value)


def _combine_discount_pair(
    product: Mapping[str, Any],
    adjustment: Mapping[str, Any],
) -> dict[str, Any]:
    combined = dict(product)
    net_line_total = _add_money(product.get("line_total"), adjustment.get("line_total"))
    quantity = product.get("quantity")
    net_unit_price = (
        _divide_money_by_quantity(net_line_total, quantity)
        if net_line_total is not None and quantity is not None
        else None
    )
    printed_price = product.get("printed_unit_price")
    if printed_price is None:
        printed_price = product.get("unit_price")
    combined.update(
        {
            "unit_price": net_unit_price,
            "net_unit_price": net_unit_price,
            "gross_unit_price": printed_price,
            "printed_unit_price": printed_price,
            "printed_unit_price_basis": "pre_discount",
            "unit_price_basis": "net",
            "unit_price_source": "derived_paired_net_line_amount_divided_by_quantity",
            "line_total": net_line_total,
            "line_total_basis": "net_after_explicit_discount",
            "gross_line_total": _add_money(
                product.get("gross_line_total"), adjustment.get("gross_line_total")
            ),
            "tax_amount": _add_money(product.get("tax_amount"), adjustment.get("tax_amount")),
            "discount_line_total": adjustment.get("line_total"),
            "discount_tax_amount": adjustment.get("tax_amount"),
            "discount_gross_line_total": adjustment.get("gross_line_total"),
            "source_adjustments": [
                {
                    field: adjustment.get(field)
                    for field in (
                        "description",
                        "quantity",
                        "printed_unit_price",
                        "line_total",
                        "gross_line_total",
                        "tax_rate",
                        "tax_amount",
                        "source_page",
                        "source_rows",
                        "source_sequence",
                    )
                }
            ],
            "paired_source_line_count": 2,
            "extraction_method": "layout_header+paired_discount",
            "derived_fields": ["line_total", "unit_price", "net_unit_price"],
            "confidence": min(float(product.get("confidence") or 0), float(adjustment.get("confidence") or 0)),
        }
    )
    product_rows = product.get("source_rows") or []
    adjustment_rows = adjustment.get("source_rows") or []
    if product_rows and adjustment_rows:
        combined["source_rows"] = [
            min(int(product_rows[0]), int(adjustment_rows[0])),
            max(int(product_rows[-1]), int(adjustment_rows[-1])),
        ]
    return combined


def _layout_cell_boundaries(line: str, anchors: Sequence[_ColumnAnchor]) -> list[int]:
    boundaries = [0]
    for left, right in zip(anchors, anchors[1:]):
        if left.field in {"description", "upc", "item_code"} or (
            left.field == "quantity" and right.field == "unit_price"
        ):
            boundary = right.start
        elif right.field == "description":
            boundary = int(round((left.end + right.start) / 2))
        else:
            boundary = int(round((left.end + right.start) / 2))
        boundaries.append(boundary)
    boundaries.append(max(len(line), anchors[-1].end + 1))
    return boundaries


def _slice_layout_cells(
    line: str,
    anchors: Sequence[_ColumnAnchor],
) -> dict[str, str]:
    boundaries = _layout_cell_boundaries(line, anchors)
    return {
        anchor.field: line[boundaries[index] : boundaries[index + 1]].strip()
        for index, anchor in enumerate(anchors)
    }


def _layout_cells_misaligned(
    line: str,
    anchors: Sequence[_ColumnAnchor],
    cells: Mapping[str, str],
) -> bool:
    """True when fixed-position slicing evidently did not land on this row's cells.

    Rows typed as space-separated text under a heading are not column
    aligned; slicing them at heading offsets cuts "250ml" into a quantity
    cell "250" and a UOM cell "ml", which is how a size token became a
    quantity.  A boundary inside a printed token, or a quantity cell holding
    letters instead of a number, are both evidence of misalignment.
    """

    for boundary in _layout_cell_boundaries(line, anchors)[1:-1]:
        if 0 < boundary < len(line) and not line[boundary - 1].isspace() and not line[boundary].isspace():
            return True
    quantity_cell = cells.get("quantity") or ""
    return bool(
        quantity_cell
        and _parse_cell_number(quantity_cell) is None
        and re.search(r"[A-Za-z]", quantity_cell)
    )


def _parse_cell_number(value: Any) -> float | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    candidate = unicodedata.normalize("NFKC", str(value)).strip()
    candidate = re.sub(r"(?i)\b(?:AED|USD|EUR|GBP|SAR|QAR|KWD|BHD|OMR)\b", "", candidate)
    candidate = candidate.replace("%", "").strip()
    if re.search(r"[A-Za-z]", candidate):
        return None
    compact_pattern = r"\(?-?\d[\d,.']*\)?"
    spaced_pattern = r"\(?-?\d{1,3}(?:[ \u00a0]\d{3})+(?:[.,]\d+)?\)?"
    if not (re.fullmatch(compact_pattern, candidate) or re.fullmatch(spaced_pattern, candidate)):
        return None
    return _parse_number(candidate)


def _parse_layout_number(value: Any) -> float | None:
    """Parse a bounded numeric cell while tolerating OCR border glyphs.

    Fixed-position slicing supplies the cell boundary.  When Tesseract joins a
    rule character or a neighbouring zero to the cell, the rightmost explicit
    numeric token is retained; no digits are repaired or synthesized.
    """

    normalized_value = (
        unicodedata.normalize("NFKC", str(value)).replace("§", "5")
        if value not in (None, "")
        else value
    )
    parsed = _parse_cell_number(normalized_value)
    if parsed is not None or value in (None, ""):
        return parsed
    tokens = re.findall(r"\(?-?\d[\d,.']*\)?", str(normalized_value))
    for token in reversed(tokens):
        parsed = _parse_number(token)
        if parsed is not None:
            return parsed
    return None


def _rate_candidates(raw_value: Any, parsed_value: float | None) -> list[float]:
    candidates: list[float] = []
    if parsed_value is not None:
        candidates.append(float(parsed_value))
    compact = re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(raw_value or "")))
    match = re.fullmatch(r"[^0-9-]*(-?\d+),(\d{3})[^0-9]*", compact)
    if match:
        decimal_candidate = float(f"{match.group(1)}.{match.group(2)}")
        if decimal_candidate not in candidates:
            candidates.append(decimal_candidate)
    return candidates


def _extension_matches(quantity: float, rate: float, extension: float) -> bool:
    try:
        calculated = (Decimal(str(quantity)) * Decimal(str(rate))).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )
        source = Decimal(str(extension)).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )
    except (InvalidOperation, TypeError, ValueError):
        return False
    return calculated == source


def _one_edit_apart(left: str, right: str) -> bool:
    if left == right:
        return True
    if abs(len(left) - len(right)) > 1:
        return False
    if len(left) == len(right):
        return sum(a != b for a, b in zip(left, right)) <= 1
    shorter, longer = (left, right) if len(left) < len(right) else (right, left)
    for index in range(len(longer)):
        if longer[:index] + longer[index + 1 :] == shorter:
            return True
    return False


def _validated_ocr_quantity(
    raw_quantity: Any,
    raw_rate: Any,
    parsed_quantity: float | None,
    parsed_rate: float | None,
    line_total: float | None,
    gross_line_total: float | None,
) -> float | None:
    """Correct one OCR digit only when two printed money fields validate it."""

    rates = [rate for rate in _rate_candidates(raw_rate, parsed_rate) if rate > 0]
    extensions = {
        float(value)
        for value in (line_total, gross_line_total)
        if value is not None and float(value) >= 0
    }
    if not rates or not extensions:
        return None
    if parsed_quantity is not None and any(
        _extension_matches(parsed_quantity, rate, extension)
        for rate in rates
        for extension in extensions
    ):
        return None

    raw_compact = re.sub(r"[^A-Za-z0-9]", "", str(raw_quantity or ""))
    translated = raw_compact.translate(
        str.maketrans({"O": "0", "o": "0", "I": "1", "i": "1", "l": "1", "T": "1"})
    )
    raw_digits = re.sub(r"\D", "", raw_compact)
    candidates: set[int] = set()
    for rate in rates:
        for extension in extensions:
            quotient = Decimal(str(extension)) / Decimal(str(rate))
            integral = quotient.to_integral_value(rounding=ROUND_HALF_UP)
            if abs(quotient - integral) <= Decimal("0.000001") and integral > 0:
                candidate = int(integral)
                candidate_text = str(candidate)
                if translated.isdigit() and translated == candidate_text:
                    candidates.add(candidate)
                elif raw_digits and _one_edit_apart(raw_digits, candidate_text):
                    candidates.add(candidate)
    if len(candidates) == 1:
        return float(next(iter(candidates)))
    return None


def _line_extension_from_consistent_source_columns(
    raw_rate: Any,
    quantity: float | None,
    parsed_rate: float | None,
    line_total: float | None,
    gross_line_total: float | None,
    discount_rate: float | None,
    discount_amount: float | None,
) -> float | None:
    """Use Gross as Net only when explicit zero-discount arithmetic proves it."""

    if (
        quantity is None
        or gross_line_total is None
        or line_total is None
        or (discount_rate not in (None, 0, 0.0))
        or (discount_amount not in (None, 0, 0.0))
    ):
        return None
    rates = _rate_candidates(raw_rate, parsed_rate)
    gross_matches = any(
        _extension_matches(quantity, rate, gross_line_total) for rate in rates
    )
    net_matches = any(_extension_matches(quantity, rate, line_total) for rate in rates)
    if gross_matches and not net_matches:
        return float(gross_line_total)
    return None


def _divide_money_by_quantity(amount: float, quantity: float) -> float | None:
    try:
        decimal_amount = Decimal(str(amount)).quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP,
        )
        decimal_quantity = Decimal(str(quantity))
        if decimal_quantity == 0:
            return None
        value = (decimal_amount / decimal_quantity).quantize(
            Decimal("0.00000001"),
            rounding=ROUND_HALF_UP,
        )
        if (value * decimal_quantity).quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP,
        ) != decimal_amount:
            return None
    except (InvalidOperation, ZeroDivisionError):
        return None
    return float(value)


def _parse_quantity_cell(value: Any) -> tuple[float | None, str | None]:
    number = _parse_cell_number(value)
    if number is not None:
        return number, None
    if value in (None, ""):
        return None, None
    candidate = unicodedata.normalize("NFKC", str(value)).strip()
    if re.fullmatch(rf"(?:\d+\s*[xX\u00d7]\s*)?{_NUM}{_SIZE_UNIT}", candidate, re.I) or re.fullmatch(
        rf"\d+\s*[xX\u00d7]\s*{_NUM}", candidate, re.I
    ):
        # "250ml", "1kg", "6x330ml": a size or pack token, not a quantity.
        return None, None
    match = re.fullmatch(
        rf"(?P<number>{_NUM})\s*(?P<uom>{_UOM})",
        candidate,
        re.I,
    )
    if not match:
        match = re.fullmatch(
            rf"(?P<uom>{_UOM})\s*(?P<number>{_NUM})",
            candidate,
            re.I,
        )
    if not match:
        if re.fullmatch(_PACK_SIZE_TOKEN, candidate, re.I):
            return None, None
        return _parse_layout_number(candidate), None
    if _is_size_unit(match.group("uom")) and not re.search(r"\d\s+[A-Za-z]", candidate):
        # "250ml" alone is a pack size that spilled out of the description,
        # not a quantity of 250 millilitres.  A spaced "2 KG" is still read
        # as a printed quantity with its unit.
        return None, None
    return _parse_number(match.group("number")), _clean_uom(match.group("uom"))


def _clean_description(value: Any) -> str | None:
    if value in (None, ""):
        return None
    cleaned = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip(" |:;-_")
    if len(cleaned) < 2 or not any(character.isalpha() for character in cleaned):
        return None
    return cleaned


def _join_description_parts(parts: Sequence[str]) -> str | None:
    cleaned = [part for part in (_clean_description(value) for value in parts) if part]
    if not cleaned:
        return None
    return re.sub(r"\s+", " ", " ".join(cleaned)).strip()


def _is_table_summary(value: str) -> bool:
    normalized = _normalized_words(value)
    return bool(
        re.match(
            r"^(?:opening(?:\s+balance)?|closing(?:\s+balance)?|balance|sub\s*total|grand\s+total|invoice\s+total|(?:t?otal|otal)(?:\s+(?:qty|quantity|pieces?|cartons?|amount|value|vat|tax))?|tax\s+total|vat\s+total|amount\s+due|total\s+due|amount\s+paid)\b",
            normalized,
        )
    )


def _extract_legacy_lines(text: str) -> list[dict[str, Any]]:
    source_lines = [line.rstrip() for line in text.splitlines()]
    header_index: int | None = None
    for index, line in enumerate(source_lines):
        normalized = _normalized_words(line)
        has_description = bool(re.search(r"\b(?:description|item|product)\b", normalized))
        has_quantity = bool(re.search(r"\b(?:qty|quantity)\b", normalized))
        has_value = bool(re.search(r"\b(?:amount|total|price|rate)\b", normalized))
        if has_description and has_quantity and has_value:
            header_index = index
            break

    extracted: list[dict[str, Any]] = []
    scan = source_lines[header_index + 1 :] if header_index is not None else source_lines
    for offset, line in enumerate(scan):
        stripped = line.strip()
        if not stripped:
            continue
        if _is_summary_label(stripped):
            if header_index is not None and extracted:
                break
            continue
        parsed = _parse_line_candidate(line, in_table=header_index is not None)
        if parsed is None:
            continue
        parsed["confidence"] = 0.91 if header_index is not None else 0.8
        parsed.update(
            {
                "net_unit_price": None,
                "gross_unit_price": None,
                "printed_unit_price": parsed.get("unit_price"),
                "printed_unit_price_basis": "unspecified",
                "unit_price_basis": "unspecified",
                "unit_price_source": "printed_unit_price",
                "line_total_basis": "unspecified",
                "tax_amount": None,
                "discount_rate": None,
                "discount_amount": None,
                "upc": None,
                "item_code": None,
                "source_page": 1,
                "source_rows": [
                    (header_index + 2 + offset) if header_index is not None else offset + 1,
                    (header_index + 2 + offset) if header_index is not None else offset + 1,
                ],
                "extraction_method": "legacy_pattern",
                "derived_fields": [],
            }
        )
        parsed["id"] = _line_id(len(extracted), parsed)
        extracted.append(parsed)
    return extracted


_NUM = r"-?\(?\d[\d,.']*\)?"
_MONEY = rf"(?:[A-Z]{{3}}\s*)?{_NUM}(?:\s*[A-Z]{{3}})?"
# Quantity units describe how many of an item were supplied.  Pack-size
# units (250ml, 500g, 1.5L) describe the item itself and belong to the
# description; a size unit glued to a number is never a quantity on its own.
_QUANTITY_UOM = (
    r"(?:EA|EACH|PC|PCS|PIECE|PIECES|UNIT|UNITS|PK|PKG|PACK|PACKS|BX|BOX|CS|CASE|"
    r"CTN|CARTON|CARTONS|BTL|BOTTLE|BOTTLES|DZ|DOZ|DOZEN|SET|SETS|ROLL|ROLLS|BAG|BAGS|TUB|TUBS|JAR|JARS|TIN|TINS|CAN|CANS|SACHET|SACHETS|TRAY|TRAYS|PAIR|PAIRS|NOS?)"
)
_UOM = rf"(?:{_QUANTITY_UOM}|{_SIZE_UNIT})"
_SIZE_TOKEN = rf"\d+(?:[.,]\d+)?\s*{_SIZE_UNIT}"
_PACK_SIZE_TOKEN = rf"(?:\d+\s*[xX×]\s*)?{_SIZE_TOKEN}"


def _is_quantity_uom(value: Any) -> bool:
    cleaned = _clean_uom(value)
    return bool(cleaned and re.fullmatch(_QUANTITY_UOM, cleaned, re.I))


def _is_size_unit(value: Any) -> bool:
    cleaned = _clean_uom(value)
    return bool(cleaned and re.fullmatch(_SIZE_UNIT, cleaned, re.I))


def _parse_line_candidate(line: str, *, in_table: bool) -> dict[str, Any] | None:
    separator = r"\s+" if in_table else r"(?:\t+|\s{2,})"
    patterns = [
        re.compile(
            rf"^\s*(?:\d{{1,4}}[.)]?\s+)?(?P<description>.+?){separator}"
            rf"(?P<quantity>{_NUM})\s+(?P<uom>{_UOM})\s+"
            rf"(?P<unit_price>{_MONEY})\s+(?P<line_total>{_MONEY})"
            rf"(?:\s+(?P<tax_rate>\d+(?:[.,]\d+)?\s*%))?\s*$",
            re.I,
        ),
        re.compile(
            rf"^\s*(?:\d{{1,4}}[.)]?\s+)?(?P<description>.+?){separator}"
            rf"(?P<quantity>{_NUM})\s+(?P<unit_price>{_MONEY})\s+(?P<line_total>{_MONEY})"
            rf"(?:\s+(?P<tax_rate>\d+(?:[.,]\d+)?\s*%))?\s*$",
            re.I,
        ),
    ]
    for pattern in patterns:
        match = pattern.match(line)
        if not match:
            continue
        description = re.sub(r"\s+", " ", match.group("description")).strip(" |")
        if len(description) < 2 or _is_summary_label(description):
            continue
        values = match.groupdict()
        quantity = _parse_number(values.get("quantity"))
        unit_price = _parse_number(values.get("unit_price"))
        line_total = _parse_number(values.get("line_total"))
        if quantity is None or unit_price is None or line_total is None:
            continue
        return {
            "description": description,
            "quantity": quantity,
            "unit_price": unit_price,
            "line_total": line_total,
            "tax_rate": _parse_percentage(values.get("tax_rate")),
            "uom": _clean_uom(values.get("uom")),
        }
    return None


def _parse_date(value: str) -> tuple[str | None, str | None]:
    cleaned = re.sub(r"\s+", " ", value.strip()).rstrip(".,")
    # ISO is intentionally handled first and is unambiguous.
    iso_match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})\b", cleaned)
    if iso_match:
        try:
            return datetime(*map(int, iso_match.groups())).date().isoformat(), None
        except ValueError:
            return None, f"Invoice date '{cleaned}' is invalid and was withheld."
    for fmt in (
        "%d %B %Y",
        "%d %b %Y",
        "%B %d %Y",
        "%b %d %Y",
        "%d-%B-%Y",
        "%d-%b-%Y",
    ):
        try:
            return datetime.strptime(cleaned.replace(",", ""), fmt).date().isoformat(), None
        except ValueError:
            pass
    numeric = re.match(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2}|\d{4})\b", cleaned)
    if numeric:
        first, second, year = (int(part) for part in numeric.groups())
        year = year + 2000 if year < 100 else year
        if first <= 12 and second <= 12:
            return None, f"Invoice date '{numeric.group(0)}' is day/month ambiguous and was withheld."
        day, month = (first, second) if first > 12 else (second, first)
        try:
            return datetime(year, month, day).date().isoformat(), None
        except ValueError:
            return None, f"Invoice date '{numeric.group(0)}' is invalid and was withheld."
    return None, f"Invoice date '{cleaned}' was not recognized and was withheld."


def _parse_currency(value: str) -> str | None:
    upper = value.strip().upper()
    aliases = {
        "DIRHAM": "AED",
        "DIRHAMS": "AED",
        "UAE DIRHAM": "AED",
        "US DOLLAR": "USD",
        "EURO": "EUR",
        "POUND": "GBP",
    }
    if upper in _CURRENCIES:
        return upper
    return aliases.get(upper)


def _parse_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = unicodedata.normalize("NFKC", str(value)).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"(?i)\b(?:AED|USD|EUR|GBP|SAR|QAR|KWD|BHD|OMR)\b", "", text)
    text = text.replace("\u00a0", "").replace(" ", "").replace("'", "")
    text = re.sub(r"[^0-9,.-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        pieces = text.split(",")
        if len(pieces) == 2 and len(pieces[1]) in {1, 2}:
            text = ".".join(pieces)
        else:
            text = "".join(pieces)
    elif text.count(".") > 1:
        pieces = text.split(".")
        text = "".join(pieces[:-1]) + "." + pieces[-1]
    try:
        number = float(text)
    except ValueError:
        return None
    return -abs(number) if negative else number


def _last_number(value: str) -> float | None:
    tokens = re.findall(r"(?:[A-Z]{3}\s*)?\(?-?\d[\d\s,.']*\)?(?:\s*[A-Z]{3})?", value, re.I)
    for token in reversed(tokens):
        number = _parse_number(token)
        if number is not None:
            return number
    return None


def _parse_percentage(value: Any) -> float | None:
    return _parse_number(value)


def _clean_uom(value: Any) -> str | None:
    if value in (None, ""):
        return None
    cleaned = re.sub(r"[^A-Za-z]", "", str(value)).upper()
    return cleaned or None


def _clean_upc(value: Any) -> str | None:
    if value in (None, ""):
        return None
    cleaned = re.sub(r"[^0-9]", "", unicodedata.normalize("NFKC", str(value)))
    return cleaned if 6 <= len(cleaned) <= 18 else None


def _valid_gtin(value: str) -> bool:
    if not value.isdigit() or len(value) not in {8, 12, 13, 14}:
        return False
    body = value[:-1]
    weighted = sum(
        int(digit) * (3 if (len(body) - index) % 2 else 1)
        for index, digit in enumerate(body)
    )
    return (10 - weighted % 10) % 10 == int(value[-1])


def _clean_item_code(value: Any) -> str | None:
    if value in (None, ""):
        return None
    cleaned = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip(" |:;,.")
    if not 2 <= len(cleaned) <= 50 or not any(character.isalnum() for character in cleaned):
        return None
    return cleaned


def _canonical_column(value: Any) -> str | None:
    if value in (None, ""):
        return None
    compact = re.sub(r"[^a-z0-9]", "", str(value).lower())
    return _COLUMN_ALIASES.get(compact)


def _unique_nonempty(values: Iterable[Any]) -> list[Any]:
    unique: dict[str, Any] = {}
    for value in values:
        if value in (None, ""):
            continue
        key = unicodedata.normalize("NFKC", str(value)).strip().casefold()
        unique.setdefault(key, value)
    return list(unique.values())


def _normalized_words(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", unicodedata.normalize("NFKC", value).lower()).strip()


def _is_summary_label(value: str) -> bool:
    normalized = _normalized_words(value)
    return bool(
        re.match(
            r"^(?:sub total|subtotal|grand total|invoice total|total|tax|vat|amount due|total due|amount paid)\b",
            normalized,
        )
    )


_NET_UNIT_PRICE_ASSUMPTION_SOURCE = "printed_unit_price_without_tax_inclusive_indication"


def resolve_net_unit_prices(lines: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Carry a usable net unit price on every line that printed a plain rate.

    Extraction records exactly what the document says: a "Unit Price" column
    has basis ``unit`` and no net price is asserted.  Downstream costing needs
    a net unit price, so when the document gives no gross or tax-inclusive
    indication the printed rate is treated as net.  The derivation is recorded
    in ``derived_fields`` and ``unit_price_source`` so it stays auditable.
    Lines that already carry a net price, or that printed a gross rate, are
    returned unchanged.
    """

    resolved: list[dict[str, Any]] = []
    for source in lines:
        line = dict(source)
        unit_price = line.get("unit_price")
        if (
            line.get("net_unit_price") is None
            and unit_price is not None
            and line.get("gross_unit_price") is None
            and line.get("printed_unit_price_basis") not in {"gross", "pre_discount"}
            and line.get("unit_price_basis") not in {"gross", "pre_discount"}
            and line.get("line_total_basis") != "gross"
        ):
            line["net_unit_price"] = unit_price
            line["unit_price_basis"] = "net"
            line["unit_price_source"] = _NET_UNIT_PRICE_ASSUMPTION_SOURCE
            derived = [str(field) for field in (line.get("derived_fields") or [])]
            if "net_unit_price" not in derived:
                derived.append("net_unit_price")
            line["derived_fields"] = derived
        resolved.append(line)
    return resolved


def _line_id(index: int, line: Mapping[str, Any]) -> str:
    evidence = "|".join(
        str(line.get(field, ""))
        for field in ("description", "quantity", "uom", "unit_price", "line_total", "tax_rate")
    )
    digest = hashlib.sha256(f"{index}|{evidence}".encode("utf-8")).hexdigest()[:16]
    return f"line-{digest}"


def _clean_source_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("Text file encoding is not UTF-8, UTF-16, or Windows-1252.")


def _check_page_count(count: int, limits: ExtractionLimits) -> None:
    if count > limits.max_pages:
        raise ExtractionLimitError(f"Document has {count} pages; the limit is {limits.max_pages}.")


def _check_page_pixels(pixels: int, page_number: int, limits: ExtractionLimits) -> None:
    if pixels > limits.max_pixels_per_page:
        raise ExtractionLimitError(
            f"Page {page_number} has {pixels} pixels; the per-page limit is {limits.max_pixels_per_page}."
        )


def _check_zip_archive(path: Path, limits: ExtractionLimits) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > limits.max_archive_members:
                raise ExtractionLimitError(
                    f"Archive has {len(members)} members; the limit is {limits.max_archive_members}."
                )
            total = sum(member.file_size for member in members)
            if total > limits.max_archive_uncompressed_bytes:
                raise ExtractionLimitError(
                    f"Archive expands to {total} bytes; the limit is {limits.max_archive_uncompressed_bytes}."
                )
            for member in members:
                compressed = max(1, member.compress_size)
                if member.file_size > 1_000_000 and member.file_size / compressed > limits.max_archive_compression_ratio:
                    raise ExtractionLimitError(
                        f"Archive member '{member.filename}' exceeds the compression-ratio safety limit."
                    )
    except ExtractionLimitError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentExtractionError(f"The archive-based document is corrupt or unreadable: {exc}") from exc


def _dedupe(values: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def _document_confidence(result: Mapping[str, Any]) -> float:
    values = [float(value) for value in result.get("field_confidence", {}).values()]
    values.extend(float(line.get("confidence", 0)) for line in result.get("lines", []))
    if not values:
        return 0.0
    return round(sum(values) / len(values), 3)


__all__ = [
    "DocumentExtractionError",
    "ExtractionLimitError",
    "ExtractionLimits",
    "SUPPORTED_EXTENSIONS",
    "UnsupportedDocumentError",
    "extract_document",
]
