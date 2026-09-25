from __future__ import annotations

from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook


SCHEMA_NAME = "RMS CONSOLIDATED INVOICE TARGET v1"
HEADER_COLUMNS = [
    "Transaction Number",
    "Document",
    "Supplier Site",
    "Order No",
    "Location",
    "Location Type",
    "Document Date",
    "Total Cost Ex Tax",
    "Tax Amount",
    "Ref No. 1",
    "Ref No. 2",
    "Ref No. 3",
    "Comment",
]
TAX_COLUMNS = ["Transaction Number", "Tax Code", "Tax Basis"]
DETAIL_COLUMNS = [
    "Transaction Number",
    "Item",
    "UPC",
    "Unit Cost",
    "Quantity",
    "Unit Tax Code",
]


def safe_text(value: Any) -> Any:
    """Preserve identifier text while preventing spreadsheet formula execution."""

    if not isinstance(value, str):
        return value
    value = value.replace("\x00", "")[:32767]
    if value.lstrip(" \t\r\n\ufeff").startswith(("=", "+", "-", "@")):
        return ("'" + value)[:32767]
    return value


def _append(sheet: Any, values: Iterable[Any]) -> None:
    sheet.append([safe_text(value) for value in values])


def _identifier(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\x00", "").strip()
    return text or None


def _number(value: Any) -> int | float | None:
    if value in (None, ""):
        return None
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _document_date(value: Any) -> date | None:
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


def _valid_order_number(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    compact = " ".join(text.upper().replace("_", " ").split())
    if "PO BOX" in compact or compact in {"ORDERABLE IND", "MIN ORDER QTY"}:
        return None
    return text


def build_target_workbook(
    invoices: list[dict[str, Any]], *, include_upc_in_export: bool = False
) -> bytes:
    """Build the exact three-sheet consolidated target workbook."""

    workbook = Workbook()
    header = workbook.active
    header.title = "Header"
    tax = workbook.create_sheet("Tax_Breakdown")
    details = workbook.create_sheet("Details")
    _append(header, HEADER_COLUMNS)
    _append(tax, TAX_COLUMNS)
    _append(details, DETAIL_COLUMNS)

    for index, invoice in enumerate(invoices, start=1):
        transaction_number = str(index)
        order_number = _valid_order_number(invoice.get("po_number"))
        if not order_number:
            master_orders = {
                value
                for line in invoice.get("lines", [])
                if (value := _valid_order_number(line.get("rms_po_number")))
            }
            if len(master_orders) == 1:
                order_number = next(iter(master_orders))
        target_total = invoice.get("subtotal")
        _append(
            header,
            [
                transaction_number,
                _identifier(invoice.get("document") or invoice.get("invoice_number")),
                _identifier(invoice.get("supplier_site")),
                _identifier(order_number),
                _identifier(invoice.get("location")),
                _identifier(invoice.get("location_type")),
                _document_date(invoice.get("invoice_date")),
                _number(target_total),
                _number(invoice.get("tax_total")),
                invoice.get("ref_no_1"),
                invoice.get("ref_no_2"),
                invoice.get("ref_no_3"),
                invoice.get("comment"),
            ],
        )
        _append(tax, [transaction_number, invoice.get("tax_code"), _number(target_total)])
        for line in invoice.get("lines", []):
            _append(
                details,
                [
                    transaction_number,
                    _identifier(line.get("rms_parent_item") or line.get("rms_item_id")),
                    (
                        _identifier(line.get("rms_upc"))
                        if include_upc_in_export
                        else None
                    ),
                    _number(line.get("target_unit_cost")),
                    _number(line.get("quantity")),
                    invoice.get("tax_code"),
                ],
            )

    for cell in header["G"][1:]:
        cell.number_format = "m/d/yyyy"
    for sheet, columns in (
        (header, (1, 2, 3, 4, 5, 6, 10, 11, 12, 13)),
        (tax, (1, 2)),
        (details, (1, 2, 3, 6)),
    ):
        for column in columns:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row=row, column=column).number_format = "@"
    for sheet, columns in ((header, (8, 9)), (tax, (3,)), (details, (4, 5))):
        for column in columns:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row=row, column=column).number_format = "0.00####"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_review_workbook(
    invoices: list[dict[str, Any]],
    *,
    export_id: str | None = None,
    created_at: str | None = None,
    include_upc_in_export: bool = False,
) -> bytes:
    """Backward-compatible callable; all export paths now emit the exact target."""

    del export_id, created_at
    return build_target_workbook(
        invoices, include_upc_in_export=include_upc_in_export
    )
