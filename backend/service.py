from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import re
import shutil
import threading
import time
import uuid
import zipfile
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook
from PIL import Image, ImageDraw

from .db import Database, INVOICE_STATUSES, json_dumps, utc_now
from .exporter import SCHEMA_NAME, build_target_workbook

logger = logging.getLogger("invoice_studio.service")


DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
    ".csv",
    ".xlsx",
    ".txt",
    ".docx",
}
CATALOG_EXTENSIONS = {".csv", ".xlsx"}
MATCH_STATUSES = {"unmatched", "suggested", "auto", "confirmed"}
DOCUMENT_TYPES = {"invoice", "credit_note", "purchase_order", "delivery_note", "unknown"}
COST_POLICY_MODES = {"invoice_only"}
# ISO 4217 minor-unit exponents used by the money gates.  The previous
# implementation rounded every currency to cents, which silently discarded
# fils for KWD/BHD/OMR (and dinars with three minor units).  Keep the map
# explicit and fail closed to the conventional two-decimal currency when a
# code is unknown.
CURRENCY_MINOR_UNITS = {
    "BHD": 3,
    "IQD": 3,
    "JOD": 3,
    "KWD": 3,
    "LYD": 3,
    "OMR": 3,
    "TND": 3,
    "CLF": 4,
    "BIF": 0,
    "CLP": 0,
    "DJF": 0,
    "GNF": 0,
    "JPY": 0,
    "KMF": 0,
    "KRW": 0,
    "MGA": 0,
    "PYG": 0,
    "RWF": 0,
    "UGX": 0,
    "VND": 0,
    "VUV": 0,
    "XAF": 0,
    "XOF": 0,
    "XPF": 0,
}
TRUSTED_MIME_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".bmp": "image/bmp",
    ".webp": "image/webp",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ServiceError(Exception):
    pass


class NotFound(ServiceError):
    pass


class Conflict(ServiceError):
    def __init__(self, message: str, *, current_version: int | None = None):
        super().__init__(message)
        self.current_version = current_version


class ValidationFailure(ServiceError):
    def __init__(self, errors: list[dict[str, str]]):
        super().__init__("validation failed")
        self.errors = errors


class UploadRejected(ServiceError):
    pass


@dataclass(slots=True)
class Settings:
    database_path: Path
    source_dir: Path
    export_dir: Path
    workers: int = 2
    max_file_bytes: int = 25 * 1024 * 1024
    max_catalog_file_bytes: int = 128 * 1024 * 1024
    max_upload_files: int = 1000
    max_pages: int = 50
    max_attempts: int = 3
    poll_seconds: float = 0.25
    db_busy_timeout_seconds: float = 30.0
    worker_stall_seconds: float = 120.0

    @classmethod
    def from_env(cls) -> "Settings":
        data_dir = Path(os.getenv("INVOICE_DATA_DIR", "data")).resolve()
        return cls(
            database_path=Path(os.getenv("INVOICE_DB_PATH", data_dir / "invoices.sqlite3")),
            source_dir=Path(os.getenv("INVOICE_SOURCE_DIR", data_dir / "sources")),
            export_dir=Path(os.getenv("INVOICE_EXPORT_DIR", data_dir / "exports")),
            workers=max(0, min(64, int(os.getenv("INVOICE_WORKERS", "4")))),
            max_file_bytes=int(os.getenv("INVOICE_MAX_FILE_BYTES", str(25 * 1024 * 1024))),
            max_catalog_file_bytes=int(
                os.getenv("INVOICE_MAX_CATALOG_FILE_BYTES", str(128 * 1024 * 1024))
            ),
            max_upload_files=int(os.getenv("INVOICE_MAX_UPLOAD_FILES", "1000")),
            max_pages=int(os.getenv("INVOICE_MAX_PAGES", "50")),
            max_attempts=max(1, int(os.getenv("INVOICE_MAX_ATTEMPTS", "3"))),
            poll_seconds=max(0.05, float(os.getenv("INVOICE_POLL_SECONDS", "0.25"))),
            db_busy_timeout_seconds=max(
                0.05, float(os.getenv("INVOICE_DB_BUSY_TIMEOUT_SECONDS", "30"))
            ),
            worker_stall_seconds=max(
                5.0, float(os.getenv("INVOICE_WORKER_STALL_SECONDS", "120"))
            ),
        )


def _load_json(value: str | None, fallback: Any) -> Any:
    if not value:
        return fallback
    try:
        return json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return fallback


def _decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise InvalidOperation
    number = Decimal(str(value).replace(",", "").strip())
    if not number.is_finite():
        raise InvalidOperation
    return number


def _decimal_text(value: Any) -> str | None:
    number = _decimal(value)
    return None if number is None else format(number, "f")


def _money_quantum(currency: str | None = None) -> Decimal:
    exponent = CURRENCY_MINOR_UNITS.get(str(currency or "").upper(), 2)
    return Decimal(1).scaleb(-exponent)


def _money_equal(left: Decimal, right: Decimal, currency: str | None = None) -> bool:
    return _money_round(left, currency) == _money_round(right, currency)


def _money_round(value: Decimal, currency: str | None = None) -> Decimal:
    return value.quantize(_money_quantum(currency), rounding=ROUND_HALF_UP)


def _number(value: Any) -> float | int | None:
    if value is None or value == "":
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        return None
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def normalize_description(value: str) -> str:
    value = value.casefold()
    value = re.sub(r"[^\w]+", " ", value, flags=re.UNICODE)
    return " ".join(value.split())


def _identifier_text(value: Any, *, limit: int = 200) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.replace("\x00", "").strip()
    return text[:limit] or None


def _normalize_upc(value: Any) -> str | None:
    text = _identifier_text(value)
    if not text:
        return None
    if text.upper().startswith("ULT_"):
        text = text[4:]
    return text or None


def _valid_order_number(value: Any) -> str | None:
    text = _identifier_text(value)
    if not text:
        return None
    compact = " ".join(text.upper().replace("_", " ").split())
    if "PO BOX" in compact or compact in {"ORDERABLE IND", "MIN ORDER QTY"}:
        return None
    return text


def _catalog_item_key(
    *, supplier_id: str | None, rms_item_id: str, upc: str | None, uom: str | None, occurrence: int
) -> str:
    identity = json_dumps(
        [supplier_id or "", rms_item_id, upc or "", (uom or "").casefold(), occurrence]
    )
    return "catalog:" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:32]


def _clean_filename(filename: str | None) -> str:
    clean = Path((filename or "upload").replace("\x00", "")).name.strip()
    if not clean or clean in {".", ".."}:
        clean = "upload"
    return clean[:240]


def _spreadsheet_safe_csv_text(value: Any) -> str:
    """Return plain CSV text that Excel cannot interpret as a formula."""

    text = str(value or "").replace("\x00", "")
    if text.lstrip(" \t\r\n\ufeff").startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


class InvoiceService:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.settings.source_dir.mkdir(parents=True, exist_ok=True)
        self.settings.export_dir.mkdir(parents=True, exist_ok=True)
        self.db = Database(
            settings.database_path,
            busy_timeout_seconds=settings.db_busy_timeout_seconds,
        )
        self.db.initialize()
        self._stop = threading.Event()
        self._wake = threading.Event()
        self._threads: list[threading.Thread] = []
        self._worker_lock = threading.Lock()
        self._worker_state: dict[str, Any] = {
            "started_at": utc_now(),
            "last_claim_at": None,
            "last_complete_at": None,
            "last_error": None,
            "last_error_at": None,
            "error_count": 0,
            "thread_exits": 0,
            "respawns": 0,
        }
        self._orphaned_jobs: set[str] = set()

    def _settings_from_row(self, row: Mapping[str, Any]) -> dict[str, Any]:
        policy = _load_json(row["target_cost_policy_json"], {})
        tolerance = _number(policy.get("maximum_absolute_difference_aed"))
        return {
            "version": row["version"],
            "brand_label": row["brand_label"],
            "location": row["location"],
            "location_type": row["location_type"],
            "supplier_rules": _load_json(row["supplier_rules_json"], []),
            "include_upc_in_export": bool(row["include_upc_in_export"]),
            "target_cost_policy": {
                "mode": "invoice_only",
                "maximum_absolute_difference_aed": 10 if tolerance is None else tolerance,
            },
            "updated_at": row["updated_at"],
        }

    def _get_settings_row(self, conn: Any) -> Any:
        row = conn.execute("SELECT * FROM app_settings WHERE id = 1").fetchone()
        if row is None:  # defensive for a manually altered database
            raise ServiceError("brand settings row is missing")
        return row

    def get_settings(self) -> dict[str, Any]:
        with self.db.connection() as conn:
            return self._settings_from_row(self._get_settings_row(conn))

    def update_settings(
        self, expected_version: int, changes: Mapping[str, Any]
    ) -> dict[str, Any]:
        errors: list[dict[str, str]] = []

        def clean(field: str, limit: int = 200) -> str:
            value = changes.get(field, "")
            if value is None:
                return ""
            if not isinstance(value, str):
                errors.append(
                    {"field": field, "code": "invalid", "message": "must be text"}
                )
                return ""
            return value.replace("\x00", "").strip()[:limit]

        brand_label = clean("brand_label")
        location = clean("location")
        location_type = clean("location_type")
        raw_rules = changes.get("supplier_rules", [])
        supplier_rules: list[dict[str, str]] = []
        seen_suppliers: set[str] = set()
        if not isinstance(raw_rules, list):
            errors.append(
                {
                    "field": "supplier_rules",
                    "code": "invalid",
                    "message": "must be a list",
                }
            )
        else:
            for index, raw_rule in enumerate(raw_rules):
                if not isinstance(raw_rule, Mapping):
                    errors.append(
                        {
                            "field": f"supplier_rules.{index}",
                            "code": "invalid",
                            "message": "must be an object",
                        }
                    )
                    continue
                rule = {
                    field: str(raw_rule.get(field) or "").replace("\x00", "").strip()[:200]
                    for field in ("supplier_id", "supplier_site", "tax_code")
                }
                missing = [field for field, value in rule.items() if not value]
                if missing:
                    errors.append(
                        {
                            "field": f"supplier_rules.{index}",
                            "code": "required",
                            "message": f"{', '.join(missing)} required",
                        }
                    )
                    continue
                supplier_key = rule["supplier_id"].casefold()
                if supplier_key in seen_suppliers:
                    errors.append(
                        {
                            "field": f"supplier_rules.{index}.supplier_id",
                            "code": "duplicate",
                            "message": "supplier_id must be unique",
                        }
                    )
                    continue
                seen_suppliers.add(supplier_key)
                supplier_rules.append(rule)

        include_upc_in_export = changes.get("include_upc_in_export", False)
        if not isinstance(include_upc_in_export, bool):
            errors.append(
                {
                    "field": "include_upc_in_export",
                    "code": "invalid",
                    "message": "must be true or false",
                }
            )

        raw_policy = changes.get("target_cost_policy", {})
        if not isinstance(raw_policy, Mapping):
            raw_policy = {}
            errors.append(
                {
                    "field": "target_cost_policy",
                    "code": "invalid",
                    "message": "must be an object",
                }
            )
        mode = str(raw_policy.get("mode") or "")
        if mode not in COST_POLICY_MODES:
            errors.append(
                {
                    "field": "target_cost_policy.mode",
                    "code": "invalid",
                    "message": "must be invoice_only",
                }
            )
        try:
            tolerance = _decimal(raw_policy.get("maximum_absolute_difference_aed"))
        except (InvalidOperation, ValueError, TypeError):
            tolerance = None
        if tolerance is None or tolerance < 0 or tolerance > Decimal("100000"):
            errors.append(
                {
                    "field": "target_cost_policy.maximum_absolute_difference_aed",
                    "code": "invalid_number",
                    "message": "must be a non-negative amount",
                }
            )
        if errors:
            raise ValidationFailure(errors)

        now = utc_now()
        with self.db.transaction(immediate=True) as conn:
            row = self._get_settings_row(conn)
            if row["version"] != expected_version:
                raise Conflict(
                    "settings were changed by another user", current_version=row["version"]
                )
            new_version = row["version"] + 1
            conn.execute(
                """
                UPDATE app_settings SET brand_label=?,location=?,location_type=?,
                    supplier_rules_json=?,include_upc_in_export=?,
                    target_cost_policy_json=?,version=?,updated_at=?
                WHERE id=1
                """,
                (
                    brand_label,
                    location,
                    location_type,
                    json_dumps(supplier_rules),
                    int(include_upc_in_export),
                    json_dumps(
                        {
                            "mode": mode,
                            "maximum_absolute_difference_aed": _number(tolerance),
                        }
                    ),
                    new_version,
                    now,
                ),
            )
            return self._settings_from_row(self._get_settings_row(conn))

    def _target_defaults(self, conn: Any, supplier_id: str | None) -> dict[str, Any]:
        settings = self._settings_from_row(self._get_settings_row(conn))
        supplier_rule = next(
            (
                rule
                for rule in settings["supplier_rules"]
                if supplier_id
                and str(rule.get("supplier_id") or "").casefold() == supplier_id.casefold()
            ),
            {},
        )
        return {
            "location": settings["location"] or None,
            "location_type": settings["location_type"] or None,
            "supplier_site": supplier_rule.get("supplier_site") or None,
            "tax_code": supplier_rule.get("tax_code") or None,
            "target_cost_policy": settings["target_cost_policy"],
        }

    @property
    def supported_formats(self) -> list[str]:
        return sorted(extension.lstrip(".") for extension in DOCUMENT_EXTENSIONS)

    def start(self) -> int:
        recovered = self.db.recover_interrupted_jobs()
        if self.settings.workers <= 0:
            return recovered
        self._stop.clear()
        with self._worker_lock:
            self._worker_state["started_at"] = utc_now()
        self.ensure_workers()
        self._wake.set()
        return recovered

    def stop(self) -> None:
        self._stop.set()
        self._wake.set()
        for thread in self._threads:
            thread.join(timeout=5)
        self._threads.clear()

    def _spawn_worker(self, index: int) -> threading.Thread:
        thread = threading.Thread(
            target=self._worker_main,
            name=f"invoice-worker-{index + 1}",
            daemon=True,
        )
        thread.start()
        return thread

    def ensure_workers(self) -> int:
        """Replace any worker thread that has exited; return the live count.

        The pool must never be silently smaller than configured: a dead worker
        leaves uploads queued forever with no error. Called at start and from
        every health probe, so a dead pool heals on the next poll.
        """

        if self._stop.is_set() or self.settings.workers <= 0:
            return sum(1 for thread in self._threads if thread.is_alive())
        with self._worker_lock:
            alive: list[threading.Thread] = []
            for thread in self._threads:
                if thread.is_alive():
                    alive.append(thread)
                else:
                    self._worker_state["respawns"] += 1
                    logger.warning("worker %s was not alive; respawning", thread.name)
            self._threads = alive
            index = len(self._threads)
            while len(self._threads) < self.settings.workers:
                self._threads.append(self._spawn_worker(index))
                index += 1
            return len(self._threads)

    def worker_status(self) -> dict[str, Any]:
        """Liveness of the extraction pool, for /api/health and operators."""

        alive = self.ensure_workers()
        with self.db.connection() as conn:
            counts = {
                row["status"]: row["count"]
                for row in conn.execute(
                    "SELECT status, COUNT(*) AS count FROM invoices "
                    "WHERE status IN ('queued','processing') GROUP BY status"
                ).fetchall()
            }
        queued = int(counts.get("queued", 0))
        processing = int(counts.get("processing", 0))
        with self._worker_lock:
            state = dict(self._worker_state)
        last_activity = max(
            (
                value
                for value in (
                    state["started_at"],
                    state["last_claim_at"],
                    state["last_complete_at"],
                )
                if value
            ),
            default=None,
        )
        idle_seconds = None
        if last_activity:
            idle_seconds = round(
                (datetime.now(UTC) - datetime.fromisoformat(last_activity.replace("Z", "+00:00")))
                .total_seconds(),
                1,
            )
        stalled = bool(
            queued > 0
            and processing == 0
            and idle_seconds is not None
            and idle_seconds > self.settings.worker_stall_seconds
        )
        problems: list[str] = []
        if self.settings.workers > 0 and alive < self.settings.workers:
            problems.append(f"{alive} of {self.settings.workers} extraction workers alive")
        if stalled:
            problems.append(
                f"{queued} queued with no worker activity for {idle_seconds:.0f}s"
            )
        return {
            "configured": self.settings.workers,
            "alive": alive,
            "queued": queued,
            "processing": processing,
            "stalled": stalled,
            "idle_seconds": idle_seconds,
            "problems": problems,
            **state,
        }

    def _record_worker_error(self, stage: str, error: BaseException) -> None:
        message = f"{stage}: {error.__class__.__name__}: {error}"[:500]
        with self._worker_lock:
            self._worker_state["error_count"] += 1
            self._worker_state["last_error"] = message
            self._worker_state["last_error_at"] = utc_now()
        logger.warning("extraction worker error: %s", message)

    def ocr_available(self) -> bool:
        return shutil.which("tesseract") is not None

    def _audit(
        self,
        conn: Any,
        invoice_id: str,
        event_type: str,
        *,
        actor: str,
        version: int,
        from_status: str | None = None,
        to_status: str | None = None,
        details: Mapping[str, Any] | None = None,
        created_at: str | None = None,
    ) -> None:
        conn.execute(
            """
            INSERT INTO audit_events
                (invoice_id,event_type,actor,created_at,from_status,to_status,version,details_json)
            VALUES (?,?,?,?,?,?,?,?)
            """,
            (
                invoice_id,
                event_type,
                actor,
                created_at or utc_now(),
                from_status,
                to_status,
                version,
                json_dumps(dict(details or {})),
            ),
        )

    def _line_from_row(self, row: Mapping[str, Any]) -> dict[str, Any]:
        return {
            "id": row["id"],
            "description": row["description"],
            "quantity": _number(row["quantity"]),
            "unit_price": _number(row["unit_price"]),
            "line_total": _number(row["line_total"]),
            "tax_rate": _number(row["tax_rate"]),
            "uom": row["uom"],
            "upc": row["upc"],
            "catalog_item_id": row["catalog_item_id"],
            "rms_item_id": row["rms_item_id"],
            "rms_parent_item": row["rms_item_id"],
            "rms_upc": row["rms_upc"],
            "rms_unit_cost": _number(row["rms_unit_cost"]),
            "rms_po_number": row["rms_po_number"],
            "target_unit_cost": _number(row["target_unit_cost"]),
            "target_cost_source": row["target_cost_source"],
            "target_cost_variance": _number(row["target_cost_variance"]),
            "target_cost_comparison_status": row["target_cost_comparison_status"],
            "target_cost_review_required": bool(row["target_cost_review_required"]),
            "match_status": row["match_status"],
            "confidence": float(row["confidence"] or 0),
            "candidates": _load_json(row["candidates_json"], []),
        }

    def _invoice_from_row(self, conn: Any, row: Mapping[str, Any], *, full: bool) -> dict[str, Any]:
        invoice = {
            "id": row["id"],
            "filename": row["filename"],
            "supplier_id": row["supplier_id"],
            "supplier_name": row["supplier_name"],
            "supplier_site": row["supplier_site"],
            "invoice_number": row["invoice_number"],
            "document_type": row["document_type"],
            "document": row["document"],
            "invoice_date": row["invoice_date"],
            "po_number": row["po_number"],
            "currency": row["currency"],
            "location": row["location"],
            "location_type": row["location_type"],
            "tax_code": row["tax_code"],
            "ref_no_1": row["ref_no_1"],
            "ref_no_2": row["ref_no_2"],
            "ref_no_3": row["ref_no_3"],
            "comment": row["comment"],
            "subtotal": _number(row["subtotal"]),
            "tax_total": _number(row["tax_total"]),
            "total": _number(row["total"]),
            "status": row["status"],
            "extraction_method": row["extraction_method"],
            "created_at": row["created_at"],
            "updated_at": row["updated_at"],
            "error": row["error"],
            "warnings": _load_json(row["warnings_json"], []),
            "version": row["version"],
            "pages": row["pages"],
            "target_cost_reviewed": bool(row["target_cost_reviewed"]),
            "target_cost_reviewed_at": row["target_cost_reviewed_at"],
        }
        counts = conn.execute(
            """
            SELECT COUNT(*) AS line_count,
                   COALESCE(SUM(CASE WHEN rms_item_id IS NOT NULL
                                     AND match_status IN ('auto','confirmed') THEN 1 ELSE 0 END),0)
                       AS matched_lines
            FROM invoice_lines WHERE invoice_id = ?
            """,
            (row["id"],),
        ).fetchone()
        invoice["line_count"] = counts["line_count"]
        invoice["matched_lines"] = counts["matched_lines"]
        line_rows = conn.execute(
            "SELECT * FROM invoice_lines WHERE invoice_id = ? ORDER BY position, id",
            (row["id"],),
        ).fetchall()
        target_total = Decimal("0")
        complete_target = bool(line_rows)
        cost_review_required = False
        currency = row["currency"]
        for line in line_rows:
            try:
                quantity = _decimal(line["quantity"])
                unit_cost = _decimal(line["target_unit_cost"])
            except (InvalidOperation, ValueError, TypeError):
                quantity = unit_cost = None
            if quantity is None or unit_cost is None:
                complete_target = False
            else:
                target_total += _money_round(quantity * unit_cost, currency)
            cost_review_required = cost_review_required or bool(
                line["target_cost_review_required"]
            )
        invoice["target_total_ex_tax"] = _number(target_total) if complete_target else None
        try:
            captured_subtotal = _decimal(row["subtotal"])
        except (InvalidOperation, ValueError, TypeError):
            captured_subtotal = None
        target_variance = (
            target_total - captured_subtotal
            if complete_target and captured_subtotal is not None
            else None
        )
        invoice["target_cost_variance"] = _number(target_variance)
        invoice["target_cost_review_required"] = bool(
            cost_review_required
            or (
                target_variance is not None
                and not _money_equal(target_variance, Decimal("0"), currency)
            )
        )
        if full:
            invoice["lines"] = [self._line_from_row(line) for line in line_rows]
            invoice["raw_text"] = row["raw_text"]
        return invoice

    def _get_row(self, conn: Any, invoice_id: str) -> Any:
        row = conn.execute("SELECT * FROM invoices WHERE id = ?", (invoice_id,)).fetchone()
        if row is None:
            raise NotFound(f"invoice {invoice_id} was not found")
        return row

    def get_invoice(self, invoice_id: str) -> dict[str, Any]:
        with self.db.connection() as conn:
            return self._invoice_from_row(conn, self._get_row(conn, invoice_id), full=True)

    def list_invoices(
        self, *, status: str | None, search: str | None, limit: int, offset: int
    ) -> dict[str, Any]:
        clauses: list[str] = []
        params: list[Any] = []
        if status:
            if status not in INVOICE_STATUSES:
                raise ValidationFailure(
                    [{"field": "status", "code": "invalid", "message": "unknown invoice status"}]
                )
            clauses.append("status = ?")
            params.append(status)
        if search and search.strip():
            needle = f"%{search.strip()}%"
            clauses.append(
                "(filename LIKE ? OR supplier_id LIKE ? OR supplier_name LIKE ? "
                "OR invoice_number LIKE ? OR po_number LIKE ?)"
            )
            params.extend([needle] * 5)
        where = " WHERE " + " AND ".join(clauses) if clauses else ""
        with self.db.connection() as conn:
            total = conn.execute(f"SELECT COUNT(*) FROM invoices{where}", params).fetchone()[0]
            rows = conn.execute(
                f"SELECT * FROM invoices{where} ORDER BY created_at DESC, id DESC LIMIT ? OFFSET ?",
                [*params, limit, offset],
            ).fetchall()
            return {
                "items": [self._invoice_from_row(conn, row, full=False) for row in rows],
                "total": total,
                "limit": limit,
                "offset": offset,
            }

    def stats(self) -> dict[str, int]:
        with self.db.connection() as conn:
            status_counts = {
                row["status"]: row["count"]
                for row in conn.execute(
                    "SELECT status, COUNT(*) AS count FROM invoices GROUP BY status"
                ).fetchall()
            }
            line_counts = conn.execute(
                """
                SELECT COUNT(*) AS total_lines,
                       COALESCE(SUM(CASE WHEN rms_item_id IS NOT NULL
                                         AND match_status IN ('auto','confirmed')
                                         THEN 1 ELSE 0 END),0) AS matched_lines
                FROM invoice_lines
                """
            ).fetchone()
            result = {"total": sum(status_counts.values())}
            for status in INVOICE_STATUSES:
                result[status] = status_counts.get(status, 0)
            result["matched_lines"] = line_counts["matched_lines"]
            result["total_lines"] = line_counts["total_lines"]
            result["catalog_items"] = conn.execute("SELECT COUNT(*) FROM catalog_items").fetchone()[0]
            result["alias_count"] = conn.execute("SELECT COUNT(*) FROM aliases").fetchone()[0]
            return result

    def _validate_document_name(self, filename: str) -> str:
        filename = _clean_filename(filename)
        extension = Path(filename).suffix.casefold()
        if extension not in DOCUMENT_EXTENSIONS:
            supported = ", ".join(self.supported_formats)
            raise UploadRejected(f"unsupported file type; supported formats: {supported}")
        return filename

    def _insert_uploaded_file(
        self,
        temp_path: Path,
        *,
        filename: str,
        digest: str,
        size: int,
        mime_type: str | None,
        supplier_id: str | None,
        supplier_name: str | None,
    ) -> tuple[dict[str, Any], bool]:
        invoice_id = str(uuid.uuid4())
        extension = Path(filename).suffix.casefold()
        final_path = self.settings.source_dir / f"{invoice_id}{extension}"
        now = utc_now()
        moved = False
        try:
            with self.db.transaction(immediate=True) as conn:
                duplicate = conn.execute(
                    "SELECT id, filename, status FROM invoices WHERE sha256 = ?", (digest,)
                ).fetchone()
                if duplicate:
                    return {
                        "id": duplicate["id"],
                        "filename": filename,
                        "status": duplicate["status"],
                    }, True
                os.replace(temp_path, final_path)
                moved = True
                conn.execute(
                    """
                    INSERT INTO invoices
                        (id,filename,stored_path,sha256,mime_type,source_size,supplier_id,
                         supplier_name,status,created_at,updated_at,version,warnings_json)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        invoice_id,
                        filename,
                        str(final_path.resolve()),
                        digest,
                        TRUSTED_MIME_TYPES[extension],
                        size,
                        supplier_id or None,
                        supplier_name or None,
                        "queued",
                        now,
                        now,
                        1,
                        "[]",
                    ),
                )
                self._audit(
                    conn,
                    invoice_id,
                    "uploaded",
                    actor="user",
                    version=1,
                    to_status="queued",
                    details={"filename": filename, "size": size, "sha256": digest},
                    created_at=now,
                )
        except BaseException:
            if moved:
                final_path.unlink(missing_ok=True)
            raise
        finally:
            temp_path.unlink(missing_ok=True)
        self._wake.set()
        return {"id": invoice_id, "filename": filename, "status": "queued"}, False

    async def ingest_upload(
        self,
        upload: Any,
        *,
        supplier_id: str | None = None,
        supplier_name: str | None = None,
    ) -> tuple[dict[str, Any], bool]:
        filename = self._validate_document_name(upload.filename or "upload")
        temp_path = self.settings.source_dir / f".{uuid.uuid4().hex}.part"
        digest = hashlib.sha256()
        size = 0
        try:
            with temp_path.open("xb") as output:
                while True:
                    chunk = await upload.read(1024 * 1024)
                    if not chunk:
                        break
                    size += len(chunk)
                    if size > self.settings.max_file_bytes:
                        raise UploadRejected(
                            f"file exceeds {self.settings.max_file_bytes // (1024 * 1024)} MB limit"
                        )
                    digest.update(chunk)
                    output.write(chunk)
            if size == 0:
                raise UploadRejected("file is empty")
            return self._insert_uploaded_file(
                temp_path,
                filename=filename,
                digest=digest.hexdigest(),
                size=size,
                mime_type=getattr(upload, "content_type", None),
                supplier_id=supplier_id,
                supplier_name=supplier_name,
            )
        finally:
            temp_path.unlink(missing_ok=True)

    def ingest_bytes(
        self,
        filename: str,
        content: bytes,
        *,
        supplier_id: str | None = None,
        supplier_name: str | None = None,
        mime_type: str | None = None,
    ) -> tuple[dict[str, Any], bool]:
        filename = self._validate_document_name(filename)
        if not content:
            raise UploadRejected("file is empty")
        if len(content) > self.settings.max_file_bytes:
            raise UploadRejected(f"file exceeds {self.settings.max_file_bytes // (1024 * 1024)} MB limit")
        temp_path = self.settings.source_dir / f".{uuid.uuid4().hex}.part"
        try:
            temp_path.write_bytes(content)
            return self._insert_uploaded_file(
                temp_path,
                filename=filename,
                digest=hashlib.sha256(content).hexdigest(),
                size=len(content),
                mime_type=mime_type,
                supplier_id=supplier_id,
                supplier_name=supplier_name,
            )
        finally:
            temp_path.unlink(missing_ok=True)

    def source_for(self, invoice_id: str) -> tuple[Path, str, str | None]:
        with self.db.connection() as conn:
            row = self._get_row(conn, invoice_id)
            path = Path(row["stored_path"])
            if not path.is_file():
                raise NotFound("stored source file is missing")
            return path, row["filename"], row["mime_type"]

    def _worker_main(self) -> None:
        try:
            self._worker_loop()
        except BaseException as error:  # pragma: no cover - last line of defence
            self._record_worker_error("loop", error)
            raise
        finally:
            with self._worker_lock:
                self._worker_state["thread_exits"] += 1
            if not self._stop.is_set():
                logger.warning("extraction worker %s exited", threading.current_thread().name)

    def _worker_loop(self) -> None:
        """Claim and process jobs until stopped.

        Every database step can raise ``sqlite3.OperationalError: database is
        locked`` when a long write (a 100k-row catalog import) outlives the busy
        timeout. That must never end the thread: log it, back off, retry.
        """

        backoff = self.settings.poll_seconds
        while not self._stop.is_set():
            self._retry_orphaned_jobs()
            try:
                invoice_id = self._claim_job()
            except Exception as error:
                self._record_worker_error("claim", error)
                self._wake.wait(backoff)
                self._wake.clear()
                backoff = min(backoff * 2, 10.0)
                continue
            backoff = self.settings.poll_seconds
            if invoice_id is None:
                self._wake.wait(self.settings.poll_seconds)
                self._wake.clear()
                continue
            with self._worker_lock:
                self._worker_state["last_claim_at"] = utc_now()
            try:
                self._process_job(invoice_id)
            except Exception as error:
                # _process_job handles extraction failures itself; reaching here
                # means recording the failure hit the database lock. Keep the id
                # so the job is not left in 'processing' until the next restart.
                self._record_worker_error("process", error)
                self._orphaned_jobs.add(invoice_id)
            with self._worker_lock:
                self._worker_state["last_complete_at"] = utc_now()

    def _retry_orphaned_jobs(self) -> None:
        if not self._orphaned_jobs:
            return
        for invoice_id in list(self._orphaned_jobs):
            try:
                self._fail_processing(
                    invoice_id,
                    RuntimeError("worker could not record the outcome; requeued"),
                    permanent=False,
                )
            except Exception as error:
                self._record_worker_error("requeue", error)
                return
            self._orphaned_jobs.discard(invoice_id)

    def _claim_job(self) -> str | None:
        now = utc_now()
        with self.db.transaction(immediate=True) as conn:
            row = conn.execute(
                """
                SELECT id, version FROM invoices
                WHERE status = 'queued' AND (next_attempt_at IS NULL OR next_attempt_at <= ?)
                ORDER BY created_at, id LIMIT 1
                """,
                (now,),
            ).fetchone()
            if row is None:
                return None
            new_version = row["version"] + 1
            changed = conn.execute(
                """
                UPDATE invoices SET status='processing', claimed_at=?, updated_at=?,
                    attempts=attempts+1, version=? WHERE id=? AND status='queued'
                """,
                (now, now, new_version, row["id"]),
            ).rowcount
            if not changed:
                return None
            self._audit(
                conn,
                row["id"],
                "processing_started",
                actor="system",
                version=new_version,
                from_status="queued",
                to_status="processing",
            )
            return row["id"]

    def _catalog_for_matching(self, conn: Any, supplier_id: str | None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        if supplier_id:
            catalog_rows = conn.execute(
                """
                SELECT * FROM catalog_items
                WHERE supplier_id IS NULL OR supplier_id = ''
                   OR supplier_id = ?
                ORDER BY rms_item_id,catalog_item_id
                """,
                (supplier_id,),
            ).fetchall()
        else:
            # A missing supplier must never trigger a 114k-row cross-supplier scan.
            catalog_rows = conn.execute(
                """
                SELECT * FROM catalog_items
                WHERE supplier_id IS NULL OR supplier_id = ''
                ORDER BY rms_item_id,catalog_item_id
                """
            ).fetchall()
        catalog = [
            {
                "catalog_item_id": row["catalog_item_id"],
                "rms_item_id": row["rms_item_id"],
                "parent_item": row["parent_item"],
                "upc": row["upc"],
                "description": row["description"],
                "supplier_id": row["supplier_id"],
                "supplier_name": row["supplier_name"],
                "uom": row["uom"],
                "unit_cost": _number(row["unit_cost"]),
                "master_po_number": row["master_po_number"],
            }
            for row in catalog_rows
        ]
        if supplier_id:
            alias_rows = conn.execute(
                "SELECT * FROM aliases WHERE supplier_scope = ?",
                (supplier_id,),
            ).fetchall()
        else:
            alias_rows = []
        aliases = [
            {
                "supplier_id": row["supplier_scope"],
                "description": row["original_description"],
                "normalized_description": row["normalized_description"],
                "uom": row["uom_scope"] or None,
                "catalog_item_id": row["catalog_item_id"],
                "rms_item_id": row["rms_item_id"],
            }
            for row in alias_rows
        ]
        return catalog, aliases

    def _prepare_lines(
        self,
        lines: Sequence[Mapping[str, Any]] | None,
        *,
        target_cost_policy: Mapping[str, Any] | None = None,
        currency: str | None = None,
    ) -> list[dict[str, Any]]:
        prepared: list[dict[str, Any]] = []
        ids: set[str] = set()
        policy = dict(
            target_cost_policy
            or {
                "mode": "invoice_only",
                "maximum_absolute_difference_aed": 10,
            }
        )
        try:
            tolerance = _decimal(policy.get("maximum_absolute_difference_aed"))
        except (InvalidOperation, ValueError, TypeError):
            tolerance = Decimal("10")
        tolerance = tolerance if tolerance is not None else Decimal("10")
        for position, source in enumerate(lines or []):
            line_id = str(source.get("id") or f"line-{position + 1}")[:100]
            if line_id in ids:
                line_id = f"{line_id}-{position + 1}"
            ids.add(line_id)
            status = str(source.get("match_status") or "unmatched")
            if status not in MATCH_STATUSES:
                status = "unmatched"
            confidence = float(source.get("confidence") or 0)
            confidence = min(100.0, max(0.0, confidence))
            candidates = source.get("candidates")
            if not isinstance(candidates, list):
                candidates = []
            invoice_unit_cost = _decimal(source.get("unit_price"))
            rms_unit_cost = _decimal(source.get("rms_unit_cost"))
            rms_usable = rms_unit_cost is not None and rms_unit_cost > Decimal("0.01")
            target_unit_cost = invoice_unit_cost
            target_source = "invoice"
            comparison_currency_available = str(currency or "").upper() == "AED"
            variance = None
            if not rms_usable:
                comparison_status = "unavailable_rms_cost"
                review_required = False
            elif not comparison_currency_available:
                comparison_status = "unavailable_currency"
                review_required = True
            elif invoice_unit_cost is None:
                comparison_status = "unavailable_invoice_cost"
                review_required = True
            else:
                variance = rms_unit_cost - invoice_unit_cost
                comparison_status = (
                    "above_tolerance" if abs(variance) > tolerance else "within_tolerance"
                )
                review_required = comparison_status == "above_tolerance"
            prepared.append(
                {
                    "id": line_id,
                    "position": position,
                    "description": str(source.get("description") or "")[:2000],
                    "quantity": _decimal_text(source.get("quantity")),
                    "unit_price": _decimal_text(source.get("unit_price")),
                    "line_total": _decimal_text(source.get("line_total")),
                    "tax_rate": _decimal_text(source.get("tax_rate")),
                    "uom": (str(source.get("uom"))[:100] if source.get("uom") is not None else None),
                    "upc": _identifier_text(source.get("upc")),
                    "catalog_item_id": _identifier_text(source.get("catalog_item_id")),
                    "rms_item_id": (
                        str(source.get("rms_item_id"))[:200]
                        if source.get("rms_item_id") is not None
                        else None
                    ),
                    "rms_upc": _normalize_upc(source.get("rms_upc")),
                    "rms_unit_cost": _decimal_text(rms_unit_cost),
                    "rms_po_number": _valid_order_number(source.get("rms_po_number")),
                    "target_unit_cost": _decimal_text(target_unit_cost),
                    "target_cost_source": target_source,
                    "target_cost_variance": _decimal_text(variance),
                    "target_cost_comparison_status": comparison_status,
                    "target_cost_review_required": review_required,
                    "match_status": status,
                    "confidence": confidence,
                    "candidates": candidates[:10],
                }
            )
        return prepared

    def _replace_lines(self, conn: Any, invoice_id: str, lines: Sequence[Mapping[str, Any]]) -> None:
        conn.execute("DELETE FROM invoice_lines WHERE invoice_id = ?", (invoice_id,))
        conn.executemany(
            """
            INSERT INTO invoice_lines
                (invoice_id,id,position,description,quantity,unit_price,line_total,tax_rate,
                 uom,upc,catalog_item_id,rms_item_id,rms_upc,rms_unit_cost,rms_po_number,
                 target_unit_cost,target_cost_source,target_cost_variance,
                 target_cost_comparison_status,target_cost_review_required,
                 match_status,confidence,candidates_json)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            [
                (
                    invoice_id,
                    line["id"],
                    line["position"],
                    line["description"],
                    line["quantity"],
                    line["unit_price"],
                    line["line_total"],
                    line["tax_rate"],
                    line["uom"],
                    line["upc"],
                    line["catalog_item_id"],
                    line["rms_item_id"],
                    line["rms_upc"],
                    line["rms_unit_cost"],
                    line["rms_po_number"],
                    line["target_unit_cost"],
                    line["target_cost_source"],
                    line["target_cost_variance"],
                    line["target_cost_comparison_status"],
                    int(line["target_cost_review_required"]),
                    line["match_status"],
                    line["confidence"],
                    json_dumps(line["candidates"]),
                )
                for line in lines
            ],
        )

    def _process_job(self, invoice_id: str) -> None:
        try:
            with self.db.connection() as conn:
                row = self._get_row(conn, invoice_id)
                source_path = row["stored_path"]
                filename = row["filename"]
                original_supplier_id = row["supplier_id"]
                original_supplier_name = row["supplier_name"]
            from . import extraction, matching

            limits = extraction.ExtractionLimits(
                max_file_bytes=self.settings.max_file_bytes,
                max_pages=self.settings.max_pages,
            )
            extracted = extraction.extract_document(source_path, filename, limits=limits)
            supplier_id = original_supplier_id or extracted.get("supplier_id")
            supplier_name = original_supplier_name or extracted.get("supplier_name")
            with self.db.connection() as conn:
                catalog, aliases = self._catalog_for_matching(conn, supplier_id)
                defaults = self._target_defaults(conn, supplier_id)
            enriched = matching.enrich_invoice(
                extracted,
                catalog,
                aliases,
                supplier_id=supplier_id,
            )
            enriched["supplier_id"] = supplier_id
            enriched["supplier_name"] = supplier_name
            enriched["supplier_site"] = (
                extracted.get("supplier_site") or defaults.get("supplier_site")
            )
            enriched["location"] = defaults.get("location")
            enriched["location_type"] = defaults.get("location_type")
            enriched["tax_code"] = defaults.get("tax_code")
            enriched["target_cost_policy"] = defaults["target_cost_policy"]
            self._complete_processing(invoice_id, enriched)
        except BaseException as error:
            permanent = error.__class__.__name__ in {
                "UnsupportedDocumentError",
                "ExtractionLimitError",
            }
            self._fail_processing(invoice_id, error, permanent=permanent)

    def _complete_processing(self, invoice_id: str, extracted: Mapping[str, Any]) -> None:
        now = utc_now()
        try:
            lines = self._prepare_lines(
                extracted.get("lines") or [],
                target_cost_policy=extracted.get("target_cost_policy"),
                currency=str(extracted.get("currency") or "") or None,
            )
            monetary = {
                field: _decimal_text(extracted.get(field)) for field in ("subtotal", "tax_total", "total")
            }
        except (InvalidOperation, ValueError, TypeError) as error:
            self._fail_processing(invoice_id, error, permanent=True)
            return
        with self.db.transaction(immediate=True) as conn:
            row = self._get_row(conn, invoice_id)
            if row["status"] != "processing":
                return
            new_version = row["version"] + 1
            warnings = extracted.get("warnings")
            if not isinstance(warnings, list):
                warnings = []
            conn.execute(
                """
                UPDATE invoices SET
                    supplier_id=?, supplier_name=?, supplier_site=?, invoice_number=?,
                    document_type=?,document=?,invoice_date=?,po_number=?,currency=?,
                    location=?,location_type=?,tax_code=?,subtotal=?,tax_total=?,total=?,
                    target_cost_reviewed=0,target_cost_reviewed_at=NULL,
                    status='needs_review', extraction_method=?, updated_at=?, error=NULL,
                    warnings_json=?, version=?, pages=?, raw_text=?, claimed_at=NULL,
                    next_attempt_at=NULL
                WHERE id=? AND status='processing'
                """,
                (
                    extracted.get("supplier_id"),
                    extracted.get("supplier_name"),
                    extracted.get("supplier_site"),
                    extracted.get("invoice_number"),
                    (
                        extracted.get("document_type")
                        if extracted.get("document_type") in DOCUMENT_TYPES
                        else "unknown"
                    ),
                    extracted.get("document") or extracted.get("invoice_number"),
                    extracted.get("invoice_date"),
                    _valid_order_number(extracted.get("po_number")),
                    str(extracted.get("currency") or "").upper() or None,
                    extracted.get("location"),
                    extracted.get("location_type"),
                    extracted.get("tax_code"),
                    monetary["subtotal"],
                    monetary["tax_total"],
                    monetary["total"],
                    extracted.get("extraction_method"),
                    now,
                    json_dumps([str(item)[:1000] for item in warnings[:100]]),
                    new_version,
                    int(extracted.get("pages") or 0),
                    str(extracted.get("raw_text") or ""),
                    invoice_id,
                ),
            )
            self._replace_lines(conn, invoice_id, lines)
            self._audit(
                conn,
                invoice_id,
                "extraction_completed",
                actor="system",
                version=new_version,
                from_status="processing",
                to_status="needs_review",
                details={"method": extracted.get("extraction_method"), "line_count": len(lines)},
                created_at=now,
            )

    def _fail_processing(self, invoice_id: str, error: BaseException, *, permanent: bool) -> None:
        now = utc_now()
        message = f"{error.__class__.__name__}: {error}"[:2000]
        with self.db.transaction(immediate=True) as conn:
            row = conn.execute("SELECT * FROM invoices WHERE id = ?", (invoice_id,)).fetchone()
            if row is None or row["status"] != "processing":
                return
            attempts = row["attempts"]
            retry = not permanent and attempts < self.settings.max_attempts
            status = "queued" if retry else "failed"
            next_attempt = None
            if retry:
                delay = min(60, 2 ** max(0, attempts - 1))
                next_attempt = (datetime.now(UTC) + timedelta(seconds=delay)).isoformat().replace(
                    "+00:00", "Z"
                )
            new_version = row["version"] + 1
            conn.execute(
                """
                UPDATE invoices SET status=?, error=?, updated_at=?, next_attempt_at=?,
                    claimed_at=NULL, version=? WHERE id=? AND status='processing'
                """,
                (status, message, now, next_attempt, new_version, invoice_id),
            )
            self._audit(
                conn,
                invoice_id,
                "processing_retry_scheduled" if retry else "processing_failed",
                actor="system",
                version=new_version,
                from_status="processing",
                to_status=status,
                details={"error": message, "attempt": attempts, "retry_at": next_attempt},
                created_at=now,
            )
        if retry:
            self._wake.set()

    def retry(self, invoice_id: str) -> dict[str, Any]:
        now = utc_now()
        with self.db.transaction(immediate=True) as conn:
            row = self._get_row(conn, invoice_id)
            if row["status"] == "processing":
                raise Conflict("invoice is currently processing", current_version=row["version"])
            if row["status"] not in {"failed", "queued"}:
                raise Conflict("only failed invoices can be retried", current_version=row["version"])
            if row["status"] == "queued":
                return self._invoice_from_row(conn, row, full=True)
            new_version = row["version"] + 1
            conn.execute(
                """
                UPDATE invoices SET status='queued', error=NULL, attempts=0, next_attempt_at=NULL,
                    claimed_at=NULL, updated_at=?, version=? WHERE id=?
                """,
                (now, new_version, invoice_id),
            )
            self._audit(
                conn,
                invoice_id,
                "retry_requested",
                actor="user",
                version=new_version,
                from_status="failed",
                to_status="queued",
                created_at=now,
            )
            result = self._invoice_from_row(conn, self._get_row(conn, invoice_id), full=True)
        self._wake.set()
        return result

    def rematch(self, invoice_id: str, expected_version: int) -> dict[str, Any]:
        from . import matching

        now = utc_now()
        with self.db.transaction(immediate=True) as conn:
            row = self._get_row(conn, invoice_id)
            if row["version"] != expected_version:
                raise Conflict(
                    "invoice was changed by another user", current_version=row["version"]
                )
            if row["status"] in {"queued", "processing"}:
                raise Conflict(
                    "invoice cannot be rematched while processing",
                    current_version=row["version"],
                )
            invoice = self._invoice_from_row(conn, row, full=True)
            catalog, aliases = self._catalog_for_matching(conn, invoice.get("supplier_id"))
            enriched_lines = matching.match_lines(
                invoice.get("lines") or [],
                catalog,
                aliases,
                invoice.get("supplier_id"),
            )
            policy = self._target_defaults(
                conn, invoice.get("supplier_id")
            )["target_cost_policy"]
            prepared = self._prepare_lines(
                enriched_lines,
                target_cost_policy=policy,
                currency=invoice.get("currency"),
            )
            self._replace_lines(conn, invoice_id, prepared)
            new_version = row["version"] + 1
            conn.execute(
                """
                UPDATE invoices SET status='needs_review',updated_at=?,version=?,
                    target_cost_reviewed=0,target_cost_reviewed_at=NULL
                WHERE id=?
                """,
                (now, new_version, invoice_id),
            )
            self._audit(
                conn,
                invoice_id,
                "invoice_rematched",
                actor="user",
                version=new_version,
                from_status=row["status"],
                to_status="needs_review",
                details={"line_count": len(prepared)},
                created_at=now,
            )
            return self._invoice_from_row(
                conn, self._get_row(conn, invoice_id), full=True
            )

    def _learn_aliases(
        self,
        conn: Any,
        *,
        invoice_id: str,
        supplier_id: str | None,
        lines: Sequence[Mapping[str, Any]],
        now: str,
    ) -> list[dict[str, str]]:
        if not supplier_id:
            return []
        learned: list[dict[str, str]] = []
        for line in lines:
            if line.get("match_status") != "confirmed" or not line.get("rms_item_id"):
                continue
            catalog_item_id = line.get("catalog_item_id")
            if not catalog_item_id:
                matches = conn.execute(
                    """
                    SELECT catalog_item_id FROM catalog_items
                    WHERE rms_item_id=? AND (supplier_id IS NULL OR supplier_id=''
                        OR lower(trim(supplier_id))=lower(trim(?)))
                    """,
                    (line["rms_item_id"], supplier_id),
                ).fetchall()
                if len(matches) != 1:
                    continue
                catalog_item_id = matches[0]["catalog_item_id"]
            normalized = normalize_description(str(line.get("description") or ""))
            if not normalized:
                continue
            uom_scope = normalize_description(str(line.get("uom") or ""))
            conn.execute(
                """
                INSERT INTO aliases
                    (supplier_scope,normalized_description,uom_scope,original_description,
                     catalog_item_id,rms_item_id,created_from_invoice_id,created_at,updated_at)
                VALUES (?,?,?,?,?,?,?,?,?)
                ON CONFLICT(supplier_scope,normalized_description,uom_scope) DO UPDATE SET
                    original_description=excluded.original_description,
                    catalog_item_id=excluded.catalog_item_id,
                    rms_item_id=excluded.rms_item_id,
                    created_from_invoice_id=excluded.created_from_invoice_id,
                    updated_at=excluded.updated_at
                """,
                (
                    supplier_id,
                    normalized,
                    uom_scope,
                    str(line.get("description") or ""),
                    catalog_item_id,
                    line["rms_item_id"],
                    invoice_id,
                    now,
                    now,
                ),
            )
            learned.append(
                {
                    "description": str(line.get("description") or ""),
                    "uom": str(line.get("uom") or ""),
                    "catalog_item_id": str(catalog_item_id),
                    "rms_item_id": str(line["rms_item_id"]),
                }
            )
        return learned

    def update_invoice(
        self, invoice_id: str, expected_version: int, changes: Mapping[str, Any]
    ) -> dict[str, Any]:
        header_fields = {
            "supplier_id",
            "supplier_name",
            "supplier_site",
            "invoice_number",
            "document_type",
            "document",
            "invoice_date",
            "po_number",
            "currency",
            "location",
            "location_type",
            "tax_code",
            "ref_no_1",
            "ref_no_2",
            "ref_no_3",
            "comment",
            "subtotal",
            "tax_total",
            "total",
        }
        now = utc_now()
        with self.db.transaction(immediate=True) as conn:
            row = self._get_row(conn, invoice_id)
            if row["version"] != expected_version:
                raise Conflict("invoice was changed by another user", current_version=row["version"])
            if row["status"] in {"queued", "processing"}:
                raise Conflict("invoice cannot be edited while processing", current_version=row["version"])
            updates: dict[str, Any] = {}
            changed_fields: list[str] = []
            for field in header_fields:
                if field not in changes:
                    continue
                value = changes[field]
                if field in {"subtotal", "tax_total", "total"}:
                    try:
                        value = _decimal_text(value)
                    except (InvalidOperation, ValueError, TypeError):
                        raise ValidationFailure(
                            [{"field": field, "code": "invalid_number", "message": "must be a valid number"}]
                        ) from None
                elif isinstance(value, str):
                    value = value.strip() or None
                    if field == "currency" and value:
                        value = value.upper()
                if field == "document_type" and value not in DOCUMENT_TYPES:
                    raise ValidationFailure(
                        [
                            {
                                "field": field,
                                "code": "invalid",
                                "message": "must be invoice, credit_note, purchase_order, delivery_note, or unknown",
                            }
                        ]
                    )
                if field == "po_number" and value and not _valid_order_number(value):
                    raise ValidationFailure(
                        [
                            {
                                "field": field,
                                "code": "invalid_order_number",
                                "message": "PO Box and non-order master fields cannot be used as an order number",
                            }
                        ]
                    )
                updates[field] = value
                if value != row[field]:
                    changed_fields.append(field)

            lines_changed = "lines" in changes
            prepared_lines: list[dict[str, Any]] | None = None
            effective_supplier_id = updates.get("supplier_id", row["supplier_id"])
            if lines_changed:
                if not isinstance(changes["lines"], list):
                    raise ValidationFailure(
                        [{"field": "lines", "code": "invalid", "message": "must be a list"}]
                    )
                resolved_lines: list[dict[str, Any]] = []
                for index, source_line in enumerate(changes["lines"]):
                    source = dict(source_line)
                    selected = None
                    catalog_item_id = _identifier_text(source.get("catalog_item_id"))
                    rms_item_id = _identifier_text(source.get("rms_item_id"))
                    if catalog_item_id:
                        selected = conn.execute(
                            "SELECT * FROM catalog_items WHERE catalog_item_id=?",
                            (catalog_item_id,),
                        ).fetchone()
                    elif rms_item_id:
                        matches = conn.execute(
                            """
                            SELECT * FROM catalog_items WHERE rms_item_id=?
                              AND (supplier_id IS NULL OR supplier_id=''
                                   OR supplier_id=?)
                            ORDER BY catalog_item_id
                            """,
                            (rms_item_id, effective_supplier_id or ""),
                        ).fetchall()
                        if len(matches) == 1:
                            selected = matches[0]
                        elif len(matches) > 1:
                            raise ValidationFailure(
                                [
                                    {
                                        "field": f"lines.{index}.catalog_item_id",
                                        "code": "ambiguous_rms_item",
                                        "message": "select the exact supplier/item/UPC catalog row",
                                    }
                                ]
                            )
                        elif conn.execute(
                            "SELECT 1 FROM catalog_items WHERE rms_item_id=? LIMIT 1",
                            (rms_item_id,),
                        ).fetchone():
                            raise ValidationFailure(
                                [
                                    {
                                        "field": f"lines.{index}.catalog_item_id",
                                        "code": "supplier_item_mismatch",
                                        "message": "RMS item belongs to another supplier scope",
                                    }
                                ]
                            )
                    if selected is not None:
                        item_supplier = selected["supplier_id"]
                        if item_supplier and (
                            not effective_supplier_id
                            or item_supplier.casefold() != str(effective_supplier_id).casefold()
                        ):
                            raise ValidationFailure(
                                [
                                    {
                                        "field": f"lines.{index}.catalog_item_id",
                                        "code": "supplier_item_mismatch",
                                        "message": "catalog row belongs to another supplier scope",
                                    }
                                ]
                            )
                        source.update(
                            {
                                "catalog_item_id": selected["catalog_item_id"],
                                "rms_item_id": selected["rms_item_id"],
                                "rms_parent_item": selected["parent_item"],
                                "rms_upc": selected["upc"],
                                "rms_unit_cost": selected["unit_cost"],
                                "rms_po_number": selected["master_po_number"],
                            }
                        )
                    elif rms_item_id or catalog_item_id:
                        raise ValidationFailure(
                            [
                                {
                                    "field": f"lines.{index}.catalog_item_id",
                                    "code": "unknown_rms_item",
                                    "message": "catalog selection was not found",
                                }
                            ]
                        )
                    resolved_lines.append(source)
                try:
                    cost_policy = self._target_defaults(
                        conn, effective_supplier_id
                    )["target_cost_policy"]
                    prepared_lines = self._prepare_lines(
                        resolved_lines,
                        target_cost_policy=cost_policy,
                        currency=updates.get("currency", row["currency"]),
                    )
                except (InvalidOperation, ValueError, TypeError):
                    raise ValidationFailure(
                        [
                            {
                                "field": "lines",
                                "code": "invalid_number",
                                "message": "line numeric values must be valid numbers",
                            }
                        ]
                    ) from None
                changed_fields.append("lines")

            if not lines_changed and "currency" in updates:
                existing_line_rows = conn.execute(
                    "SELECT * FROM invoice_lines WHERE invoice_id=? ORDER BY position,id",
                    (invoice_id,),
                ).fetchall()
                cost_policy = self._target_defaults(
                    conn, effective_supplier_id
                )["target_cost_policy"]
                existing_lines: list[dict[str, Any]] = []
                for line_row in existing_line_rows:
                    source_line = self._line_from_row(line_row)
                    for numeric_field in (
                        "quantity",
                        "unit_price",
                        "line_total",
                        "tax_rate",
                        "rms_unit_cost",
                    ):
                        source_line[numeric_field] = line_row[numeric_field]
                    existing_lines.append(source_line)
                prepared_lines = self._prepare_lines(
                    existing_lines,
                    target_cost_policy=cost_policy,
                    currency=updates["currency"],
                )

            if "invoice_number" in updates and "document" not in updates:
                if not row["document"] or row["document"] == row["invoice_number"]:
                    updates["document"] = updates["invoice_number"]
                    changed_fields.append("document")

            cost_relevant_change = lines_changed or bool(
                {"subtotal", "supplier_id", "currency"}.intersection(changed_fields)
            )

            new_status = "needs_review"
            new_version = row["version"] + 1
            assignments = [f"{field} = ?" for field in updates]
            values = list(updates.values())
            assignments.extend(["status = ?", "updated_at = ?", "version = ?", "error = NULL"])
            values.extend([new_status, now, new_version, invoice_id])
            if cost_relevant_change:
                assignments.extend(
                    ["target_cost_reviewed = 0", "target_cost_reviewed_at = NULL"]
                )
            conn.execute(f"UPDATE invoices SET {', '.join(assignments)} WHERE id = ?", values)
            if prepared_lines is not None:
                self._replace_lines(conn, invoice_id, prepared_lines)
            supplier_id = effective_supplier_id
            alias_lines = prepared_lines or []
            learned = self._learn_aliases(
                conn,
                invoice_id=invoice_id,
                supplier_id=supplier_id,
                lines=alias_lines,
                now=now,
            )
            self._audit(
                conn,
                invoice_id,
                "invoice_edited",
                actor="user",
                version=new_version,
                from_status=row["status"],
                to_status=new_status,
                details={"changed_fields": sorted(set(changed_fields)), "aliases_learned": learned},
                created_at=now,
            )
            return self._invoice_from_row(conn, self._get_row(conn, invoice_id), full=True)

    def _validation_errors(
        self,
        conn: Any,
        invoice: Mapping[str, Any],
        *,
        acknowledge_target_cost_variance: bool = False,
    ) -> list[dict[str, str]]:
        errors: list[dict[str, str]] = []

        def error(field: str, code: str, message: str) -> None:
            errors.append({"field": field, "code": code, "message": message})

        for field, label in (
            ("supplier_id", "supplier ID"),
            ("supplier_name", "supplier name"),
            ("invoice_number", "invoice number"),
            ("invoice_date", "invoice date"),
            ("currency", "currency"),
            ("document", "target document"),
            ("supplier_site", "supplier site"),
            ("location", "location"),
            ("location_type", "location type"),
            ("tax_code", "tax code"),
        ):
            if not invoice.get(field):
                error(field, "required", f"{label} is required")
        if invoice.get("invoice_date"):
            try:
                date.fromisoformat(str(invoice["invoice_date"]))
            except ValueError:
                error("invoice_date", "invalid_date", "must use YYYY-MM-DD")
        if invoice.get("currency") and not re.fullmatch(r"[A-Z]{3}", str(invoice["currency"])):
                error("currency", "invalid_currency", "must be a three-letter ISO currency code")
        if invoice.get("document_type") != "invoice":
            error(
                "document_type",
                "not_invoice",
                "only records explicitly classified as invoices can be approved",
            )
        if invoice.get("po_number") and not _valid_order_number(invoice.get("po_number")):
            error(
                "po_number",
                "invalid_order_number",
                "PO Box and non-order master fields cannot be used as an order number",
            )

        numbers: dict[str, Decimal | None] = {}
        for field in ("subtotal", "tax_total", "total"):
            try:
                numbers[field] = _decimal(invoice.get(field))
            except (InvalidOperation, ValueError, TypeError):
                numbers[field] = None
                error(field, "invalid_number", "must be a valid number")
                continue
            if numbers[field] is None:
                error(field, "required", f"{field.replace('_', ' ')} is required")
            elif numbers[field] < 0:
                error(field, "negative_money", "negative monetary values are not allowed")

        lines = invoice.get("lines") or []
        if not lines:
            error("lines", "required", "at least one invoice line is required")
        line_sum = Decimal("0")
        have_all_line_totals = bool(lines)
        for index, line in enumerate(lines):
            prefix = f"lines.{index}"
            if not str(line.get("description") or "").strip():
                error(f"{prefix}.description", "required", "description is required")
            try:
                quantity = _decimal(line.get("quantity"))
            except (InvalidOperation, ValueError, TypeError):
                quantity = None
                error(f"{prefix}.quantity", "invalid_number", "quantity must be a valid number")
            if quantity is None:
                error(f"{prefix}.quantity", "required", "quantity is required")
            elif quantity <= 0:
                error(f"{prefix}.quantity", "nonpositive_quantity", "quantity must be greater than zero")
            line_numbers: dict[str, Decimal | None] = {}
            for field in ("unit_price", "line_total"):
                try:
                    value = _decimal(line.get(field))
                except (InvalidOperation, ValueError, TypeError):
                    value = None
                    error(f"{prefix}.{field}", "invalid_number", "must be a valid number")
                line_numbers[field] = value
                if value is None:
                    error(f"{prefix}.{field}", "required", f"{field.replace('_', ' ')} is required")
                    if field == "line_total":
                        have_all_line_totals = False
                elif value < 0:
                    error(f"{prefix}.{field}", "negative_money", "negative monetary values are not allowed")
            if (
                quantity is not None
                and line_numbers.get("unit_price") is not None
                and line_numbers.get("line_total") is not None
                and not _money_equal(
                    quantity * line_numbers["unit_price"],
                    line_numbers["line_total"],
                    invoice.get("currency"),
                )
            ):
                calculated = quantity * line_numbers["unit_price"]
                error(
                    f"{prefix}.line_total",
                    "line_calculation_mismatch",
                    f"quantity times unit price ({calculated}) does not match line total ({line_numbers['line_total']})",
                )
            if quantity is not None and line_numbers.get("unit_price") is not None:
                line_sum += _money_round(
                    quantity * line_numbers["unit_price"], invoice.get("currency")
                )
            elif line_numbers.get("line_total") is not None:
                line_sum += _money_round(line_numbers["line_total"], invoice.get("currency"))
            rms_item_id = line.get("rms_item_id")
            catalog_item_id = line.get("catalog_item_id")
            if (
                not rms_item_id
                or not catalog_item_id
                or line.get("match_status") not in {"auto", "confirmed"}
            ):
                error(f"{prefix}.rms_item_id", "unmapped", "line must be mapped to an RMS item")
            else:
                catalog_item = conn.execute(
                    "SELECT * FROM catalog_items WHERE catalog_item_id = ?",
                    (catalog_item_id,),
                ).fetchone()
                if catalog_item is None:
                    error(f"{prefix}.rms_item_id", "unknown_rms_item", "RMS item is not in the catalog")
                elif str(catalog_item["rms_item_id"]) != str(rms_item_id):
                    error(
                        f"{prefix}.catalog_item_id",
                        "catalog_item_mismatch",
                        "catalog row does not match the selected RMS item",
                    )
                elif catalog_item["supplier_id"] and (
                    not invoice.get("supplier_id")
                    or catalog_item["supplier_id"].casefold()
                    != str(invoice.get("supplier_id")).casefold()
                ):
                    error(
                        f"{prefix}.rms_item_id",
                        "supplier_item_mismatch",
                        "RMS item belongs to a different supplier scope",
                    )
            try:
                target_unit_cost = _decimal(line.get("target_unit_cost"))
            except (InvalidOperation, ValueError, TypeError):
                target_unit_cost = None
            if (
                target_unit_cost is None
                or line_numbers.get("unit_price") is None
                or target_unit_cost != line_numbers["unit_price"]
            ):
                error(
                    f"{prefix}.target_unit_cost",
                    "target_cost_mismatch",
                    "target unit cost must equal invoice net unit cost",
                )

        if have_all_line_totals and numbers.get("subtotal") is not None:
            if not _money_equal(line_sum, numbers["subtotal"], invoice.get("currency")):
                error(
                    "subtotal",
                    "sum_mismatch",
                    f"line totals ({line_sum}) do not match subtotal ({numbers['subtotal']})",
                )

        if invoice.get("target_cost_review_required") and not (
            invoice.get("target_cost_reviewed") or acknowledge_target_cost_variance
        ):
            error(
                "target_cost_reviewed",
                "target_cost_review_required",
                "RMS comparison exceeds tolerance or is unavailable; explicit acknowledgement is required",
            )
        if all(numbers.get(field) is not None for field in ("subtotal", "tax_total", "total")):
            expected_total = numbers["subtotal"] + numbers["tax_total"]
            if not _money_equal(expected_total, numbers["total"], invoice.get("currency")):
                error(
                    "total",
                    "sum_mismatch",
                    f"subtotal plus tax ({expected_total}) does not match total ({numbers['total']})",
                )

        if invoice.get("supplier_id") and invoice.get("invoice_number"):
            duplicate = conn.execute(
                """
                SELECT id FROM invoices
                WHERE id <> ? AND lower(trim(supplier_id)) = lower(trim(?))
                  AND lower(trim(invoice_number)) = lower(trim(?)) AND status <> 'failed'
                LIMIT 1
                """,
                (invoice["id"], invoice["supplier_id"], invoice["invoice_number"]),
            ).fetchone()
            if duplicate:
                error(
                    "invoice_number",
                    "duplicate_supplier_invoice",
                    f"supplier invoice number already exists on {duplicate['id']}",
                )
        return errors

    def approve(
        self,
        invoice_id: str,
        expected_version: int,
        *,
        acknowledge_target_cost_variance: bool = False,
    ) -> dict[str, Any]:
        now = utc_now()
        with self.db.transaction(immediate=True) as conn:
            row = self._get_row(conn, invoice_id)
            if row["version"] != expected_version:
                raise Conflict("invoice was changed by another user", current_version=row["version"])
            if row["status"] not in {"needs_review", "ready"}:
                raise Conflict(
                    "invoice must finish processing before approval", current_version=row["version"]
                )
            invoice = self._invoice_from_row(conn, row, full=True)
            errors = self._validation_errors(
                conn,
                invoice,
                acknowledge_target_cost_variance=acknowledge_target_cost_variance,
            )
            if errors:
                raise ValidationFailure(errors)
            new_version = row["version"] + 1
            reviewed = bool(
                invoice.get("target_cost_reviewed")
                or (
                    invoice.get("target_cost_review_required")
                    and acknowledge_target_cost_variance
                )
            )
            conn.execute(
                """
                UPDATE invoices SET status='ready',updated_at=?,version=?,
                    target_cost_reviewed=?,target_cost_reviewed_at=?
                WHERE id=?
                """,
                (now, new_version, int(reviewed), now if reviewed else None, invoice_id),
            )
            learned = self._learn_aliases(
                conn,
                invoice_id=invoice_id,
                supplier_id=invoice.get("supplier_id"),
                lines=invoice["lines"],
                now=now,
            )
            self._audit(
                conn,
                invoice_id,
                "invoice_approved",
                actor="user",
                version=new_version,
                from_status=row["status"],
                to_status="ready",
                details={
                    "aliases_learned": learned,
                    "target_cost_review_required": invoice.get(
                        "target_cost_review_required", False
                    ),
                    "target_cost_acknowledged": reviewed,
                },
                created_at=now,
            )
            return self._invoice_from_row(conn, self._get_row(conn, invoice_id), full=True)

    def audit(self, invoice_id: str) -> dict[str, Any]:
        with self.db.connection() as conn:
            self._get_row(conn, invoice_id)
            rows = conn.execute(
                "SELECT * FROM audit_events WHERE invoice_id = ? ORDER BY id DESC", (invoice_id,)
            ).fetchall()
            return {
                "items": [
                    {
                        "id": row["id"],
                        "invoice_id": row["invoice_id"],
                        "event_type": row["event_type"],
                        "actor": row["actor"],
                        "created_at": row["created_at"],
                        "from_status": row["from_status"],
                        "to_status": row["to_status"],
                        "version": row["version"],
                        "details": _load_json(row["details_json"], {}),
                    }
                    for row in rows
                ]
            }

    def list_catalog(
        self, *, search: str | None, limit: int, supplier_id: str | None = None
    ) -> dict[str, Any]:
        clauses: list[str] = []
        params: list[Any] = []
        if supplier_id and supplier_id.strip():
            clauses.append("lower(trim(supplier_id)) = lower(trim(?))")
            params.append(supplier_id.strip())
        if search and search.strip():
            clauses.append(
                "(rms_item_id LIKE ? OR upc LIKE ? OR description LIKE ? OR supplier_id LIKE ?)"
            )
            needle = f"%{search.strip()}%"
            params.extend([needle, needle, needle, needle])
        where = " WHERE " + " AND ".join(clauses) if clauses else ""
        with self.db.connection() as conn:
            total = conn.execute(f"SELECT COUNT(*) FROM catalog_items{where}", params).fetchone()[0]
            rows = conn.execute(
                f"""SELECT * FROM catalog_items{where}
                    ORDER BY description,rms_item_id,catalog_item_id LIMIT ?""",
                [*params, limit],
            ).fetchall()
            return {
                "items": [
                    {
                        "catalog_item_id": row["catalog_item_id"],
                        "rms_item_id": row["rms_item_id"],
                        "parent_item": row["parent_item"],
                        "rms_parent_item": row["parent_item"],
                        "upc": row["upc"],
                        "description": row["description"],
                        "supplier_id": row["supplier_id"],
                        "supplier_name": row["supplier_name"],
                        "uom": row["uom"],
                        "unit_cost": _number(row["unit_cost"]),
                        "master_po_number": row["master_po_number"],
                    }
                    for row in rows
                ],
                "total": total,
            }

    @staticmethod
    def _canonical_catalog_field(key: Any) -> str | None:
        normalized = re.sub(
            r"[\s\-]+", "_", str(key or "").strip().casefold()
        ).strip("_")
        aliases = {
            "rms_item_id": {
                "rms_item_id",
                "rms_id",
                "rms_item",
                "item_id",
                "item_parent",
                "parent_item",
                "sku",
            },
            "upc": {"upc", "barcode", "item", "item_code"},
            "description": {
                "description",
                "item_description",
                "item_desc",
                "item_name",
            },
            "supplier_id": {"supplier_id", "supplier", "vendor_id", "vendor"},
            "supplier_name": {"supplier_name", "vendor_name"},
            "uom": {"uom", "standard_uom", "unit", "unit_of_measure"},
            "unit_cost": {"unit_cost", "supplier_unit_cost", "cost", "price"},
            "master_po_number": {
                "po_number",
                "purchase_order",
                "purchase_order_number",
                "order_number",
            },
        }
        return next(
            (target for target, names in aliases.items() if normalized in names), None
        )

    def _catalog_rows(self, filename: str, content: bytes) -> Iterable[dict[str, Any]]:
        extension = Path(filename).suffix.casefold()
        if extension == ".csv":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp1252")
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise UploadRejected("catalog CSV has no header row")
            if len(reader.fieldnames) > 256:
                raise UploadRejected("catalog exceeds the 256 column safety limit")
            fields = {
                self._canonical_catalog_field(name) for name in reader.fieldnames
            }
            if not {"rms_item_id", "description"} <= fields:
                raise UploadRejected(
                    "catalog requires a parent/RMS item column and description column"
                )
            for row_number, row in enumerate(reader, start=2):
                if row_number > 250_001:
                    raise UploadRejected("catalog exceeds the 250,000 row safety limit")
                yield row
            return
        if extension == ".xlsx":
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    members = archive.infolist()
                    if len(members) > 10_000:
                        raise UploadRejected("catalog workbook contains too many archive members")
                    expanded = 0
                    for member in members:
                        expanded += member.file_size
                        if expanded > 1024 * 1024 * 1024:
                            raise UploadRejected("catalog workbook expands beyond the 1 GB safety limit")
                        if member.file_size > 0 and member.file_size / max(1, member.compress_size) > 200:
                            raise UploadRejected("catalog workbook has an unsafe compression ratio")
            except zipfile.BadZipFile as error:
                raise UploadRejected("catalog XLSX is not a valid Office archive") from error
            try:
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            except Exception as error:
                raise UploadRejected(f"invalid XLSX catalog: {error}") from error
            try:
                selected: tuple[Any, tuple[Any, ...]] | None = None
                for sheet in workbook.worksheets:
                    headers = next(
                        sheet.iter_rows(min_row=1, max_row=1, values_only=True), None
                    )
                    if not headers or len(headers) > 256:
                        continue
                    fields = {self._canonical_catalog_field(value) for value in headers}
                    if {"rms_item_id", "description"} <= fields:
                        selected = (sheet, headers)
                        break
                if selected is None:
                    raise UploadRejected(
                        "catalog workbook has no sheet with parent/RMS item and description columns"
                    )
                sheet, headers = selected
                names = [str(value or "") for value in headers]
                for row_number, values in enumerate(
                    sheet.iter_rows(min_row=2, values_only=True), start=2
                ):
                    if row_number > 250_001:
                        raise UploadRejected("catalog exceeds the 250,000 row safety limit")
                    yield {name: value for name, value in zip(names, values)}
            finally:
                workbook.close()
            return
        raise UploadRejected("catalog must be CSV or XLSX")

    def import_catalog(self, filename: str, content: bytes) -> dict[str, Any]:
        filename = _clean_filename(filename)
        if Path(filename).suffix.casefold() not in CATALOG_EXTENSIONS:
            raise UploadRejected("catalog must be CSV or XLSX")
        if not content:
            raise UploadRejected("catalog file is empty")
        if len(content) > self.settings.max_catalog_file_bytes:
            raise UploadRejected(
                f"catalog exceeds {self.settings.max_catalog_file_bytes // (1024 * 1024)} MB limit"
            )

        imported = 0
        skipped = 0
        warnings: list[str] = []
        occurrences: dict[tuple[str, str, str, str], int] = {}
        now = utc_now()
        insert_sql = """
            INSERT INTO catalog_items
                (catalog_item_id,rms_item_id,parent_item,upc,description,
                 normalized_description,supplier_id,supplier_name,uom,unit_cost,
                 master_po_number,source_row,source_name,created_at,updated_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(catalog_item_id) DO UPDATE SET
                rms_item_id=excluded.rms_item_id,
                parent_item=excluded.parent_item,
                upc=excluded.upc,
                description=excluded.description,
                normalized_description=excluded.normalized_description,
                supplier_id=excluded.supplier_id,
                supplier_name=excluded.supplier_name,
                uom=excluded.uom,
                unit_cost=excluded.unit_cost,
                master_po_number=excluded.master_po_number,
                source_row=excluded.source_row,
                source_name=excluded.source_name,
                updated_at=excluded.updated_at
        """
        # Parse the whole file before taking the write lock. Holding BEGIN
        # IMMEDIATE across a multi-minute parse blocks every worker's job claim
        # past the busy timeout; the insert itself takes seconds.
        pending: list[tuple[Any, ...]] = []
        if True:
            for row_number, source in enumerate(self._catalog_rows(filename, content), start=2):
                row: dict[str, Any] = {}
                for key, value in source.items():
                    target = self._canonical_catalog_field(key)
                    if target and target not in row:
                        row[target] = value
                item_id = _identifier_text(row.get("rms_item_id"))
                description = _identifier_text(row.get("description"), limit=2000)
                if not item_id and not description and not any(source.values()):
                    skipped += 1
                    continue
                if not item_id or not description:
                    skipped += 1
                    if len(warnings) < 100:
                        warnings.append(
                            f"row {row_number}: parent/RMS item and description are required"
                        )
                    continue
                supplier_id = _identifier_text(row.get("supplier_id"))
                supplier_name = _identifier_text(row.get("supplier_name"), limit=500)
                upc = _normalize_upc(row.get("upc"))
                uom = _identifier_text(row.get("uom"), limit=100)
                try:
                    unit_cost = _decimal_text(row.get("unit_cost"))
                    if unit_cost is not None and Decimal(unit_cost) < 0:
                        raise InvalidOperation
                except (InvalidOperation, ValueError, TypeError):
                    skipped += 1
                    if len(warnings) < 100:
                        warnings.append(f"row {row_number}: invalid unit_cost")
                    continue
                raw_order = _identifier_text(row.get("master_po_number"))
                master_po_number = _valid_order_number(raw_order)
                if raw_order and not master_po_number and len(warnings) < 100:
                    warnings.append(
                        f"row {row_number}: ignored non-order value in explicit order field"
                    )
                identity = (
                    (supplier_id or "").casefold(),
                    item_id,
                    upc or "",
                    (uom or "").casefold(),
                )
                occurrence = occurrences.get(identity, 0) + 1
                occurrences[identity] = occurrence
                catalog_item_id = _catalog_item_key(
                    supplier_id=supplier_id,
                    rms_item_id=item_id,
                    upc=upc,
                    uom=uom,
                    occurrence=occurrence,
                )
                pending.append(
                    (
                        catalog_item_id,
                        item_id,
                        item_id,
                        upc,
                        description,
                        normalize_description(description),
                        supplier_id,
                        supplier_name,
                        uom,
                        unit_cost,
                        master_po_number,
                        row_number,
                        filename,
                        now,
                        now,
                    )
                )
                imported += 1
        with self.db.transaction(immediate=True) as conn:
            for start in range(0, len(pending), 1000):
                conn.executemany(insert_sql, pending[start : start + 1000])
        self._wake.set()
        return {"imported": imported, "skipped": skipped, "warnings": warnings}

    def import_aliases(self, filename: str, content: bytes) -> dict[str, Any]:
        filename = _clean_filename(filename)
        if Path(filename).suffix.casefold() != ".json":
            raise UploadRejected("aliases must be a JSON file")
        if not content:
            raise UploadRejected("alias file is empty")
        if len(content) > 5 * 1024 * 1024:
            raise UploadRejected("alias file exceeds the 5 MB limit")
        try:
            payload = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as error:
            raise UploadRejected(f"invalid alias JSON: {error}") from error
        records = payload.get("aliases") if isinstance(payload, Mapping) else payload
        if not isinstance(records, list):
            raise UploadRejected("alias JSON must contain an aliases list")
        imported = 0
        skipped = 0
        warnings: list[str] = []
        now = utc_now()
        with self.db.transaction(immediate=True) as conn:
            for index, raw in enumerate(records):
                if not isinstance(raw, Mapping):
                    skipped += 1
                    if len(warnings) < 100:
                        warnings.append(f"alias {index + 1}: must be an object")
                    continue
                supplier_id = _identifier_text(
                    raw.get("supplier_id") or raw.get("supplier_site")
                )
                description = _identifier_text(
                    raw.get("description")
                    or raw.get("alias")
                    or raw.get("invoice_product_name"),
                    limit=2000,
                )
                catalog_item_id = _identifier_text(raw.get("catalog_item_id"))
                rms_item_id = _identifier_text(
                    raw.get("rms_item_id") or raw.get("rms_item")
                )
                uom_scope = normalize_description(str(raw.get("uom") or ""))
                if not supplier_id or not description or not (
                    catalog_item_id or rms_item_id
                ):
                    skipped += 1
                    if len(warnings) < 100:
                        warnings.append(
                            f"alias {index + 1}: supplier, description, and catalog target are required"
                        )
                    continue
                if catalog_item_id:
                    matches = conn.execute(
                        """
                        SELECT * FROM catalog_items WHERE catalog_item_id=?
                          AND lower(trim(supplier_id))=lower(trim(?))
                        """,
                        (catalog_item_id, supplier_id),
                    ).fetchall()
                else:
                    matches = conn.execute(
                        """
                        SELECT * FROM catalog_items WHERE rms_item_id=?
                          AND lower(trim(supplier_id))=lower(trim(?))
                        ORDER BY catalog_item_id
                        """,
                        (rms_item_id, supplier_id),
                    ).fetchall()
                if len(matches) != 1:
                    skipped += 1
                    if len(warnings) < 100:
                        warnings.append(
                            f"alias {index + 1}: catalog target is missing or ambiguous for supplier"
                        )
                    continue
                selected = matches[0]
                conn.execute(
                    """
                    INSERT INTO aliases
                        (supplier_scope,normalized_description,uom_scope,
                         original_description,catalog_item_id,rms_item_id,
                         created_from_invoice_id,created_at,updated_at)
                    VALUES (?,?,?,?,?,?,NULL,?,?)
                    ON CONFLICT(supplier_scope,normalized_description,uom_scope) DO UPDATE SET
                        original_description=excluded.original_description,
                        catalog_item_id=excluded.catalog_item_id,
                        rms_item_id=excluded.rms_item_id,
                        created_from_invoice_id=NULL,
                        updated_at=excluded.updated_at
                    """,
                    (
                        supplier_id,
                        normalize_description(description),
                        uom_scope,
                        description,
                        selected["catalog_item_id"],
                        selected["rms_item_id"],
                        now,
                        now,
                    ),
                )
                imported += 1
        return {"imported": imported, "skipped": skipped, "warnings": warnings}

    def list_aliases(self) -> dict[str, Any]:
        with self.db.connection() as conn:
            rows = conn.execute(
                "SELECT * FROM aliases ORDER BY updated_at DESC, id DESC"
            ).fetchall()
            return {
                "items": [
                    {
                        "id": row["id"],
                        "supplier_id": row["supplier_scope"],
                        "description": row["original_description"],
                        "normalized_description": row["normalized_description"],
                        "uom": row["uom_scope"] or None,
                        "catalog_item_id": row["catalog_item_id"],
                        "rms_item_id": row["rms_item_id"],
                        "created_from_invoice_id": row["created_from_invoice_id"],
                        "created_at": row["created_at"],
                        "updated_at": row["updated_at"],
                    }
                    for row in rows
                ],
                "total": len(rows),
            }

    def create_export(self, invoice_ids: Sequence[str]) -> dict[str, Any]:
        unique_ids = list(dict.fromkeys(str(value) for value in invoice_ids if value))
        if not unique_ids:
            raise ValidationFailure(
                [{"field": "invoice_ids", "code": "required", "message": "select at least one invoice"}]
            )
        if len(unique_ids) > 5000:
            raise ValidationFailure(
                [{"field": "invoice_ids", "code": "too_many", "message": "maximum export is 5,000 invoices"}]
            )
        export_id = str(uuid.uuid4())
        now = utc_now()
        timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
        filename = f"consolidated-invoices-{timestamp}-{export_id[:8]}.xlsx"
        final_path = self.settings.export_dir / filename
        temp_path = self.settings.export_dir / f".{export_id}.part"
        written = False
        try:
            with self.db.transaction(immediate=True) as conn:
                export_settings = self._settings_from_row(self._get_settings_row(conn))
                include_upc_in_export = export_settings["include_upc_in_export"]
                invoices: list[dict[str, Any]] = []
                eligibility_errors: list[dict[str, str]] = []
                for index, invoice_id in enumerate(unique_ids):
                    row = conn.execute("SELECT * FROM invoices WHERE id = ?", (invoice_id,)).fetchone()
                    if row is None:
                        eligibility_errors.append(
                            {
                                "field": f"invoice_ids.{index}",
                                "code": "not_found",
                                "message": f"invoice {invoice_id} was not found",
                            }
                        )
                        continue
                    invoice = self._invoice_from_row(conn, row, full=True)
                    if row["status"] not in {"ready", "exported"}:
                        eligibility_errors.append(
                            {
                                "field": f"invoice_ids.{index}",
                                "code": "not_approved",
                                "message": f"invoice {invoice_id} is not approved",
                            }
                        )
                        continue
                    errors = self._validation_errors(conn, invoice)
                    for error in errors:
                        eligibility_errors.append(
                            {
                                "field": f"invoice_ids.{index}.{error['field']}",
                                "code": error["code"],
                                "message": error["message"],
                            }
                        )
                    invoices.append(invoice)
                if eligibility_errors:
                    raise ValidationFailure(eligibility_errors)
                workbook = build_target_workbook(
                    invoices, include_upc_in_export=include_upc_in_export
                )
                temp_path.write_bytes(workbook)
                os.replace(temp_path, final_path)
                written = True
                digest = hashlib.sha256(workbook).hexdigest()
                conn.execute(
                    """
                    INSERT INTO exports
                        (id,filename,stored_path,created_at,invoice_count,line_count,
                         invoice_ids_json,schema_name,sha256)
                    VALUES (?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        export_id,
                        filename,
                        str(final_path.resolve()),
                        now,
                        len(invoices),
                        sum(len(invoice["lines"]) for invoice in invoices),
                        json_dumps(unique_ids),
                        SCHEMA_NAME,
                        digest,
                    ),
                )
                for invoice in invoices:
                    approved_version = invoice["version"]
                    new_version = approved_version + 1
                    conn.execute(
                        "UPDATE invoices SET status='exported', updated_at=?, version=? WHERE id=?",
                        (now, new_version, invoice["id"]),
                    )
                    conn.execute(
                        "INSERT INTO export_invoices(export_id,invoice_id,invoice_version) VALUES (?,?,?)",
                        (export_id, invoice["id"], approved_version),
                    )
                    self._audit(
                        conn,
                        invoice["id"],
                        "invoice_exported",
                        actor="user",
                        version=new_version,
                        from_status=invoice["status"],
                        to_status="exported",
                        details={
                            "export_id": export_id,
                            "schema": SCHEMA_NAME,
                            "include_upc_in_export": include_upc_in_export,
                        },
                        created_at=now,
                    )
            return {
                "id": export_id,
                "filename": filename,
                "path": final_path,
                "created_at": now,
                "invoice_count": len(unique_ids),
                "sha256": digest,
                "schema_name": SCHEMA_NAME,
            }
        except BaseException:
            temp_path.unlink(missing_ok=True)
            if written:
                final_path.unlink(missing_ok=True)
            raise

    def create_exception_report(self) -> dict[str, Any]:
        """Build a read-only CSV of records that are not target-export eligible."""

        status_reasons = {
            "queued": ("status_queued", "awaiting extraction"),
            "processing": ("status_processing", "extraction is in progress"),
            "needs_review": ("status_needs_review", "manual review and approval required"),
            "failed": ("status_failed", "document processing failed"),
        }
        rows: list[list[str]] = []
        with self.db.connection() as conn:
            invoice_rows = conn.execute(
                "SELECT * FROM invoices ORDER BY created_at,id"
            ).fetchall()
            for row in invoice_rows:
                invoice = self._invoice_from_row(conn, row, full=True)
                reasons: list[str] = []
                status_reason = status_reasons.get(row["status"])
                if status_reason:
                    code, message = status_reason
                    if row["status"] == "failed" and row["error"]:
                        message = f"{message}: {row['error']}"
                    reasons.append(f"[{code}] {message}")

                # Records still being processed do not yet have reviewable
                # target metadata.  Once processing finishes, report every
                # concrete approval/export gate using its stable field/code.
                if row["status"] not in {"queued", "processing", "failed"}:
                    for error in self._validation_errors(conn, invoice):
                        reasons.append(
                            f"[{error['field']}.{error['code']}] {error['message']}"
                        )

                eligible = row["status"] in {"ready", "exported"} and not reasons
                if eligible:
                    continue
                rows.append(
                    [
                        _spreadsheet_safe_csv_text(invoice["id"]),
                        _spreadsheet_safe_csv_text(invoice["filename"]),
                        _spreadsheet_safe_csv_text(invoice.get("document_type") or "unknown"),
                        _spreadsheet_safe_csv_text(invoice["status"]),
                        _spreadsheet_safe_csv_text(" | ".join(reasons)),
                    ]
                )

        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator="\r\n")
        writer.writerow(["Invoice ID", "Filename", "Document Type", "Status", "Reasons"])
        writer.writerows(rows)
        timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
        return {
            "filename": f"invoice-exceptions-{timestamp}.csv",
            "content": output.getvalue().encode("utf-8-sig"),
            "count": len(rows),
        }

    def list_exports(self) -> dict[str, Any]:
        with self.db.connection() as conn:
            rows = conn.execute("SELECT * FROM exports ORDER BY created_at DESC, id DESC").fetchall()
            return {
                "items": [
                    {
                        "id": row["id"],
                        "filename": row["filename"],
                        "created_at": row["created_at"],
                        "invoice_count": row["invoice_count"],
                        "line_count": row["line_count"],
                        "invoice_ids": _load_json(row["invoice_ids_json"], []),
                        "schema_name": row["schema_name"],
                        "sha256": row["sha256"],
                    }
                    for row in rows
                ]
            }

    def seed_demo(self) -> dict[str, Any]:
        catalog = [
            ("RMS-1001", "Fictional Hydrating Cleanser 250ml", "DEMO-SUPPLIER", "EA", "45.00"),
            ("RMS-1002", "Fictional Velvet Lip Colour Ruby", "DEMO-SUPPLIER", "EA", "62.50"),
            ("RMS-1003", "Fictional Daily Face Cream 50ml", "DEMO-SUPPLIER", "EA", "80.00"),
            ("RMS-2001", "Fictional Brightening Serum 30ml", "DEMO-SECOND", "EA", "95.00"),
        ]
        demos = [
            {
                "id": "demo-invoice-001",
                "filename": "FICTIONAL-demo-invoice-001.png",
                "supplier_id": "DEMO-SUPPLIER",
                "supplier_name": "Fictional Beauty Supplies LLC",
                "supplier_site": "DEMO-DUBAI",
                "invoice_number": "DEMO-INV-1001",
                "invoice_date": "2026-09-20",
                "po_number": "DEMO-PO-501",
                "currency": "AED",
                "subtotal": "215.00",
                "tax_total": "10.75",
                "total": "225.75",
                "lines": [
                    {
                        "id": "1",
                        "description": "Fictional Hydrating Cleanser 250ml",
                        "quantity": "2",
                        "unit_price": "45.00",
                        "line_total": "90.00",
                        "tax_rate": "5",
                        "uom": "EA",
                        "rms_item_id": "RMS-1001",
                        "match_status": "auto",
                        "confidence": 100,
                        "candidates": [],
                    },
                    {
                        "id": "2",
                        "description": "Fictional Velvet Lip Colour Ruby",
                        "quantity": "2",
                        "unit_price": "62.50",
                        "line_total": "125.00",
                        "tax_rate": "5",
                        "uom": "EA",
                        "rms_item_id": "RMS-1002",
                        "match_status": "auto",
                        "confidence": 100,
                        "candidates": [],
                    },
                ],
            },
            {
                "id": "demo-invoice-002",
                "filename": "FICTIONAL-demo-invoice-002.png",
                "supplier_id": "DEMO-SECOND",
                "supplier_name": "Fictional Skin Lab FZCO",
                "supplier_site": "DEMO-JEBEL-ALI",
                "invoice_number": "DEMO-INV-2001",
                "invoice_date": "2026-09-21",
                "po_number": "DEMO-PO-502",
                "currency": "AED",
                "subtotal": "190.00",
                "tax_total": "9.50",
                "total": "199.50",
                "lines": [
                    {
                        "id": "1",
                        "description": "Bright Serum Special 30 ml",
                        "quantity": "2",
                        "unit_price": "95.00",
                        "line_total": "190.00",
                        "tax_rate": "5",
                        "uom": "EA",
                        "rms_item_id": None,
                        "match_status": "suggested",
                        "confidence": 84,
                        "candidates": [
                            {
                                "rms_item_id": "RMS-2001",
                                "description": "Fictional Brightening Serum 30ml",
                                "score": 84,
                                "reason": "fuzzy description suggestion; user confirmation required",
                            }
                        ],
                    }
                ],
            },
        ]
        now = utc_now()
        catalog_upserted = 0
        invoices_created = 0
        with self.db.transaction(immediate=True) as conn:
            for item in catalog:
                catalog_item_id = f"demo:{item[0]}"
                conn.execute(
                    """
                    INSERT INTO catalog_items
                        (catalog_item_id,rms_item_id,parent_item,upc,description,
                         normalized_description,supplier_id,supplier_name,uom,unit_cost,
                         master_po_number,source_row,source_name,created_at,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(catalog_item_id) DO UPDATE SET
                        description=excluded.description,supplier_id=excluded.supplier_id,
                        uom=excluded.uom,unit_cost=excluded.unit_cost,updated_at=excluded.updated_at
                    """,
                    (
                        catalog_item_id,
                        item[0],
                        item[0],
                        None,
                        item[1],
                        normalize_description(item[1]),
                        item[2],
                        None,
                        item[3],
                        item[4],
                        None,
                        None,
                        "fictional demo",
                        now,
                        now,
                    ),
                )
                catalog_upserted += 1
            for demo in demos:
                exists = conn.execute("SELECT 1 FROM invoices WHERE id = ?", (demo["id"],)).fetchone()
                if exists:
                    continue
                image = Image.new("RGB", (1240, 1754), "white")
                drawing = ImageDraw.Draw(image)
                text_lines = [
                    "FICTIONAL DEMO INVOICE — NOT A REAL SUPPLIER DOCUMENT",
                    demo["supplier_name"],
                    f"Invoice: {demo['invoice_number']}   Date: {demo['invoice_date']}",
                    f"PO: {demo['po_number']}   Currency: {demo['currency']}",
                    "",
                ]
                text_lines.extend(
                    f"{line['description']} | {line['quantity']} x {line['unit_price']} = {line['line_total']}"
                    for line in demo["lines"]
                )
                text_lines.extend(
                    ["", f"Subtotal {demo['subtotal']}", f"Tax {demo['tax_total']}", f"Total {demo['total']}"]
                )
                drawing.multiline_text((80, 90), "\n".join(text_lines), fill="black", spacing=18)
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                content = buffer.getvalue()
                source_path = self.settings.source_dir / demo["filename"]
                source_path.write_bytes(content)
                raw_text = "\n".join(text_lines)
                conn.execute(
                    """
                    INSERT INTO invoices
                        (id,filename,stored_path,sha256,mime_type,source_size,supplier_id,
                         supplier_name,supplier_site,invoice_number,document_type,document,
                         invoice_date,po_number,currency,location,location_type,tax_code,
                         subtotal,tax_total,total,status,extraction_method,created_at,updated_at,
                         warnings_json,version,pages,raw_text)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        demo["id"],
                        demo["filename"],
                        str(source_path.resolve()),
                        hashlib.sha256(content).hexdigest(),
                        "image/png",
                        len(content),
                        demo["supplier_id"],
                        demo["supplier_name"],
                        demo["supplier_site"],
                        demo["invoice_number"],
                        "invoice",
                        demo["invoice_number"],
                        demo["invoice_date"],
                        demo["po_number"],
                        demo["currency"],
                        "DEMO-LOCATION",
                        "Store (S)",
                        "DEMO-TAX",
                        demo["subtotal"],
                        demo["tax_total"],
                        demo["total"],
                        "needs_review",
                        "fictional_demo_seed",
                        now,
                        now,
                        json_dumps(["Fictional demo data. Do not use for payment."]),
                        1,
                        1,
                        raw_text,
                    ),
                )
                prepared_demo_lines = []
                costs = {item[0]: item[4] for item in catalog}
                for source_line in demo["lines"]:
                    line = dict(source_line)
                    if line.get("rms_item_id"):
                        line["catalog_item_id"] = f"demo:{line['rms_item_id']}"
                        line["rms_unit_cost"] = costs[line["rms_item_id"]]
                    prepared_demo_lines.append(line)
                self._replace_lines(
                    conn,
                    demo["id"],
                    self._prepare_lines(prepared_demo_lines, currency=demo["currency"]),
                )
                self._audit(
                    conn,
                    demo["id"],
                    "fictional_demo_seeded",
                    actor="demo_seed",
                    version=1,
                    to_status="needs_review",
                    details={"fictional": True},
                    created_at=now,
                )
                invoices_created += 1
        return {
            "invoices_created": invoices_created,
            "catalog_items_upserted": catalog_upserted,
            "message": "Fictional demo data is ready for review.",
        }
