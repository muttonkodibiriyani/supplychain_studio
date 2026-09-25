from __future__ import annotations

import io
import os
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.exporter import (
    DETAIL_COLUMNS,
    HEADER_COLUMNS,
    TAX_COLUMNS,
    build_target_workbook,
)
from backend.service import Conflict, InvoiceService, Settings, ValidationFailure


PRIVATE_MASTER = Path(os.getenv("PRIVATE_RMS_MASTER_PATH", ""))
PRIVATE_ALIASES = Path(os.getenv("PRIVATE_ALIASES_PATH", ""))


def make_service(tmp_path: Path) -> InvoiceService:
    return InvoiceService(
        Settings(
            database_path=tmp_path / "invoice.sqlite3",
            source_dir=tmp_path / "sources",
            export_dir=tmp_path / "exports",
            workers=0,
        )
    )


def headers(sheet: object) -> list[object]:
    return [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]


def test_settings_are_versioned_single_brand_and_reject_non_invoice_policy(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)
    initial = service.get_settings()
    assert initial["brand_label"] == ""
    assert initial["include_upc_in_export"] is False
    assert initial["target_cost_policy"] == {
        "mode": "invoice_only",
        "maximum_absolute_difference_aed": 10,
    }
    updated = service.update_settings(
        initial["version"],
        {
            "brand_label": "Synthetic Brand",
            "location": "00001",
            "location_type": "Store (S)",
            "include_upc_in_export": True,
            "supplier_rules": [
                {
                    "supplier_id": "00077",
                    "supplier_site": "00088",
                    "tax_code": "TEST05",
                }
            ],
            "target_cost_policy": {
                "mode": "invoice_only",
                "maximum_absolute_difference_aed": "10.00",
            },
        },
    )
    assert updated["location"] == "00001"
    assert updated["include_upc_in_export"] is True
    assert updated["supplier_rules"][0]["supplier_site"] == "00088"
    with pytest.raises(Conflict):
        service.update_settings(initial["version"], updated)
    with pytest.raises(ValidationFailure) as policy:
        service.update_settings(
            updated["version"],
            {
                **updated,
                "target_cost_policy": {
                    "mode": "rms_within_tolerance",
                    "maximum_absolute_difference_aed": 10,
                },
            },
        )
    assert policy.value.errors[0]["field"] == "target_cost_policy.mode"


def test_pure_target_builder_defaults_upc_blank_and_opt_in_preserves_text() -> None:
    invoice = {
        "invoice_number": "INV-1",
        "document": "INV-1",
        "supplier_site": "SITE-1",
        "location": "LOC-1",
        "location_type": "S",
        "invoice_date": "2026-09-25",
        "subtotal": "10.00",
        "tax_total": "0.50",
        "tax_code": "VAT5",
        "lines": [
            {
                "rms_parent_item": "ITEM-1",
                "rms_upc": "000045678901",
                "target_unit_cost": "10.00",
                "quantity": "1",
            }
        ],
    }
    default_workbook = load_workbook(
        io.BytesIO(build_target_workbook([invoice])), data_only=True
    )
    enabled_workbook = load_workbook(
        io.BytesIO(
            build_target_workbook([invoice], include_upc_in_export=True)
        ),
        data_only=True,
    )

    assert default_workbook["Details"]["C2"].value is None
    assert default_workbook["Details"]["C2"].number_format == "@"
    assert enabled_workbook["Details"]["C2"].value == "000045678901"
    assert enabled_workbook["Details"]["C2"].data_type == "s"
    assert enabled_workbook["Details"]["C2"].number_format == "@"


def test_existing_settings_database_migrates_to_default_blank_upc_policy(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(database_path) as connection:
        connection.executescript(
            """
            CREATE TABLE app_settings (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                brand_label TEXT NOT NULL DEFAULT '',
                location TEXT NOT NULL DEFAULT '',
                location_type TEXT NOT NULL DEFAULT '',
                supplier_rules_json TEXT NOT NULL DEFAULT '[]',
                target_cost_policy_json TEXT NOT NULL,
                version INTEGER NOT NULL DEFAULT 1,
                updated_at TEXT NOT NULL
            );
            INSERT INTO app_settings
                (id,brand_label,location,location_type,supplier_rules_json,
                 target_cost_policy_json,version,updated_at)
            VALUES
                (1,'Legacy brand','','','[]',
                 '{"mode":"invoice_only","maximum_absolute_difference_aed":10}',
                 7,'2026-09-25T00:00:00.000Z');
            """
        )
    service = InvoiceService(
        Settings(
            database_path=database_path,
            source_dir=tmp_path / "sources",
            export_dir=tmp_path / "exports",
            workers=0,
        )
    )

    settings = service.get_settings()
    assert settings["brand_label"] == "Legacy brand"
    assert settings["version"] == 7
    assert settings["include_upc_in_export"] is False


def test_exact_consolidated_workbook_multi_invoice_links_and_text_identifiers(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    service.import_catalog(
        "leading.csv",
        b"rms_item_id,upc,description,supplier_id,uom,unit_cost\n"
        b"0000123,000000456789,Fictional Hydrating Cleanser 250ml,DEMO-SUPPLIER,EA,45.00\n",
    )
    selected = service.list_catalog(
        search="0000123", supplier_id="DEMO-SUPPLIER", limit=10
    )["items"][0]

    first = service.get_invoice("demo-invoice-001")
    first_lines = first["lines"]
    first_lines[0].update(
        {
            "catalog_item_id": selected["catalog_item_id"],
            "rms_item_id": selected["rms_item_id"],
            "match_status": "confirmed",
        }
    )
    first = service.update_invoice(
        first["id"],
        first["version"],
        {
            "document": "000INV-1",
            "supplier_site": "000077",
            "po_number": "0000456",
            "location": "00001",
            "lines": first_lines,
        },
    )
    first = service.approve(first["id"], first["version"])

    second = service.get_invoice("demo-invoice-002")
    second_line = second["lines"][0]
    second_line.update(
        {
            "rms_item_id": "RMS-2001",
            "match_status": "confirmed",
        }
    )
    second = service.update_invoice(
        second["id"], second["version"], {"lines": [second_line]}
    )
    second = service.approve(second["id"], second["version"])

    result = service.create_export([first["id"], second["id"]])
    workbook = load_workbook(result["path"], data_only=False)
    assert workbook.sheetnames == ["Header", "Tax_Breakdown", "Details"]
    assert headers(workbook["Header"]) == HEADER_COLUMNS
    assert headers(workbook["Tax_Breakdown"]) == TAX_COLUMNS
    assert headers(workbook["Details"]) == DETAIL_COLUMNS
    assert workbook["Header"].max_row == 3
    assert workbook["Tax_Breakdown"].max_row == 3
    assert workbook["Details"].max_row == 4

    header_transactions = [workbook["Header"].cell(row, 1).value for row in (2, 3)]
    tax_transactions = [workbook["Tax_Breakdown"].cell(row, 1).value for row in (2, 3)]
    detail_transactions = [
        workbook["Details"].cell(row, 1).value
        for row in range(2, workbook["Details"].max_row + 1)
    ]
    assert header_transactions == ["1", "2"]
    assert tax_transactions == ["1", "2"]
    assert set(detail_transactions) == {"1", "2"}
    assert detail_transactions.count("1") == 2
    assert detail_transactions.count("2") == 1

    header = workbook["Header"]
    details = workbook["Details"]
    assert header["B2"].value == "000INV-1"
    assert header["C2"].value == "000077"
    assert header["D2"].value == "0000456"
    assert header["E2"].value == "00001"
    assert header["G2"].data_type == "d"
    assert header["G2"].value.date().isoformat() == "2026-09-20"
    assert header["H2"].value == 215
    assert header["I2"].value == pytest.approx(10.75)
    assert details["B2"].value == "0000123"
    assert details["C2"].value is None
    for coordinate in ("A2", "B2", "C2", "D2", "E2"):
        assert header[coordinate].data_type == "s"
        assert header[coordinate].number_format == "@"
    for coordinate in ("A2", "B2"):
        assert details[coordinate].data_type == "s"
        assert details[coordinate].number_format == "@"
    assert details["C2"].number_format == "@"

    settings = service.get_settings()
    opted_in = service.update_settings(
        settings["version"], {**settings, "include_upc_in_export": True}
    )
    assert opted_in["include_upc_in_export"] is True
    opt_in_result = service.create_export([first["id"], second["id"]])
    opt_in_workbook = load_workbook(opt_in_result["path"], data_only=False)
    assert opt_in_workbook["Details"]["C2"].value == "000000456789"
    assert opt_in_workbook["Details"]["C2"].data_type == "s"
    assert opt_in_workbook["Details"]["C2"].number_format == "@"
    export_events = [
        event
        for event in service.audit(first["id"])["items"]
        if event["event_type"] == "invoice_exported"
    ]
    assert [
        event["details"]["include_upc_in_export"] for event in export_events[:2]
    ] == [True, False]


def test_invoice_cost_is_always_target_and_rms_variance_requires_review(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    invoice = service.get_invoice("demo-invoice-001")
    lines = invoice["lines"]
    lines[0].update(unit_price="40.00", line_total="80.00")
    within = service.update_invoice(
        invoice["id"],
        invoice["version"],
        {
            "subtotal": "205.00",
            "tax_total": "10.25",
            "total": "215.25",
            "lines": lines,
        },
    )
    assert within["lines"][0]["rms_unit_cost"] == 45
    assert within["lines"][0]["target_unit_cost"] == 40
    assert within["lines"][0]["target_cost_variance"] == 5
    assert within["lines"][0]["target_cost_comparison_status"] == "within_tolerance"
    assert not within["target_cost_review_required"]

    lines = within["lines"]
    lines[0].update(unit_price="20.00", line_total="40.00")
    above = service.update_invoice(
        within["id"],
        within["version"],
        {
            "subtotal": "165.00",
            "tax_total": "8.25",
            "total": "173.25",
            "lines": lines,
        },
    )
    assert above["lines"][0]["target_unit_cost"] == 20
    assert above["lines"][0]["rms_unit_cost"] == 45
    assert above["lines"][0]["target_cost_variance"] == 25
    assert above["lines"][0]["target_cost_comparison_status"] == "above_tolerance"
    assert above["target_cost_review_required"]
    with pytest.raises(ValidationFailure) as review:
        service.approve(above["id"], above["version"])
    assert any(
        error["code"] == "target_cost_review_required"
        for error in review.value.errors
    )
    approved = service.approve(
        above["id"],
        above["version"],
        acknowledge_target_cost_variance=True,
    )
    result = service.create_export([approved["id"]])
    workbook = load_workbook(result["path"], data_only=True)
    assert workbook["Header"]["H2"].value == pytest.approx(165.00)
    assert workbook["Tax_Breakdown"]["C2"].value == pytest.approx(165.00)
    assert workbook["Details"]["D2"].value == pytest.approx(20.00)


def test_non_aed_rms_comparison_is_unavailable_not_converted(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    invoice = service.get_invoice("demo-invoice-001")
    edited = service.update_invoice(
        invoice["id"],
        invoice["version"],
        {"currency": "USD", "lines": invoice["lines"]},
    )
    assert edited["lines"][0]["target_unit_cost"] == 45
    assert edited["lines"][0]["target_cost_variance"] is None
    assert edited["lines"][0]["target_cost_comparison_status"] == "unavailable_currency"
    assert edited["target_cost_review_required"]


def test_repeating_net_unit_cost_rounds_each_extension_half_up_to_cents(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    service.import_catalog(
        "repeating.csv",
        b"rms_item_id,upc,description,supplier_id,uom,unit_cost\n"
        b"REPEAT-ITEM,000333,Fictional Hydrating Cleanser 250ml,DEMO-SUPPLIER,EA,3.3333333333\n",
    )
    selected = service.list_catalog(
        search="REPEAT-ITEM", supplier_id="DEMO-SUPPLIER", limit=10
    )["items"][0]
    invoice = service.get_invoice("demo-invoice-001")
    lines = invoice["lines"]
    lines[0].update(
        quantity="300",
        unit_price="3.3333333333",
        line_total="1000.00",
        catalog_item_id=selected["catalog_item_id"],
        rms_item_id=selected["rms_item_id"],
        match_status="confirmed",
    )
    edited = service.update_invoice(
        invoice["id"],
        invoice["version"],
        {
            "subtotal": "1125.00",
            "tax_total": "56.25",
            "total": "1181.25",
            "lines": lines,
        },
    )
    assert edited["lines"][0]["target_unit_cost"] == pytest.approx(3.3333333333)
    assert edited["lines"][0]["rms_unit_cost"] == pytest.approx(3.3333333333)
    assert edited["lines"][0]["target_cost_variance"] == 0
    assert edited["lines"][0]["target_cost_comparison_status"] == "within_tolerance"
    assert edited["target_total_ex_tax"] == 1125
    approved = service.approve(edited["id"], edited["version"])
    exported = service.create_export([approved["id"]])
    workbook = load_workbook(exported["path"], data_only=True)
    assert workbook["Header"]["H2"].value == 1125
    assert workbook["Details"]["D2"].value == pytest.approx(3.3333333333)


def test_non_invoice_and_any_unresolved_line_block_whole_invoice(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    invoice = service.get_invoice("demo-invoice-002")
    with pytest.raises(ValidationFailure) as unresolved:
        service.approve(invoice["id"], invoice["version"])
    assert any(error["code"] == "unmapped" for error in unresolved.value.errors)

    line = invoice["lines"][0]
    line.update(rms_item_id="RMS-2001", match_status="confirmed")
    credit = service.update_invoice(
        invoice["id"],
        invoice["version"],
        {"document_type": "credit_note", "lines": [line]},
    )
    with pytest.raises(ValidationFailure) as non_invoice:
        service.approve(credit["id"], credit["version"])
    assert any(error["code"] == "not_invoice" for error in non_invoice.value.errors)


def test_supplier_duplicate_item_and_upc_rows_are_preserved_and_filterable(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)
    result = service.import_catalog(
        "duplicates.csv",
        b"rms_item_id,upc,description,supplier_id,uom,unit_cost\n"
        b"0001,000111,Shared Item,SUP-A,EA,1.00\n"
        b"0001,000222,Shared Item,SUP-B,EA,2.00\n"
        b"0001,000333,Shared Item,SUP-A,EA,1.10\n",
    )
    assert result["imported"] == 3
    assert service.stats()["catalog_items"] == 3
    supplier_a = service.list_catalog(search=None, supplier_id="SUP-A", limit=2)
    supplier_b = service.list_catalog(search=None, supplier_id="SUP-B", limit=2)
    assert supplier_a["total"] == 2
    assert supplier_b["total"] == 1
    assert {item["upc"] for item in supplier_a["items"]} == {"000111", "000333"}
    assert len({item["catalog_item_id"] for item in supplier_a["items"]}) == 2


def test_explicit_po_precedence_and_po_box_is_never_exported(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    service.seed_demo()
    imported = service.import_catalog(
        "orders.csv",
        b"rms_item_id,upc,description,supplier_id,uom,unit_cost,po_number\n"
        b"MASTER-ITEM,000999,Fictional Hydrating Cleanser 250ml,DEMO-SUPPLIER,EA,45.00,MASTER-0001\n"
        b"BAD-ITEM,001000,Bad order field,DEMO-SUPPLIER,EA,1.00,PO Box 123\n",
    )
    assert imported["imported"] == 2
    assert any("ignored non-order" in warning for warning in imported["warnings"])
    selected = service.list_catalog(
        search="MASTER-ITEM", supplier_id="DEMO-SUPPLIER", limit=10
    )["items"][0]
    invoice = service.get_invoice("demo-invoice-001")
    lines = invoice["lines"]
    lines[0].update(
        catalog_item_id=selected["catalog_item_id"],
        rms_item_id=selected["rms_item_id"],
        match_status="confirmed",
    )
    fallback = service.update_invoice(
        invoice["id"], invoice["version"], {"po_number": None, "lines": lines}
    )
    fallback = service.approve(fallback["id"], fallback["version"])
    exported = service.create_export([fallback["id"]])
    workbook = load_workbook(exported["path"], data_only=True)
    assert workbook["Header"]["D2"].value == "MASTER-0001"

    current = service.get_invoice(fallback["id"])
    with pytest.raises(ValidationFailure) as po_box:
        service.update_invoice(
            current["id"], current["version"], {"po_number": "PO Box 123"}
        )
    assert po_box.value.errors[0]["code"] == "invalid_order_number"


def test_sparse_185_column_rms_projection_uses_parent_item_and_strips_known_upc_prefix(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)
    required = {
        1: "BRAND",
        6: "ITEM_PARENT",
        7: "ITEM",
        16: "ITEM_DESC",
        19: "STANDARD_UOM",
        28: "SUPPLIER",
        33: "UNIT_COST",
        178: "SUPPLIER_NAME",
        179: "SUPPLIER_COUNTRY",
        185: "EXTRA_WIDE_COLUMN",
    }
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Item Master")
    headers_row = [required.get(index, f"UNUSED_{index}") for index in range(1, 186)]
    sheet.append(headers_row)
    for supplier, item, upc, cost in (
        ("00077", "000123", "ULT_000045678901", "12.34"),
        ("00088", "000123", "ULT_000045678901", "13.34"),
    ):
        row = [None] * 185
        row[5] = item
        row[6] = upc
        row[15] = "Synthetic projected item"
        row[18] = "EA"
        row[27] = supplier
        row[32] = cost
        row[177] = "Synthetic supplier"
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    result = service.import_catalog("wide-rms.xlsx", output.getvalue())
    assert result == {"imported": 2, "skipped": 0, "warnings": []}
    assert service.stats()["catalog_items"] == 2
    first = service.list_catalog(search=None, supplier_id="00077", limit=10)["items"][0]
    assert first["rms_item_id"] == "000123"
    assert first["upc"] == "000045678901"
    assert first["unit_cost"] == pytest.approx(12.34)


def test_catalog_row_guard_accepts_more_than_old_100k_limit(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    row_count = 100_005
    content = "rms_item_id,description,supplier_id,uom,unit_cost\n" + "".join(
        f"SYN-{index},Synthetic item {index},SYN-SUP,EA,1.00\n"
        for index in range(row_count)
    )
    result = service.import_catalog("large-synthetic.csv", content.encode())
    assert result == {"imported": row_count, "skipped": 0, "warnings": []}
    assert service.stats()["catalog_items"] == row_count


@pytest.mark.skipif(
    os.getenv("RUN_PRIVATE_RMS_TESTS") != "1"
    or not PRIVATE_MASTER.is_file()
    or not PRIVATE_ALIASES.is_file(),
    reason="explicit private acceptance run only",
)
def test_supplied_full_rms_master_and_confirmed_alias_import(tmp_path: Path) -> None:
    expected_rows = int(os.environ["PRIVATE_RMS_EXPECTED_ROWS"])
    service = make_service(tmp_path)
    result = service.import_catalog(PRIVATE_MASTER.name, PRIVATE_MASTER.read_bytes())
    assert result == {"imported": expected_rows, "skipped": 0, "warnings": []}
    assert service.stats()["catalog_items"] == expected_rows
    aliases = service.import_aliases(PRIVATE_ALIASES.name, PRIVATE_ALIASES.read_bytes())
    assert aliases == {"imported": 5, "skipped": 0, "warnings": []}
    assert service.list_aliases()["total"] == 5
    assert service.get_settings()["target_cost_policy"]["mode"] == "invoice_only"
