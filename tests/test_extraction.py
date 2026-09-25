from __future__ import annotations

import importlib.util
import shutil
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from backend.extraction import (
    DocumentExtractionError,
    ExtractionLimitError,
    ExtractionLimits,
    UnsupportedDocumentError,
    _ocr_pdf_pages,
    _run_tesseract,
    extract_document,
)


FIXTURES = Path(__file__).parent / "fixtures"
OCR_AVAILABLE = shutil.which("tesseract") is not None
PDFIUM_AVAILABLE = importlib.util.find_spec("pypdfium2") is not None


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _write_text_pdf(path: Path, lines: list[str]) -> None:
    """Write a small standards-compliant PDF with a real text content stream."""

    commands = ["BT", "/F1 11 Tf", "50 780 Td", "14 TL"]
    for index, line in enumerate(lines):
        if index:
            commands.append("T*")
        commands.append(f"({_pdf_escape(line)}) Tj")
    commands.append("ET")
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    document = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(document))
        document.extend(f"{number} 0 obj\n".encode("ascii"))
        document.extend(body)
        document.extend(b"\nendobj\n")
    xref_offset = len(document)
    document.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    document.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        document.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    document.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode(
            "ascii"
        )
    )
    path.write_bytes(document)


def _write_facturx_pdf(path: Path, *, type_code: str = "380") -> None:
    from pypdf import PdfReader, PdfWriter

    base_path = path.with_name(path.stem + "-base.pdf")
    _write_text_pdf(base_path, ["TAX INVOICE", "Embedded structured invoice data"])
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<CrossIndustryInvoice>
  <ExchangedDocument>
    <ID>SYN-XML-01</ID><TypeCode>{type_code}</TypeCode>
    <IssueDateTime><DateTimeString format="102">20260925</DateTimeString></IssueDateTime>
  </ExchangedDocument>
  <SupplyChainTradeTransaction>
    <IncludedSupplyChainTradeLineItem>
      <SpecifiedTradeProduct>
        <GlobalID schemeID="0160">0123456789012</GlobalID>
        <SellerAssignedID>SYN-X1</SellerAssignedID><Name>Synthetic Barrier Cream</Name>
      </SpecifiedTradeProduct>
      <SpecifiedLineTradeAgreement>
        <GrossPriceProductTradePrice><ChargeAmount>9.25</ChargeAmount></GrossPriceProductTradePrice>
        <NetPriceProductTradePrice><ChargeAmount>8.75</ChargeAmount></NetPriceProductTradePrice>
      </SpecifiedLineTradeAgreement>
      <SpecifiedLineTradeDelivery><BilledQuantity unitCode="C62">4</BilledQuantity></SpecifiedLineTradeDelivery>
      <SpecifiedLineTradeSettlement>
        <ApplicableTradeTax><RateApplicablePercent>5</RateApplicablePercent></ApplicableTradeTax>
        <SpecifiedTradeSettlementLineMonetarySummation><LineTotalAmount>35.00</LineTotalAmount></SpecifiedTradeSettlementLineMonetarySummation>
      </SpecifiedLineTradeSettlement>
    </IncludedSupplyChainTradeLineItem>
    <ApplicableHeaderTradeAgreement>
      <SellerTradeParty><Name>Synthetic Trading</Name><SpecifiedTaxRegistration><ID>SYN-TAX</ID></SpecifiedTaxRegistration></SellerTradeParty>
      <BuyerOrderReferencedDocument><IssuerAssignedID>Campaign - PO - SYN-PO-77</IssuerAssignedID></BuyerOrderReferencedDocument>
    </ApplicableHeaderTradeAgreement>
    <ApplicableHeaderTradeSettlement>
      <InvoiceCurrencyCode>AED</InvoiceCurrencyCode>
      <SpecifiedTradeSettlementHeaderMonetarySummation>
        <LineTotalAmount>35.00</LineTotalAmount><TaxBasisTotalAmount>35.00</TaxBasisTotalAmount>
        <TaxTotalAmount currencyID="AED">1.75</TaxTotalAmount><GrandTotalAmount>36.75</GrandTotalAmount>
      </SpecifiedTradeSettlementHeaderMonetarySummation>
    </ApplicableHeaderTradeSettlement>
  </SupplyChainTradeTransaction>
</CrossIndustryInvoice>""".encode("utf-8")
    reader = PdfReader(base_path)
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.add_attachment("factur-x.xml", xml)
    with path.open("wb") as handle:
        writer.write(handle)
    base_path.unlink()


def _invoice_lines(invoice_number: str = "PDF-100") -> list[str]:
    return [
        "INVOICE",
        "Supplier: Desert Beauty Trading LLC",
        f"Invoice No: {invoice_number}",
        "Invoice Date: 24 September 2026",
        "PO Number: PO-880021",
        "Currency: AED",
        "Description                     Qty UOM Unit Price Line Total VAT",
        "Hydrating Cleanser 50ml          2 EA 25.00 50.00 5%",
        "Subtotal AED 50.00",
        "VAT Total AED 2.50",
        "Grand Total AED 52.50",
    ]


class ExtractionTests(unittest.TestCase):
    def test_text_invoice_captures_evidence_headers_lines_and_totals(self) -> None:
        result = extract_document(FIXTURES / "sample_invoice.txt")

        self.assertEqual(result["supplier_name"], "Desert Beauty Trading LLC")
        self.assertEqual(result["invoice_number"], "DB-2026-1042")
        self.assertEqual(result["invoice_date"], "2026-09-24")
        self.assertEqual(result["currency"], "AED")
        self.assertEqual(result["total"], 68.25)
        self.assertEqual(len(result["lines"]), 2)
        self.assertEqual(result["lines"][0]["description"], "Hydrating Cleanser 50ml")
        self.assertEqual(result["lines"][0]["quantity"], 2.0)
        self.assertEqual(result["lines"][0]["uom"], "EA")
        self.assertIn("Supplier: Desert Beauty", result["raw_text"])
        self.assertEqual(result["page_texts"][0]["raw_text"], result["raw_text"])
        self.assertTrue(result["lines"][0]["id"].startswith("line-"))

    def test_fixed_width_layout_uses_header_columns_and_joins_wrapped_description(self) -> None:
        result = extract_document(FIXTURES / "layouts" / "fixed_width_wrapped.txt")

        self.assertEqual(result["document_type"], "invoice")
        self.assertEqual(result["table_extraction_method"], "layout_header")
        self.assertEqual(len(result["lines"]), 2)
        first = result["lines"][0]
        self.assertEqual(
            first["description"],
            "Botanical Face Cleanser Refill pouch, fragrance free",
        )
        self.assertEqual(first["upc"], "012345678901")
        self.assertEqual(first["item_code"], "SYN-001")
        self.assertEqual(first["quantity"], 2.0)
        self.assertEqual(first["uom"], "EA")
        self.assertEqual(first["unit_price"], 25.0)
        self.assertEqual(first["net_unit_price"], 25.0)
        self.assertIsNone(first["gross_unit_price"])
        self.assertEqual(first["line_total"], 50.0)
        self.assertGreater(first["source_rows"][1], first["source_rows"][0])

    def test_non_invoice_types_stay_explicit_and_reviewable(self) -> None:
        credit = extract_document(FIXTURES / "layouts" / "credit_note.txt")
        plain_credit = extract_document(
            FIXTURES / "layouts" / "plain_credit_note.txt"
        )
        credit_memo = extract_document(FIXTURES / "layouts" / "credit_memo.txt")
        delivery = extract_document(FIXTURES / "layouts" / "delivery_note.txt")
        unknown = extract_document(FIXTURES / "layouts" / "unknown_numeric_report.txt")

        self.assertEqual(credit["document_type"], "credit_note")
        self.assertGreaterEqual(credit["document_type_confidence"], 0.9)
        self.assertEqual(plain_credit["document_type"], "credit_note")
        self.assertGreaterEqual(plain_credit["document_type_confidence"], 0.9)
        self.assertEqual(credit_memo["document_type"], "credit_note")
        self.assertGreaterEqual(credit_memo["document_type_confidence"], 0.9)
        self.assertEqual(delivery["document_type"], "delivery_note")
        self.assertEqual(delivery["lines"], [])
        self.assertEqual(unknown["document_type"], "unknown")
        self.assertEqual(unknown["document_type_confidence"], 0.0)
        self.assertEqual(len(unknown["lines"]), 2)
        self.assertTrue(any("route" in warning.lower() for warning in unknown["warnings"]))

    def test_credit_invoice_ocr_layout_reconciles_without_conflating_credit_note(self) -> None:
        result = extract_document(
            FIXTURES / "layouts" / "credit_invoice_ocr_zero_tax.txt"
        )

        self.assertEqual(result["document_type"], "invoice")
        self.assertGreaterEqual(result["document_type_confidence"], 0.9)
        self.assertEqual(result["supplier_name"], "Synthetic Carbon Company")
        self.assertEqual(result["invoice_number"], "SYN-CARBON-01")
        self.assertEqual(result["po_number"], "SYN-PO-900")
        self.assertEqual(result["currency"], "KWD")
        self.assertIsNone(result["invoice_date"])
        self.assertEqual(result["total"], 636.5)
        self.assertEqual(result["tax_total"], 0.0)
        self.assertEqual(result["subtotal"], 636.5)
        self.assertEqual(
            result["field_provenance"]["tax_total"],
            "absent_tax_treated_as_zero",
        )
        self.assertEqual(
            result["field_provenance"]["subtotal"],
            "derived_grand_total_minus_tax_total",
        )
        self.assertEqual(result["table_diagnostics"]["possible_unparsed_rows"], 0)
        self.assertEqual(len(result["lines"]), 3)
        self.assertEqual(
            [line["quantity"] for line in result["lines"]],
            [96.0, 110.0, 140.0],
        )
        self.assertEqual(
            [line["net_unit_price"] for line in result["lines"]],
            [2.0, 1.75, 1.8],
        )
        self.assertEqual(
            [line["line_total"] for line in result["lines"]],
            [192.0, 192.5, 252.0],
        )
        self.assertEqual(
            result["lines"][2]["line_total_source"],
            "gross_extension_validated_by_zero_discount_and_printed_rate",
        )
        self.assertIn("quantity", result["lines"][0]["derived_fields"])
        self.assertIn("quantity", result["lines"][1]["derived_fields"])

    def test_numeric_text_without_table_headers_does_not_create_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.txt"
            path.write_text(
                "ACCOUNT ACTIVITY\nReference 12345\nOpening balance 2 50.00 100.00",
                encoding="utf-8",
            )
            result = extract_document(path)

        self.assertEqual(result["document_type"], "unknown")
        self.assertEqual(result["lines"], [])
        self.assertEqual(result["table_extraction_method"], "none")

    def test_explicit_net_amount_derives_high_precision_net_unit_cost(self) -> None:
        result = extract_document(FIXTURES / "layouts" / "net_amount_derived.txt")

        self.assertEqual(result["subtotal"], 10.0)
        self.assertEqual(
            result["field_provenance"]["subtotal"],
            "derived_grand_total_minus_tax_total",
        )
        line = result["lines"][0]
        self.assertEqual(line["printed_unit_price"], 4.5)
        self.assertEqual(line["printed_unit_price_basis"], "gross")
        self.assertEqual(line["gross_unit_price"], 4.5)
        self.assertEqual(line["line_total"], 10.0)
        self.assertEqual(line["line_total_basis"], "net")
        self.assertEqual(line["unit_price"], 3.33333333)
        self.assertEqual(line["net_unit_price"], 3.33333333)
        self.assertEqual(
            line["unit_price_source"],
            "derived_net_line_amount_divided_by_quantity",
        )
        self.assertEqual(line["derived_fields"], ["unit_price", "net_unit_price"])

    def test_embedded_facturx_xml_preserves_structured_lines_and_document_type(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            invoice_path = Path(directory) / "structured-invoice.pdf"
            credit_path = Path(directory) / "structured-credit.pdf"
            _write_facturx_pdf(invoice_path)
            _write_facturx_pdf(credit_path, type_code="381")
            invoice = extract_document(invoice_path)
            credit = extract_document(credit_path)

        self.assertEqual(invoice["extraction_method"], "pdf_facturx_xml")
        self.assertEqual(invoice["document_type"], "invoice")
        self.assertEqual(invoice["invoice_number"], "SYN-XML-01")
        self.assertEqual(invoice["invoice_date"], "2026-09-25")
        self.assertEqual(invoice["po_number"], "SYN-PO-77")
        self.assertEqual(invoice["buyer_order_reference"], "Campaign - PO - SYN-PO-77")
        self.assertEqual(invoice["subtotal"], 35.0)
        self.assertEqual(invoice["tax_total"], 1.75)
        self.assertEqual(len(invoice["lines"]), 1)
        line = invoice["lines"][0]
        self.assertEqual(line["quantity"], 4.0)
        self.assertEqual(line["uom"], "EA")
        self.assertEqual(line["unit_price"], 8.75)
        self.assertEqual(line["net_unit_price"], 8.75)
        self.assertEqual(line["gross_unit_price"], 9.25)
        self.assertEqual(line["line_total"], 35.0)
        self.assertEqual(line["upc"], "0123456789012")
        self.assertEqual(credit["document_type"], "credit_note")

    def test_adjacent_explicit_discount_rows_pair_without_losing_source_evidence(self) -> None:
        result = extract_document(FIXTURES / "layouts" / "paired_discount_rows.txt")

        self.assertEqual(result["subtotal"], 34.0)
        self.assertEqual(result["tax_total"], 1.7)
        self.assertEqual(
            result["field_provenance"]["tax_total"],
            "derived_sum_complete_explicit_line_tax_amounts",
        )
        self.assertEqual(result["supplier_name"], "Synthetic Distribution LLC")
        self.assertEqual(len(result["lines"]), 2)
        self.assertEqual(result["table_diagnostics"]["source_data_rows"], 4)
        self.assertEqual(result["table_diagnostics"]["paired_adjustment_rows"], 2)
        self.assertEqual(result["table_diagnostics"]["unpaired_adjustment_rows"], 0)
        self.assertEqual([line["quantity"] for line in result["lines"]], [3.0, 2.0])
        self.assertEqual([line["net_unit_price"] for line in result["lines"]], [8.0, 5.0])
        self.assertEqual([line["gross_unit_price"] for line in result["lines"]], [10.0, 5.0])
        self.assertEqual([line["line_total"] for line in result["lines"]], [24.0, 10.0])
        self.assertTrue(all(len(line["source_adjustments"]) == 1 for line in result["lines"]))
        self.assertTrue(all(line["paired_source_line_count"] == 2 for line in result["lines"]))

    def test_missing_values_are_not_calculated_or_invented(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "partial.txt"
            path.write_text(
                "Invoice No: PART-1\nDescription Qty Unit Price Line Total\nWidget 2 3.00 6.00",
                encoding="utf-8",
            )
            result = extract_document(path)

        self.assertIsNone(result["supplier_name"])
        self.assertIsNone(result["subtotal"])
        self.assertIsNone(result["tax_total"])
        self.assertIsNone(result["total"])
        self.assertEqual(result["lines"][0]["line_total"], 6.0)
        self.assertIsNone(result["lines"][0]["net_unit_price"])
        self.assertEqual(result["lines"][0]["derived_fields"], [])

    def test_ambiguous_numeric_date_is_withheld(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ambiguous.txt"
            path.write_text("Invoice No: D-1\nDate: 01/02/2026", encoding="utf-8")
            result = extract_document(path)

        self.assertIsNone(result["invoice_date"])
        self.assertTrue(any("ambiguous" in warning for warning in result["warnings"]))

    def test_decimal_comma_and_grouping_are_parsed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "european.txt"
            path.write_text(
                "Invoice No: EU-1\nCurrency: EUR\n"
                "Description  Qty UOM Unit Price Line Total\n"
                "Professional Serum  2 EA 1.234,56 2.469,12\n"
                "Grand Total EUR 2.469,12",
                encoding="utf-8",
            )
            result = extract_document(path)

        self.assertEqual(result["lines"][0]["unit_price"], 1234.56)
        self.assertEqual(result["lines"][0]["line_total"], 2469.12)
        self.assertEqual(result["total"], 2469.12)

    def test_structured_csv_extracts_one_invoice(self) -> None:
        result = extract_document(FIXTURES / "sample_invoice.csv")

        self.assertEqual(result["extraction_method"], "csv_structured")
        self.assertEqual(result["supplier_id"], "SUP-1007")
        self.assertEqual(result["total"], 68.25)
        self.assertEqual(len(result["lines"]), 2)

    def test_mixed_invoice_csv_is_rejected_instead_of_merged(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.csv"
            path.write_text(
                "invoice_number,supplier_name,description,quantity,unit_price,line_total\n"
                "INV-1,One Ltd,Item A,1,2,2\n"
                "INV-2,One Ltd,Item B,1,3,3\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(DocumentExtractionError, "multiple invoices"):
                extract_document(path)

    def test_xlsx_structured_columns(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invoice.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "invoice_number",
                    "invoice_date",
                    "supplier_name",
                    "currency",
                    "description",
                    "quantity",
                    "uom",
                    "unit_price",
                    "line_total",
                    "total",
                ]
            )
            sheet.append(
                ["XL-1", "2026-09-24", "Sheet Supplier", "AED", "Cleanser 50ml", 3, "EA", 10, 30, 30]
            )
            workbook.save(path)
            workbook.close()
            result = extract_document(path)

        self.assertEqual(result["extraction_method"], "xlsx_structured")
        self.assertEqual(result["invoice_number"], "XL-1")
        self.assertEqual(result["lines"][0]["quantity"], 3.0)

    def test_embedded_text_pdf_uses_pdf_text_and_preserves_page_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invoice.pdf"
            _write_text_pdf(path, _invoice_lines())
            result = extract_document(path)

        self.assertEqual(result["extraction_method"], "pdf_text")
        self.assertEqual(result["invoice_number"], "PDF-100")
        self.assertEqual(result["pages"], 1)
        self.assertEqual(result["page_texts"][0]["method"], "pdf_text")
        self.assertEqual(result["lines"][0]["description"], "Hydrating Cleanser 50ml")

    def test_encrypted_pdf_has_actionable_error(self) -> None:
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "encrypted.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.encrypt("secret")
            with path.open("wb") as handle:
                writer.write(handle)
            with self.assertRaisesRegex(DocumentExtractionError, "Encrypted PDF"):
                extract_document(path)

    def test_corrupt_pdf_has_actionable_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "corrupt.pdf"
            path.write_bytes(b"%PDF-1.4\nthis is not a document")
            with self.assertRaisesRegex(DocumentExtractionError, "corrupt or unreadable"):
                extract_document(path)

    def test_page_limit_fails_before_partial_extraction(self) -> None:
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "two-pages.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.add_blank_page(width=100, height=100)
            with path.open("wb") as handle:
                writer.write(handle)
            with self.assertRaisesRegex(ExtractionLimitError, "2 pages"):
                extract_document(path, limits=ExtractionLimits(max_pages=1))

    def test_file_size_limit_fails_before_parsing(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "large.txt"
            path.write_text("Invoice No: TOO-BIG\n" + ("x" * 200), encoding="utf-8")
            with self.assertRaisesRegex(ExtractionLimitError, "limit is 32 bytes"):
                extract_document(path, limits=ExtractionLimits(max_file_bytes=32))

    def test_docx_text_and_table_are_extracted(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invoice.docx"
            document = Document()
            document.add_paragraph("Supplier: DOCX Beauty LLC")
            document.add_paragraph("Invoice No: DOCX-1")
            document.add_paragraph("Invoice Date: 24 September 2026")
            table = document.add_table(rows=2, cols=5)
            for cell, value in zip(
                table.rows[0].cells,
                ["Description", "Qty", "UOM", "Unit Price", "Line Total"],
            ):
                cell.text = value
            for cell, value in zip(
                table.rows[1].cells,
                ["Repair Serum 50ml", "2", "EA", "20.00", "40.00"],
            ):
                cell.text = value
            document.save(path)
            result = extract_document(path)

        self.assertEqual(result["extraction_method"], "docx_text")
        self.assertEqual(result["invoice_number"], "DOCX-1")
        self.assertEqual(result["lines"][0]["description"], "Repair Serum 50ml")
        self.assertTrue(any("Embedded images" in warning for warning in result["warnings"]))

    def test_unsupported_type_is_explicit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invoice.svg"
            path.write_text("<svg/>", encoding="utf-8")
            with self.assertRaisesRegex(UnsupportedDocumentError, "Unsupported document type"):
                extract_document(path)

    def test_tesseract_staging_png_is_rgb(self) -> None:
        from PIL import Image

        staged: dict[str, object] = {}

        def inspect_staged_file(command: list[str], **_: object) -> SimpleNamespace:
            with Image.open(command[1]) as image:
                image.load()
                staged["mode"] = image.mode
                staged["pixel"] = image.getpixel((0, 0))
            return SimpleNamespace(
                returncode=0,
                stdout="level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n",
                stderr="",
            )

        source = Image.new("L", (600, 400), 127)
        try:
            with patch("backend.extraction.subprocess.run", side_effect=inspect_staged_file):
                _run_tesseract(source, timeout=5, max_pixels=2_000_000)
            self.assertEqual(staged["mode"], "RGB")
            self.assertEqual(staged["pixel"], (127, 127, 127))
            # Internal conversion and cleanup must not close the caller's image.
            self.assertEqual(source.getpixel((0, 0)), 127)
        finally:
            source.close()

    @unittest.skipUnless(PDFIUM_AVAILABLE, "pypdfium2 is required")
    def test_large_scanned_packet_preserves_single_pass_page_coverage(self) -> None:
        from pypdf import PdfWriter

        calls: list[int] = []

        def fake_ocr(*_: object, psm: int = 6, **__: object) -> tuple[str, float]:
            calls.append(psm)
            return "INVOICE", 90.0

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "large-scanned-packet.pdf"
            writer = PdfWriter()
            for _ in range(13):
                writer.add_blank_page(width=100, height=100)
            with path.open("wb") as handle:
                writer.write(handle)
            with (
                patch("backend.extraction.shutil.which", return_value="/usr/bin/tesseract"),
                patch("backend.extraction._run_tesseract", side_effect=fake_ocr),
            ):
                results = _ocr_pdf_pages(path, range(13), ExtractionLimits())

        self.assertEqual(len(results), 13)
        self.assertEqual(calls, [6] * 13)

    @unittest.skipUnless(OCR_AVAILABLE, "native Tesseract is not installed")
    def test_real_png_ocr_extracts_source_text(self) -> None:
        from PIL import Image, ImageDraw, ImageFont

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invoice.png"
            image = Image.new("RGB", (1800, 1100), "white")
            draw = ImageDraw.Draw(image)
            font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
            font = ImageFont.truetype(font_path, 34) if Path(font_path).exists() else ImageFont.load_default()
            draw.multiline_text((80, 70), "\n".join(_invoice_lines("OCR-100")), fill="black", font=font, spacing=14)
            image.save(path)
            image.close()
            result = extract_document(path)

        self.assertEqual(result["extraction_method"], "image_ocr")
        self.assertIn("OCR-100", result["raw_text"])
        self.assertEqual(result["invoice_number"], "OCR-100")
        self.assertGreaterEqual(len(result["lines"]), 1)
        self.assertEqual(result["pages"], 1)
        self.assertIsNotNone(result["page_texts"][0]["ocr_confidence"])
        self.assertTrue(any("English language" in warning for warning in result["warnings"]))

    @unittest.skipUnless(
        OCR_AVAILABLE and PDFIUM_AVAILABLE,
        "native Tesseract and pypdfium2 are required",
    )
    def test_real_multipage_scanned_pdf_uses_ocr_for_every_page(self) -> None:
        from PIL import Image, ImageDraw, ImageFont

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scanned.pdf"
            font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
            font = ImageFont.truetype(font_path, 36) if Path(font_path).exists() else ImageFont.load_default()
            images = []
            for lines in (
                ["INVOICE", "Supplier: OCR Beauty LLC", "Invoice No: SCAN-200"],
                ["Description Qty UOM Unit Price Line Total", "Cleanser 50ml 2 EA 25.00 50.00", "Grand Total AED 50.00"],
            ):
                image = Image.new("RGB", (1800, 1100), "white")
                ImageDraw.Draw(image).multiline_text((80, 80), "\n".join(lines), fill="black", font=font, spacing=20)
                images.append(image)
            images[0].save(path, "PDF", save_all=True, append_images=images[1:], resolution=150)
            for image in images:
                image.close()
            result = extract_document(path)

        self.assertEqual(result["pages"], 2)
        self.assertEqual(result["extraction_method"], "pdf_ocr")
        self.assertEqual([page["method"] for page in result["page_texts"]], ["ocr", "ocr"])
        self.assertIn("SCAN-200", result["raw_text"])
        self.assertEqual(result["invoice_number"], "SCAN-200")
        self.assertEqual(len(result["lines"]), 1)


if __name__ == "__main__":
    unittest.main()
